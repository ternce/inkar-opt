from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.app.services.references.stock_diagnostics import compare_stock_quantities, normalized_quantity_map


def _key(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().split())


def load_columns(path: Path, *, sku_header: str, quantity_header: str) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, ())
    by_name = {_key(value): index for index, value in enumerate(headers)}
    sku_index = by_name.get(_key(sku_header))
    quantity_index = by_name.get(_key(quantity_header))
    if sku_index is None or quantity_index is None:
        raise ValueError(
            f"{path}: required columns not found; sku={sku_header!r}, quantity={quantity_header!r}; "
            f"available={list(by_name)}"
        )
    values = []
    for row in rows:
        sku = row[sku_index] if sku_index < len(row) else None
        quantity = row[quantity_index] if quantity_index < len(row) else None
        values.append((sku, quantity))
    workbook.close()
    return normalized_quantity_map(values)


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare reference and generated XLSX stock quantities by normalized SKU.")
    parser.add_argument("reference_xlsx", type=Path)
    parser.add_argument("generated_xlsx", type=Path)
    parser.add_argument("--reference-sku-header", default="Материал")
    parser.add_argument("--reference-quantity-header", default="Остаток")
    parser.add_argument("--generated-sku-header", default="SKU")
    parser.add_argument("--generated-quantity-header", default="Остаток")
    parser.add_argument("--example-limit", type=int, default=100)
    args = parser.parse_args()
    expected = load_columns(
        args.reference_xlsx,
        sku_header=args.reference_sku_header,
        quantity_header=args.reference_quantity_header,
    )
    actual = load_columns(
        args.generated_xlsx,
        sku_header=args.generated_sku_header,
        quantity_header=args.generated_quantity_header,
    )
    report = compare_stock_quantities(expected, actual, example_limit=max(0, args.example_limit))
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["mismatchCount"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
