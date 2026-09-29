from __future__ import annotations

import hashlib
import io
import json
import logging
import math
import re
import shutil
import subprocess
import tempfile
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePath
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from ..models import (
    CompetitorPriceList,
    CompetitorPriceListItem,
    ManualPriceListImport,
    ManualPriceListImportError,
    PriceFormat,
    Product,
    VidmanAccount,
    VidmanCanonicalProduct,
    VidmanCompetitorPriceListSource,
    VidmanImportPage,
    VidmanImportRun,
    VidmanPriceList,
    VidmanProductMatch,
    VidmanRawCanonicalLink,
    VidmanRawItem,
)
from ..timezone import now_kz_naive
from .provisor_auto_refresh import new_owner_token, release_lock, try_acquire_lock
from .vidman_competitor_price_lists import (
    TRUSTED_VIDMAN_PRICE_STATUSES,
    build_vidman_competitor_price_list,
    vidman_competitor_source_key,
)
from .competitor_assignments import selected_price_format_ids_for_competitor_price_list
from .competitor_matching import rebuild_competitor_prices_for_selected
from .competitor_price_lists import sync_selected_competitor_configs
from .competitor_read_models import refresh_price_list_item_counters
from .percentile_preparation import enqueue_percentile_preparation
from .vidman_normalization import _canonical_from_parsed, parse_vidman_product, process_vidman_stage2
from .vidman_product_matching import (
    AUTO_MATCHED,
    MANUALLY_APPROVED,
    REVIEW_REQUIRED,
    build_product_indexes,
    decide_match,
    process_vidman_product_matches,
)
from .vidman_raw_collector import VidmanRawRow, make_row_hash, parse_expiry_date


MAX_VIDMAN_FILE_SIZE_BYTES = 15 * 1024 * 1024
MAX_VIDMAN_ROWS = 100_000
LIBREOFFICE_CONVERSION_TIMEOUT_SECONDS = 45
HEADER_SCAN_ROWS = 50
COMPLETENESS_MIN_CURRENT_ROWS = 100
COMPLETENESS_WARN_RATIO = Decimal("0.50")
logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class VidmanFileRow:
    row_number: int
    raw_name: str
    raw_manufacturer: str
    raw_expiry_text: str
    expiry_date: date | None
    raw_price_text: str
    price: Decimal
    raw_pack_qty: str
    pack_qty: Decimal | None
    raw_min_order: str
    min_order: Decimal | None
    raw_stock: str
    stock: Decimal | None
    file_account: str = ""
    file_main_id: str = ""
    raw: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class VidmanFileError:
    row_number: int | None
    field: str
    raw_value: str
    error_code: str
    message: str


@dataclass(frozen=True)
class ParsedVidmanFile:
    file_type: str
    checksum: str
    filename: str
    sheet: str
    total_rows: int
    empty_rows: int
    duplicate_rows: int
    valid_rows: list[VidmanFileRow]
    errors: list[VidmanFileError]
    headers: dict[str, str]
    detected_accounts: list[str]
    detected_main_ids: list[str]


HEADER_ALIASES: dict[str, set[str]] = {
    "name": {"name", "product name", "product_name", "наименование", "наименование товара", "название", "товар"},
    "manufacturer": {"manufacturer", "producer", "brand", "производитель", "изготовитель"},
    "expiry": {"expiry", "expiry date", "expiry_date", "shelf life", "срок годности", "срок"},
    "price": {"price", "goodsprice", "goods price", "цена", "цена с ндс", "цена тг"},
    "pack_qty": {"pack qty", "pack quantity", "qty package", "кол в уп", "количество в упаковке", "упаковка"},
    "min_order": {"min order", "minimum order", "мин заказ", "минимальный заказ"},
    "stock": {"stock", "stored", "остаток", "остатки"},
    "account": {"account", "account id", "login", "аккаунт", "логин"},
    "main_id": {"main id", "main_id", "price id", "price list id", "id прайса", "id прайс листа"},
    "price_list_name": {"price list", "price list name", "прайс", "прайс лист"},
}


def safe_upload_filename(value: object) -> str:
    text = str(value or "vidman.xlsx").replace("\\", "/")
    text = PurePath(text).name
    text = re.sub(r"[\x00-\x1f\x7f]+", "", text).strip(" .")
    return (text or "vidman.xlsx")[:255]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return str(value).replace("\ufeff", "").replace("\u00a0", " ").strip()


def _header(value: object) -> str:
    text = _text(value).casefold().replace("ё", "е")
    text = re.sub(r"[_./\\-]+", " ", text)
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-zа-я ]+", "", text)).strip()


def _header_map(row: Iterable[object]) -> dict[str, int] | None:
    normalized = [_header(value) for value in row]
    found: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_header(alias) for alias in aliases}
        index = next((idx for idx, value in enumerate(normalized) if value in normalized_aliases), None)
        if index is not None:
            found[field] = index
    return found if {"name", "price"}.issubset(found) else None


def _decimal(value: object, *, positive: bool = False) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            result = Decimal(str(value))
        except InvalidOperation:
            return None
    else:
        text = _text(value).replace(" ", "").replace("\u00a0", "")
        text = re.sub(r"[^0-9,.-]", "", text)
        if not text:
            return None
        if "," in text and "." in text:
            decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
            thousands_separator = "." if decimal_separator == "," else ","
            text = text.replace(thousands_separator, "").replace(decimal_separator, ".")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            result = Decimal(text)
        except InvalidOperation:
            return None
    if not result.is_finite() or (positive and result <= 0):
        return None
    return result.quantize(Decimal("0.0001"))


def _date_value(value: object) -> tuple[str, date | None]:
    if isinstance(value, datetime):
        return value.date().isoformat(), value.date()
    if isinstance(value, date):
        return value.isoformat(), value
    raw = _text(value)
    return raw, parse_expiry_date(raw)


def _xlsx_sheets(content: bytes):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"malformed .xlsx file: {exc}") from exc
    for sheet in workbook.worksheets:
        yield sheet.title, sheet.iter_rows(values_only=True)


class _LibreOfficeUnavailable(RuntimeError):
    pass


class _LibreOfficeConversionFailed(RuntimeError):
    pass


def _libreoffice_executable() -> str | None:
    return shutil.which("libreoffice") or shutil.which("soffice")


def _parser_failure(exc: BaseException) -> str:
    message = re.sub(r"\s+", " ", str(exc)).strip()
    return f"{type(exc).__name__}: {message or 'no error details'}"


def _convert_legacy_xls_to_xlsx(content: bytes) -> bytes:
    # The public parsers check this before selecting a reader. Keep the guard here
    # too so this safety boundary cannot accidentally be bypassed by a future caller.
    if len(content) > MAX_VIDMAN_FILE_SIZE_BYTES:
        raise _LibreOfficeConversionFailed("file is too large")

    executable = _libreoffice_executable()
    if not executable:
        raise _LibreOfficeUnavailable

    with tempfile.TemporaryDirectory(prefix="vidman-xls-") as temporary_directory:
        temp_dir = Path(temporary_directory)
        # Deliberately use a fixed server-generated name, never the uploaded filename.
        input_path = temp_dir / "upload.xls"
        output_path = temp_dir / "upload.xlsx"
        input_path.write_bytes(content)
        command = [
            executable,
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(temp_dir),
            str(input_path),
        ]
        try:
            completed = subprocess.run(
                command,
                shell=False,
                capture_output=True,
                text=True,
                timeout=LIBREOFFICE_CONVERSION_TIMEOUT_SECONDS,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise _LibreOfficeConversionFailed(
                f"conversion timed out after {LIBREOFFICE_CONVERSION_TIMEOUT_SECONDS} seconds"
            ) from exc
        except FileNotFoundError as exc:
            raise _LibreOfficeUnavailable from exc
        except OSError as exc:
            raise _LibreOfficeConversionFailed(f"LibreOffice could not be started ({type(exc).__name__})") from exc

        if completed.returncode != 0:
            raise _LibreOfficeConversionFailed(f"LibreOffice exited with code {completed.returncode}")
        if not output_path.is_file():
            raise _LibreOfficeConversionFailed("LibreOffice did not produce an .xlsx file")
        converted = output_path.read_bytes()
        if not converted:
            raise _LibreOfficeConversionFailed("LibreOffice produced an empty .xlsx file")
        return converted


def _xls_sheets(content: bytes):
    try:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        materialized_sheets: list[tuple[str, list[tuple[object, ...]]]] = []
        try:
            for sheet in workbook.sheets():
                materialized_rows: list[tuple[object, ...]] = []
                for row_index in range(sheet.nrows):
                    values: list[object] = []
                    for cell in sheet.row(row_index):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            values.append(xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode))
                        else:
                            values.append(cell.value)
                    materialized_rows.append(tuple(values))
                materialized_sheets.append((sheet.name, materialized_rows))
        finally:
            workbook.release_resources()
    except Exception as exc:
        logger.warning(
            "xlrd failed to parse an uploaded legacy .xls; attempting LibreOffice conversion: %s",
            _parser_failure(exc),
            exc_info=True,
        )
        try:
            converted = _convert_legacy_xls_to_xlsx(content)
        except _LibreOfficeUnavailable as conversion_exc:
            raise ValueError(
                "Legacy .xls parsing failed and LibreOffice conversion is unavailable."
            ) from conversion_exc
        except _LibreOfficeConversionFailed as conversion_exc:
            raise ValueError(
                f"Legacy .xls parsing failed ({_parser_failure(exc)}); "
                f"LibreOffice conversion failed ({conversion_exc})."
            ) from conversion_exc

        try:
            yield from _xlsx_sheets(converted)
        except Exception as converted_exc:
            raise ValueError(
                f"Legacy .xls parsing failed ({_parser_failure(exc)}); "
                f"LibreOffice conversion failed (converted .xlsx could not be parsed: "
                f"{_parser_failure(converted_exc)})."
            ) from converted_exc
        return
    for sheet_name, rows in materialized_sheets:
        yield sheet_name, iter(rows)


def parse_vidman_file(content: bytes, filename: str) -> ParsedVidmanFile:
    filename = safe_upload_filename(filename)
    if not content:
        raise ValueError("empty file")
    if len(content) > MAX_VIDMAN_FILE_SIZE_BYTES:
        raise ValueError("file is too large")
    suffix = filename.casefold().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "xlsx":
        file_type = "xlsx"
        sheets = _xlsx_sheets(content)
    elif suffix == "xls":
        file_type = "xls"
        sheets = _xls_sheets(content)
    else:
        raise ValueError("only .xlsx and .xls Vidman files are supported")

    selected_sheet = ""
    mapping: dict[str, int] | None = None
    header_values: tuple[object, ...] = ()
    data_rows: list[tuple[int, tuple[object, ...]]] = []
    for sheet_name, iterator in sheets:
        buffered: list[tuple[int, tuple[object, ...]]] = []
        header_position: int | None = None
        for row_number, raw_row in enumerate(iterator, start=1):
            row = tuple(raw_row or ())
            if row_number <= HEADER_SCAN_ROWS:
                candidate = _header_map(row)
                if candidate is not None and header_position is None:
                    mapping = candidate
                    header_position = row_number
                    header_values = row
                    selected_sheet = sheet_name
                    continue
            if header_position is not None and row_number > header_position:
                buffered.append((row_number, row))
                if len(buffered) > MAX_VIDMAN_ROWS:
                    raise ValueError("file contains too many rows")
        if header_position is not None:
            data_rows = buffered
            break
    if mapping is None:
        raise ValueError("required Vidman columns were not found: product name, price")

    headers = {field: _text(header_values[index]) if index < len(header_values) else field for field, index in mapping.items()}
    valid: list[VidmanFileRow] = []
    errors: list[VidmanFileError] = []
    empty_rows = 0
    duplicates = 0
    seen: set[tuple[str, ...]] = set()
    accounts: set[str] = set()
    main_ids: set[str] = set()
    total_rows = 0

    def cell(row: tuple[object, ...], field: str) -> object:
        index = mapping.get(field)
        return row[index] if index is not None and index < len(row) else None

    for row_number, row in data_rows:
        if not any(_text(value) for value in row):
            empty_rows += 1
            continue
        total_rows += 1
        name = _text(cell(row, "name"))
        manufacturer = _text(cell(row, "manufacturer"))
        raw_price = _text(cell(row, "price"))
        price = _decimal(cell(row, "price"), positive=True)
        expiry_raw, expiry = _date_value(cell(row, "expiry"))
        pack_raw = _text(cell(row, "pack_qty"))
        min_raw = _text(cell(row, "min_order"))
        stock_raw = _text(cell(row, "stock"))
        file_account = _text(cell(row, "account"))
        file_main_id = _text(cell(row, "main_id"))
        if file_account:
            accounts.add(file_account)
        if file_main_id:
            main_ids.add(file_main_id)
        if not name:
            errors.append(VidmanFileError(row_number, "name", "", "missing_name", "product name is required"))
            continue
        if price is None:
            errors.append(VidmanFileError(row_number, "price", raw_price, "invalid_price", "price must be positive"))
            continue
        signature = tuple(value.casefold() for value in (name, manufacturer, expiry_raw, raw_price, pack_raw, min_raw, stock_raw))
        if signature in seen:
            duplicates += 1
            continue
        seen.add(signature)
        valid.append(
            VidmanFileRow(
                row_number=row_number,
                raw_name=name,
                raw_manufacturer=manufacturer,
                raw_expiry_text=expiry_raw,
                expiry_date=expiry,
                raw_price_text=raw_price,
                price=price,
                raw_pack_qty=pack_raw,
                pack_qty=_decimal(cell(row, "pack_qty")),
                raw_min_order=min_raw,
                min_order=_decimal(cell(row, "min_order")),
                raw_stock=stock_raw,
                stock=_decimal(cell(row, "stock")),
                file_account=file_account,
                file_main_id=file_main_id,
                raw={"row": list(row), "headers": headers},
            )
        )

    return ParsedVidmanFile(
        file_type=file_type,
        checksum=hashlib.sha256(content).hexdigest(),
        filename=filename,
        sheet=selected_sheet,
        total_rows=total_rows,
        empty_rows=empty_rows,
        duplicate_rows=duplicates,
        valid_rows=valid,
        errors=errors,
        headers=headers,
        detected_accounts=sorted(accounts),
        detected_main_ids=sorted(main_ids),
    )


def _target(
    *, db: Session, competitor_price_list_id: int
) -> tuple[CompetitorPriceList, VidmanCompetitorPriceListSource, VidmanAccount, VidmanPriceList]:
    price_list = db.get(CompetitorPriceList, int(competitor_price_list_id))
    if price_list is None or str(price_list.source_type or "").strip().lower() != "vidman":
        raise ValueError("canonical Vidman competitor price list not found")
    account_id = int(str(price_list.account_id or "0"))
    main_id = int(str(price_list.external_price_list_id or "0"))
    expected_key = vidman_competitor_source_key(account_id, main_id)
    if not account_id or not main_id or price_list.source_key != expected_key:
        raise ValueError("legacy or non-canonical Vidman source identity is not supported")
    source = db.execute(
        select(VidmanCompetitorPriceListSource)
        .where(VidmanCompetitorPriceListSource.account_id == account_id)
        .where(VidmanCompetitorPriceListSource.main_id == main_id)
        .where(VidmanCompetitorPriceListSource.competitor_price_list_id == price_list.id)
    ).scalar_one_or_none()
    account = db.get(VidmanAccount, account_id)
    vidman_price_list = db.execute(
        select(VidmanPriceList)
        .where(VidmanPriceList.account_id == account_id)
        .where(VidmanPriceList.main_id == main_id)
    ).scalar_one_or_none()
    if source is None or account is None or vidman_price_list is None:
        raise ValueError("Vidman source mapping is incomplete")
    if not source.is_active:
        raise ValueError("Vidman source is inactive")
    return price_list, source, account, vidman_price_list


def _validate_file_target(parsed: ParsedVidmanFile, *, account: VidmanAccount, main_id: int) -> list[str]:
    warnings: list[str] = []
    main_values = {str(value).strip() for value in parsed.detected_main_ids if str(value).strip()}
    if main_values and main_values != {str(main_id)}:
        raise ValueError(f"file PLK id does not match selected target: expected {main_id}, found {sorted(main_values)}")
    account_values = {value.strip().casefold() for value in parsed.detected_accounts if value.strip()}
    account_aliases = {str(account.id), str(account.login or "").strip().casefold(), str(account.display_name or "").strip().casefold()} - {""}
    if account_values and not account_values.intersection(account_aliases):
        raise ValueError("file account does not match selected Vidman account")
    if not main_values:
        warnings.append("file_does_not_contain_plk_id; selected target will be used")
    if not account_values:
        warnings.append("file_does_not_contain_account; selected target will be used")
    return warnings


def _preview_matches(db: Session, rows: list[VidmanFileRow]) -> tuple[int, int, int]:
    parsed_rows = [(row, parse_vidman_product(row.raw_name, row.raw_manufacturer)) for row in rows]
    signatures = sorted({parsed.normalized_signature for _row, parsed in parsed_rows if parsed.normalized_signature})
    canonicals = {
        row.canonical_signature: row
        for row in db.execute(
            select(VidmanCanonicalProduct).where(VidmanCanonicalProduct.canonical_signature.in_(signatures or ["__none__"]))
        ).scalars()
    }
    matches = {
        int(row.canonical_product_id): row
        for row in db.execute(
            select(VidmanProductMatch).where(VidmanProductMatch.canonical_product_id.in_([item.id for item in canonicals.values()] or [-1]))
        ).scalars()
    }
    indexes = None
    matched = review = unmatched = 0
    for _row, parsed in parsed_rows:
        canonical = canonicals.get(parsed.normalized_signature)
        existing = matches.get(int(canonical.id)) if canonical is not None and canonical.id is not None else None
        status = str(existing.status) if existing is not None else ""
        if status in TRUSTED_VIDMAN_PRICE_STATUSES and existing.product_id is not None:
            matched += 1
            continue
        if status == REVIEW_REQUIRED:
            review += 1
            continue
        if not parsed.normalized_signature:
            unmatched += 1
            continue
        if indexes is None:
            indexes = build_product_indexes(db)
        decision = decide_match(canonical or _canonical_from_parsed(parsed), indexes)
        if decision.status == AUTO_MATCHED and decision.product_id is not None:
            matched += 1
        elif decision.status == REVIEW_REQUIRED:
            review += 1
        else:
            unmatched += 1
    return matched, review, unmatched


def _published_rows_count(db: Session, price_list: CompetitorPriceList) -> int:
    return int(
        price_list.items_count
        or db.scalar(
            select(func.count(CompetitorPriceListItem.id)).where(
                CompetitorPriceListItem.price_list_id == price_list.id
            )
        )
        or 0
    )


def preview_vidman_file(
    *, db: Session, competitor_price_list_id: int, content: bytes, filename: str
) -> dict[str, Any]:
    price_list, source, account, vidman_price_list = _target(db=db, competitor_price_list_id=competitor_price_list_id)
    parsed = parse_vidman_file(content, filename)
    warnings = _validate_file_target(parsed, account=account, main_id=int(vidman_price_list.main_id))
    matched, review, unmatched = _preview_matches(db, parsed.valid_rows)
    current_rows = _published_rows_count(db, price_list)
    incoming_rows = len(parsed.valid_rows)
    suspicious_drop = bool(
        current_rows >= COMPLETENESS_MIN_CURRENT_ROWS
        and Decimal(incoming_rows) < Decimal(current_rows) * COMPLETENESS_WARN_RATIO
    )
    if suspicious_drop:
        warnings.append(f"suspicious_row_count_drop:{current_rows}->{incoming_rows}")
    if parsed.errors:
        warnings.append("invalid_rows_will_not_be_imported")
    return {
        "ok": True,
        "mode": "dry_run",
        "filename": parsed.filename,
        "fileType": parsed.file_type,
        "checksum": parsed.checksum,
        "confirmationToken": parsed.checksum,
        "sheet": parsed.sheet,
        "target": {
            "competitorPriceListId": int(price_list.id),
            "sourceKey": price_list.source_key,
            "accountId": int(account.id),
            "accountLogin": account.login,
            "mainId": int(vidman_price_list.main_id),
            "priceListName": vidman_price_list.name,
            "updateMode": source.update_mode or "auto",
        },
        "totalRows": parsed.total_rows,
        "emptyRows": parsed.empty_rows,
        "validRows": incoming_rows,
        "invalidRows": len(parsed.errors),
        "duplicateRows": parsed.duplicate_rows,
        "knownMatchedRows": matched,
        "reviewRows": review,
        "unmatchedRows": unmatched,
        "currentPublishedRows": current_rows,
        "requiresCompletenessOverride": suspicious_drop,
        "warnings": warnings,
        "headers": parsed.headers,
        "errors": [_error_payload(item) for item in parsed.errors[:100]],
    }


def _error_payload(error: VidmanFileError) -> dict[str, Any]:
    return {
        "rowNumber": error.row_number,
        "field": error.field,
        "rawValue": error.raw_value,
        "errorCode": error.error_code,
        "message": error.message,
    }


def _history(
    *,
    db: Session,
    parsed: ParsedVidmanFile | None,
    price_list_id: int,
    source_key: str,
    filename: str,
    requested_by: str,
    status: str,
    matched: int = 0,
    persisted: int = 0,
    preserved: bool = False,
    error: str = "",
    metadata: dict[str, Any] | None = None,
) -> ManualPriceListImport:
    now = now_kz_naive()
    row = ManualPriceListImport(
        competitor_price_list_id=price_list_id,
        source_key=source_key,
        original_filename=safe_upload_filename(filename),
        file_type=parsed.file_type if parsed else "",
        file_checksum=parsed.checksum if parsed else "",
        detected_sheet=parsed.sheet if parsed else "",
        status=status,
        total_rows=parsed.total_rows if parsed else 0,
        empty_rows=parsed.empty_rows if parsed else 0,
        valid_rows=len(parsed.valid_rows) if parsed else 0,
        invalid_rows=len(parsed.errors) if parsed else 0,
        duplicate_rows=parsed.duplicate_rows if parsed else 0,
        matched_rows=matched,
        unmatched_rows=max((len(parsed.valid_rows) if parsed else 0) - matched, 0),
        persisted_rows=persisted,
        preserved_previous_snapshot=preserved,
        requested_by=requested_by,
        started_at=now,
        finished_at=now,
        error_summary=error,
        metadata_json=json.dumps(metadata or {}, ensure_ascii=False, default=str),
    )
    db.add(row)
    db.flush()
    for item in (parsed.errors if parsed else [])[:1000]:
        db.add(
            ManualPriceListImportError(
                import_id=row.id,
                row_number=item.row_number,
                field=item.field,
                raw_value=item.raw_value[:1000],
                error_code=item.error_code,
                message=item.message,
            )
        )
    return row


def import_vidman_file(
    *,
    db: Session,
    competitor_price_list_id: int,
    content: bytes,
    filename: str,
    expected_checksum: str,
    allow_incomplete: bool,
    requested_by: str,
) -> dict[str, Any]:
    price_list, source, account, vidman_price_list = _target(db=db, competitor_price_list_id=competitor_price_list_id)
    parsed = parse_vidman_file(content, filename)
    if not expected_checksum or parsed.checksum != str(expected_checksum).strip().lower():
        raise ValueError("file checksum does not match the confirmed preview")
    warnings = _validate_file_target(parsed, account=account, main_id=int(vidman_price_list.main_id))
    if not parsed.valid_rows:
        raise ValueError("file has no valid Vidman rows")
    current_rows = _published_rows_count(db, price_list)
    suspicious_drop = bool(
        current_rows >= COMPLETENESS_MIN_CURRENT_ROWS
        and Decimal(len(parsed.valid_rows)) < Decimal(current_rows) * COMPLETENESS_WARN_RATIO
    )
    if suspicious_drop and not allow_incomplete:
        raise ValueError("incoming file is suspiciously smaller than the published source; explicit override is required")

    lock_name = f"vidman_manual_import:{account.id}:{vidman_price_list.main_id}"
    owner_token = new_owner_token()
    if not try_acquire_lock(
        db,
        name=lock_name,
        lock_type="vidman_manual_import",
        owner_token=owner_token,
        lease=timedelta(minutes=30),
        metadata={"requested_by": requested_by, "checksum": parsed.checksum},
    ):
        raise ValueError("another import is already running for this Vidman PLK")

    import_run_id: int | None = None
    try:
        run = VidmanImportRun(
            account_id=int(account.id),
            status="running",
            total_plks=1,
            completed_plks=0,
            failed_plks=0,
            total_pages=1,
            total_rows=len(parsed.valid_rows),
            metadata_json=json.dumps(
                {"origin": "manual_file", "filename": parsed.filename, "checksum": parsed.checksum, "requestedBy": requested_by},
                ensure_ascii=False,
            ),
        )
        db.add(run)
        db.flush()
        import_run_id = int(run.id)
        page = VidmanImportPage(
            import_run_id=run.id,
            price_list_id=vidman_price_list.id,
            main_id=vidman_price_list.main_id,
            page_number=1,
            status="success",
            rows_count=len(parsed.valid_rows),
            attempts=1,
            started_at=now_kz_naive(),
            finished_at=now_kz_naive(),
        )
        db.add(page)
        for row in parsed.valid_rows:
            raw_row = VidmanRawRow(
                row_number=row.row_number,
                raw_name=row.raw_name,
                raw_manufacturer=row.raw_manufacturer,
                raw_expiry_text=row.raw_expiry_text,
                expiry_date=row.expiry_date,
                raw_price_text=row.raw_price_text,
                price=row.price,
                raw_pack_qty=row.raw_pack_qty,
                pack_qty=row.pack_qty,
                raw_min_order=row.raw_min_order,
                min_order=row.min_order,
                raw_stock=row.raw_stock,
                stock=row.stock,
                raw_html=json.dumps(row.raw, ensure_ascii=False, default=str),
            )
            db.add(
                VidmanRawItem(
                    import_run_id=run.id,
                    account_id=account.id,
                    price_list_id=vidman_price_list.id,
                    main_id=vidman_price_list.main_id,
                    page_number=1,
                    row_number=row.row_number,
                    raw_name=row.raw_name,
                    raw_manufacturer=row.raw_manufacturer,
                    raw_expiry_text=row.raw_expiry_text,
                    expiry_date=row.expiry_date,
                    raw_price_text=row.raw_price_text,
                    price=row.price,
                    raw_pack_qty=row.raw_pack_qty,
                    pack_qty=row.pack_qty,
                    raw_min_order=row.raw_min_order,
                    min_order=row.min_order,
                    raw_stock=row.raw_stock,
                    stock=row.stock,
                    raw_html=raw_row.raw_html,
                    row_hash=make_row_hash(
                        import_run_id=run.id,
                        account_id=account.id,
                        main_id=vidman_price_list.main_id,
                        page_number=1,
                        row=raw_row,
                    ),
                )
            )
        run.status = "success"
        run.completed_plks = 1
        run.finished_at = now_kz_naive()
        db.flush()

        process_vidman_stage2(db, account_id=account.id, main_id=vidman_price_list.main_id, only_unprocessed=True)
        canonical_ids = sorted(
            {
                int(value)
                for value in db.execute(
                    select(VidmanRawCanonicalLink.canonical_product_id)
                    .join(VidmanRawItem, VidmanRawItem.id == VidmanRawCanonicalLink.raw_item_id)
                    .where(VidmanRawItem.import_run_id == import_run_id)
                ).scalars()
            }
        )
        for canonical_id in canonical_ids:
            process_vidman_product_matches(db, canonical_id=canonical_id, only_unmatched=True, apply=True)

        summary = build_vidman_competitor_price_list(
            db=db,
            account_id=account.id,
            main_id=vidman_price_list.main_id,
            price_format_code=source.price_format_code,
            import_run_id=import_run_id,
            apply=True,
            require_active=True,
            commit=False,
        )
        if summary.skipped_reason:
            raise ValueError(f"Vidman publication rejected the snapshot: {summary.skipped_reason}")
        source.update_mode = "manual"
        source.updated_at = now_kz_naive()
        history = _history(
            db=db,
            parsed=parsed,
            price_list_id=int(price_list.id),
            source_key=price_list.source_key,
            filename=parsed.filename,
            requested_by=requested_by,
            status="partial_success" if parsed.errors else "success",
            matched=int(summary.trusted_match_rows),
            persisted=int(summary.rows_written),
            metadata={
                "origin": "manual_vidman_file",
                "accountId": int(account.id),
                "mainId": int(vidman_price_list.main_id),
                "vidmanImportRunId": import_run_id,
                "canonicalSourceKey": price_list.source_key,
                "warnings": warnings,
            },
        )
        db.commit()
        try:
            for price_format_id in selected_price_format_ids_for_competitor_price_list(
                db=db, competitor_price_list_id=int(price_list.id)
            ):
                enqueue_percentile_preparation(
                    db=db,
                    price_format_id=price_format_id,
                    reason="vidman_manual_file_imported",
                )
        except Exception as exc:
            db.rollback()
            warnings.append("percentile_preparation_enqueue_failed")
            logger.exception("Failed to enqueue percentile preparation after Vidman manual import: %s", exc)
        return {
            "ok": True,
            "status": history.status,
            "id": int(price_list.id),
            "sourceKey": price_list.source_key,
            "updateMode": "manual",
            "importId": int(history.id),
            "vidmanImportRunId": import_run_id,
            "rowsWritten": int(summary.rows_written),
            "trustedMatchedRows": int(summary.trusted_match_rows),
            "reviewOrUnmatchedRows": int(summary.unresolved_match_rows),
            "warnings": warnings,
        }
    except Exception as exc:
        db.rollback()
        if import_run_id is not None:
            failed_run = db.get(VidmanImportRun, import_run_id)
            if failed_run is not None:
                failed_run.status = "error"
                failed_run.failed_plks = 1
                failed_run.error_message = str(exc)[:2000]
                failed_run.finished_at = now_kz_naive()
        _history(
            db=db,
            parsed=parsed,
            price_list_id=int(price_list.id),
            source_key=price_list.source_key,
            filename=parsed.filename,
            requested_by=requested_by,
            status="error",
            preserved=True,
            error=str(exc)[:2000],
            metadata={"origin": "manual_vidman_file", "accountId": account.id, "mainId": vidman_price_list.main_id, "vidmanImportRunId": import_run_id},
        )
        db.commit()
        raise
    finally:
        try:
            release_lock(db, name=lock_name, owner_token=owner_token)
        except Exception:
            db.rollback()


def set_vidman_update_mode(*, db: Session, competitor_price_list_id: int, update_mode: str) -> dict[str, Any]:
    mode = str(update_mode or "").strip().lower()
    if mode not in {"auto", "manual"}:
        raise ValueError("update_mode must be auto or manual")
    price_list = db.get(CompetitorPriceList, int(competitor_price_list_id))
    if price_list is None or str(price_list.source_type or "").lower() not in {"vidman", "manual_vidman"}:
        raise ValueError("Vidman price list not found")
    price_list.update_mode = mode
    price_list.updated_at = now_kz_naive()
    if price_list.source_type == "vidman":
        try:
            _price_list, source, _account, _vidman_price_list = _target(
                db=db, competitor_price_list_id=competitor_price_list_id
            )
            source.update_mode = mode
            source.updated_at = now_kz_naive()
        except ValueError:
            pass
    db.commit()
    return {"ok": True, "id": int(price_list.id), "sourceKey": price_list.source_key, "updateMode": mode}


# Client workbooks exist in two pivot variants: the original supplier-per-column
# report and a newer price/stock-pair report. Both publish directly to canonical
# competitor lists because Product.code matching is authoritative here.
MULTI_SOURCE_TYPE = "manual_vidman"
PRICE_PREFIXES = ("цена ", "price ")
STOCK_PREFIXES = ("остаток ", "stock ")
PIVOT_CODE_HEADERS = {_header("\u041a\u043e\u0434"), "sku", "code"}
PIVOT_NAME_HEADERS = {_header("\u041d\u0430\u0438\u043c\u0435\u043d\u043e\u0432\u0430\u043d\u0438\u0435"), "name", "product name"}
PIVOT_MANUFACTURER_HEADERS = {_header("\u041f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c"), "manufacturer", "producer"}


@dataclass(frozen=True)
class VidmanPivotRow:
    row_number: int
    primary_sku: str
    primary_sku_numeric: bool
    fallback_sku: str
    distributor_sku: str
    name: str
    expiry: str
    manufacturer: str
    price: Decimal
    stock: Decimal | None


@dataclass
class DetectedVidmanSource:
    name: str
    stable_key: str
    price_header: str
    stock_header: str
    rows: list[VidmanPivotRow] = field(default_factory=list)
    errors: list[VidmanFileError] = field(default_factory=list)


@dataclass
class ParsedVidmanPivot:
    filename: str
    file_type: str
    checksum: str
    sheet: str
    header_row: int
    format_variant: str
    total_rows: int
    empty_rows: int
    sources: list[DetectedVidmanSource]


def _sku_cell(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return str(int(value)) if value.is_integer() else str(value).strip()
    # Text cells retain significant leading zeroes. Whitespace is not a valid
    # part of an internal code in the current product catalogue.
    return re.sub(r"\s+", "", str(value).replace("\ufeff", "").strip())


def _source_name_from_header(value: object, prefixes: tuple[str, ...]) -> str:
    raw = _text(value)
    folded = raw.casefold()
    for prefix in prefixes:
        if folded.startswith(prefix):
            return raw[len(prefix):].strip()
    return ""


def _stable_file_source_key(*, context: str, source_name: str) -> str:
    normalized_context = re.sub(r"\s+", " ", str(context or "global").strip().casefold())
    normalized_name = re.sub(r"\s+", " ", str(source_name or "").strip().casefold())
    digest = hashlib.sha256(f"{normalized_context}\0{normalized_name}".encode("utf-8")).hexdigest()[:20]
    slug = re.sub(r"[^0-9a-zа-я]+", "-", normalized_name, flags=re.IGNORECASE).strip("-")[:48] or "source"
    return f"vidman-file:{digest}:{slug}"


def _numeric_sku_cell(value: object) -> bool:
    return (
        not isinstance(value, bool)
        and isinstance(value, (int, float))
        and math.isfinite(float(value))
        and float(value).is_integer()
    )


def _pivot_header_diagnostics(
    candidates: list[tuple[str, int, tuple[object, ...]]],
) -> str:
    ranked: list[tuple[int, str, int, list[str], list[str], list[str]]] = []
    for sheet_name, row_number, row in candidates:
        normalized = [_header(value) for value in row]
        fixed: list[str] = []
        if any(value in PIVOT_CODE_HEADERS for value in normalized):
            fixed.append("Код/SKU")
        if any(value in PIVOT_NAME_HEADERS for value in normalized):
            fixed.append("Наименование")
        manufacturer_index = next(
            (index for index, value in enumerate(normalized) if value in PIVOT_MANUFACTURER_HEADERS),
            None,
        )
        if manufacturer_index is not None:
            fixed.append("Производитель")
        supplier_headers = (
            [_text(value) for value in row[manufacturer_index + 1 :] if _text(value)]
            if manufacturer_index is not None
            else []
        )
        visible = [_text(value)[:60] for value in row if _text(value)][:10]
        ranked.append((len(fixed), sheet_name, row_number, fixed, supplier_headers[:10], visible))
    ranked.sort(key=lambda item: (item[0], len(item[5])), reverse=True)
    best = ranked[0] if ranked else (0, "", 0, [], [], [])
    descriptions = [
        f"{sheet}!row {row_number} [{', '.join(visible)}]"
        for _score, sheet, row_number, _fixed, _suppliers, visible in ranked[:3]
        if visible
    ]
    return (
        f"header candidates: {'; '.join(descriptions) or 'none'}; "
        f"detected fixed metadata columns: {', '.join(best[3]) or 'none'}; "
        f"supplier columns found: {', '.join(best[4]) or 'none'}"
    )


def parse_vidman_pivot_file(content: bytes, filename: str, *, context: str) -> ParsedVidmanPivot:
    filename = safe_upload_filename(filename)
    if not content:
        raise ValueError("empty file")
    if len(content) > MAX_VIDMAN_FILE_SIZE_BYTES:
        raise ValueError("file is too large")
    suffix = filename.casefold().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "xlsx":
        sheets = _xlsx_sheets(content)
    elif suffix == "xls":
        sheets = _xls_sheets(content)
    else:
        raise ValueError("only .xlsx and .xls Vidman files are supported")

    selected_sheet = ""
    header_row = 0
    format_variant = ""
    header: tuple[object, ...] = ()
    rows: list[tuple[int, tuple[object, ...]]] = []
    scanned_candidates: list[tuple[str, int, tuple[object, ...]]] = []
    for sheet_name, iterator in sheets:
        buffered = list(iterator)
        for index, candidate in enumerate(buffered[:HEADER_SCAN_ROWS]):
            candidate_tuple = tuple(candidate)
            if any(_text(value) for value in candidate_tuple):
                scanned_candidates.append((sheet_name, index + 1, candidate_tuple))
            normalized = [_header(value) for value in candidate_tuple]
            has_supplier_layout = (
                any(value in PIVOT_CODE_HEADERS for value in normalized)
                and any(value in PIVOT_NAME_HEADERS for value in normalized)
                and any(value in PIVOT_MANUFACTURER_HEADERS for value in normalized)
            )
            has_pair_layout = "sku" in normalized and any(
                value in {"goodsid", "goods id"} for value in normalized
            )
            if has_supplier_layout or has_pair_layout:
                selected_sheet = sheet_name
                header_row = index + 1
                format_variant = "supplier_columns" if has_supplier_layout else "price_stock_pairs"
                header = candidate_tuple
                rows = [(row_number, tuple(row)) for row_number, row in enumerate(buffered[index + 1 :], start=index + 2)]
                break
        if header:
            break
    if not header:
        raise ValueError(f"Vidman pivot table header was not found; {_pivot_header_diagnostics(scanned_candidates)}")
    if len(rows) > MAX_VIDMAN_ROWS:
        raise ValueError("file contains too many rows")

    normalized_headers = [_header(value) for value in header]
    detected: list[tuple[int, int | None, DetectedVidmanSource]] = []
    if format_variant == "supplier_columns":
        primary_index = next(index for index, value in enumerate(normalized_headers) if value in PIVOT_CODE_HEADERS)
        fallback_index = None
        distributor_index = None
        name_index = next(index for index, value in enumerate(normalized_headers) if value in PIVOT_NAME_HEADERS)
        expiry_index = None
        manufacturer_index = next(
            index for index, value in enumerate(normalized_headers) if value in PIVOT_MANUFACTURER_HEADERS
        )
        for index in range(manufacturer_index + 1, len(header)):
            source_name = re.sub(r"\s+", " ", _text(header[index])).strip()
            if source_name:
                detected.append(
                    (
                        index,
                        None,
                        DetectedVidmanSource(
                            name=source_name,
                            stable_key=_stable_file_source_key(context=context, source_name=source_name),
                            price_header=source_name,
                            stock_header="",
                        ),
                    )
                )
    else:
        primary_index = normalized_headers.index("sku")
        fallback_index = next(i for i, value in enumerate(normalized_headers) if value in {"goodsid", "goods id"})
        distributor_index = next((i for i, value in enumerate(normalized_headers) if value in {"sku дистрибьютора", "distributor sku"}), None)
        name_index = next((i for i, value in enumerate(normalized_headers) if value in {"название", "name", "product name"}), None)
        expiry_index = next((i for i, value in enumerate(normalized_headers) if value in {"срок годности", "expiry", "expiry date"}), None)
        manufacturer_index = next((i for i, value in enumerate(normalized_headers) if value in {"производитель дистрибьютора", "manufacturer"}), None)
        stock_by_name: dict[str, tuple[int, str]] = {}
        for index, value in enumerate(header):
            source_name = _source_name_from_header(value, STOCK_PREFIXES)
            if source_name:
                stock_by_name[source_name.casefold()] = (index, _text(value))
        for index, value in enumerate(header):
            source_name = _source_name_from_header(value, PRICE_PREFIXES)
            if not source_name:
                continue
            stock = stock_by_name.get(source_name.casefold())
            detected.append(
                (
                    index,
                    stock[0] if stock else None,
                    DetectedVidmanSource(
                        name=source_name,
                        stable_key=_stable_file_source_key(context=context, source_name=source_name),
                        price_header=_text(value),
                        stock_header=stock[1] if stock else "",
                    ),
                )
            )
    if not detected:
        raise ValueError(
            f"no supplier price columns were found after the fixed metadata columns; "
            f"{_pivot_header_diagnostics([(selected_sheet, header_row, header)])}"
        )

    empty_rows = 0
    total_rows = 0
    for row_number, row in rows:
        if not any(_text(value) for value in row):
            empty_rows += 1
            continue
        total_rows += 1
        raw_primary_sku = row[primary_index] if primary_index < len(row) else None
        primary_sku = _sku_cell(raw_primary_sku)
        fallback_sku = _sku_cell(row[fallback_index] if fallback_index is not None and fallback_index < len(row) else None)
        for price_index, stock_index, source in detected:
            raw_price = row[price_index] if price_index < len(row) else None
            if not _text(raw_price):
                continue
            price = _decimal(raw_price, positive=True)
            if price is None:
                source.errors.append(VidmanFileError(row_number, source.price_header, _text(raw_price), "invalid_price", "price must be positive"))
                continue
            if not primary_sku and not fallback_sku:
                source.errors.append(
                    VidmanFileError(
                        row_number,
                        _text(header[primary_index]) or "SKU",
                        "",
                        "missing_sku",
                        "product code is empty",
                    )
                )
                continue
            source.rows.append(
                VidmanPivotRow(
                    row_number=row_number,
                    primary_sku=primary_sku,
                    primary_sku_numeric=(
                        format_variant == "supplier_columns" and _numeric_sku_cell(raw_primary_sku)
                    ),
                    fallback_sku=fallback_sku,
                    distributor_sku=_sku_cell(row[distributor_index]) if distributor_index is not None and distributor_index < len(row) else "",
                    name=_text(row[name_index]) if name_index is not None and name_index < len(row) else "",
                    expiry=_text(row[expiry_index]) if expiry_index is not None and expiry_index < len(row) else "",
                    manufacturer=_text(row[manufacturer_index]) if manufacturer_index is not None and manufacturer_index < len(row) else "",
                    price=price,
                    stock=_decimal(row[stock_index]) if stock_index is not None and stock_index < len(row) else None,
                )
            )
    return ParsedVidmanPivot(
        filename=filename,
        file_type=suffix,
        checksum=hashlib.sha256(content).hexdigest(),
        sheet=selected_sheet,
        header_row=header_row,
        format_variant=format_variant,
        total_rows=total_rows,
        empty_rows=empty_rows,
        sources=[item[2] for item in detected],
    )


def _product_code_index(
    db: Session,
) -> tuple[dict[str, Product], dict[str, Product], dict[str, Product], dict[int, Product]]:
    exact: dict[str, Product] = {}
    folded: dict[str, Product] = {}
    numeric: dict[str, Product] = {}
    by_goods_id: dict[int, Product] = {}
    folded_duplicates: set[str] = set()
    numeric_duplicates: set[str] = set()
    for product in db.execute(select(Product)).scalars():
        code = _sku_cell(product.code)
        if not code:
            continue
        exact[code] = product
        key = code.casefold()
        if key in folded:
            folded_duplicates.add(key)
        else:
            folded[key] = product
        if re.fullmatch(r"\d+", code):
            numeric_key = str(int(code))
            if numeric_key in numeric:
                numeric_duplicates.add(numeric_key)
            else:
                numeric[numeric_key] = product
        if product.provisor_goods_id is not None:
            by_goods_id[int(product.provisor_goods_id)] = product
    for key in folded_duplicates:
        folded.pop(key, None)
    for key in numeric_duplicates:
        numeric.pop(key, None)
    return exact, folded, numeric, by_goods_id


def _match_pivot_row(
    row: VidmanPivotRow,
    exact: dict[str, Product],
    folded: dict[str, Product],
    numeric: dict[str, Product],
    by_goods_id: dict[int, Product],
) -> tuple[Product | None, str, str]:
    if row.primary_sku:
        product = exact.get(row.primary_sku) or folded.get(row.primary_sku.casefold())
        if product is not None:
            return product, row.primary_sku, "SKU"
        if row.primary_sku_numeric and re.fullmatch(r"\d+", row.primary_sku):
            product = numeric.get(str(int(row.primary_sku)))
            if product is not None:
                return product, str(product.code), "SKU"
    if row.fallback_sku and re.fullmatch(r"\d+", row.fallback_sku):
        product = by_goods_id.get(int(row.fallback_sku))
        if product is not None:
            return product, row.fallback_sku, "goodsId"
    return None, row.primary_sku or row.fallback_sku, ""


def _existing_file_source(db: Session, stable_key: str) -> CompetitorPriceList | None:
    return db.execute(
        select(CompetitorPriceList)
        .where(CompetitorPriceList.source_type == MULTI_SOURCE_TYPE)
        .where(CompetitorPriceList.source_key == stable_key)
        .order_by(CompetitorPriceList.id.asc())
    ).scalars().first()


def preview_vidman_multi_source_file(
    *, db: Session, price_format: PriceFormat, branch: str, content: bytes, filename: str
) -> dict[str, Any]:
    target_branch = str(branch or "").strip()
    if not target_branch:
        raise ValueError("Не выбран филиал для импорта Vidman-файла.")
    parsed = parse_vidman_pivot_file(content, filename, context=target_branch)
    exact, folded, numeric, by_goods_id = _product_code_index(db)
    source_reports: list[dict[str, Any]] = []
    totals = {"validRows": 0, "matchedRows": 0, "unmatchedRows": 0, "invalidRows": 0, "duplicateRows": 0}
    matched_product_ids: set[int] = set()
    unmatched_codes: set[str] = set()
    invalid_prices = 0
    for source in parsed.sources:
        match_results = [
            (row, _match_pivot_row(row, exact, folded, numeric, by_goods_id))
            for row in source.rows
        ]
        matched = sum(1 for _row, result in match_results if result[0] is not None)
        matched_product_ids.update(
            int(result[0].id) for _row, result in match_results if result[0] is not None
        )
        unmatched_codes.update(
            row.primary_sku or row.fallback_sku
            for row, result in match_results
            if result[0] is None and (row.primary_sku or row.fallback_sku)
        )
        invalid_prices += sum(1 for error in source.errors if error.error_code == "invalid_price")
        unmatched = len(source.rows) - matched
        duplicate_rows = len(source.rows) - len(
            {(row.primary_sku, row.fallback_sku, row.price, row.stock) for row in source.rows}
        )
        existing = _existing_file_source(db, source.stable_key)
        current_rows = _published_rows_count(db, existing) if existing is not None else 0
        suspicious = bool(current_rows >= COMPLETENESS_MIN_CURRENT_ROWS and Decimal(matched) < Decimal(current_rows) * COMPLETENESS_WARN_RATIO)
        source_reports.append(
            {
                "name": source.name,
                "stableKey": source.stable_key,
                "priceHeader": source.price_header,
                "stockHeader": source.stock_header,
                "validPriceRows": len(source.rows),
                "status": "ready" if source.rows else "skipped_empty",
                "skuMatchedRows": matched,
                "skuUnmatchedRows": unmatched,
                "invalidRows": len(source.errors),
                "duplicateRows": duplicate_rows,
                "currentPublishedRows": current_rows,
                "competitorPriceListId": int(existing.id) if existing is not None else None,
                "updateMode": existing.update_mode if existing is not None else "manual",
                "requiresCompletenessOverride": suspicious,
                "errors": [_error_payload(error) for error in source.errors[:25]],
            }
        )
        totals["validRows"] += len(source.rows)
        totals["matchedRows"] += matched
        totals["unmatchedRows"] += unmatched
        totals["invalidRows"] += len(source.errors)
        totals["duplicateRows"] += duplicate_rows
    return {
        "ok": True,
        "mode": "dry_run",
        "branch": target_branch,
        "filename": parsed.filename,
        "fileType": parsed.file_type,
        "sheet": parsed.sheet,
        "headerRow": parsed.header_row,
        "formatVariant": parsed.format_variant,
        "checksum": parsed.checksum,
        "confirmationToken": parsed.checksum,
        "detectedSources": len(source_reports),
        "sources": source_reports,
        "totalRows": parsed.total_rows,
        "emptyRows": parsed.empty_rows,
        "matchedProducts": len(matched_product_ids),
        "unmatchedCodes": len(unmatched_codes),
        "unmatchedProductCodes": sorted(unmatched_codes)[:100],
        "invalidPrices": invalid_prices,
        **totals,
        "requiresCompletenessOverride": any(row["requiresCompletenessOverride"] for row in source_reports),
    }


def import_vidman_multi_source_file(
    *,
    db: Session,
    price_format: PriceFormat,
    branch: str,
    content: bytes,
    filename: str,
    expected_checksum: str,
    allow_incomplete: bool,
    requested_by: str,
) -> dict[str, Any]:
    target_branch = str(branch or "").strip()
    if not target_branch:
        raise ValueError("Не выбран филиал для импорта Vidman-файла.")
    parsed = parse_vidman_pivot_file(content, filename, context=target_branch)
    if not expected_checksum or parsed.checksum != str(expected_checksum).strip().lower():
        raise ValueError("file checksum does not match the confirmed preview")
    exact, folded, numeric, by_goods_id = _product_code_index(db)
    batch_id = f"vidman-file-{uuid.uuid4().hex}"
    owner_token = new_owner_token()
    lock_name = f"vidman_multi_file:{target_branch}"
    if not try_acquire_lock(
        db,
        name=lock_name,
        lock_type="vidman_manual_import",
        owner_token=owner_token,
        lease=timedelta(minutes=30),
        metadata={"requested_by": requested_by, "checksum": parsed.checksum},
    ):
        raise ValueError("another Vidman file import is already running for this branch")
    results: list[dict[str, Any]] = []
    affected_price_formats: set[int] = set()
    try:
        for source in parsed.sources:
            existing = _existing_file_source(db, source.stable_key)
            if not source.rows:
                now = now_kz_naive()
                preserved = existing is not None
                skip_status = (
                    "skipped_empty_existing_preserved" if preserved else "skipped_empty"
                )
                message = "No valid prices found; source was not created or updated."
                history = ManualPriceListImport(
                    competitor_price_list_id=existing.id if existing is not None else None,
                    source_key=source.stable_key,
                    original_filename=parsed.filename,
                    file_type=parsed.file_type,
                    file_checksum=parsed.checksum,
                    detected_sheet=parsed.sheet,
                    status=skip_status,
                    total_rows=parsed.total_rows,
                    empty_rows=parsed.empty_rows,
                    valid_rows=0,
                    invalid_rows=len(source.errors),
                    duplicate_rows=0,
                    matched_rows=0,
                    unmatched_rows=0,
                    persisted_rows=0,
                    preserved_previous_snapshot=preserved,
                    requested_by=requested_by,
                    started_at=now,
                    finished_at=now,
                    error_summary=message,
                    metadata_json=json.dumps(
                        {
                            "origin": "manual_vidman_multi_source",
                            "batchId": batch_id,
                            "sourceName": source.name,
                            "stableKey": source.stable_key,
                            "outcome": skip_status,
                            "existingSourcePreserved": preserved,
                            "validPriceRows": 0,
                        },
                        ensure_ascii=False,
                    ),
                )
                db.add(history)
                db.flush()
                results.append(
                    {
                        "name": source.name,
                        "stableKey": source.stable_key,
                        "competitorPriceListId": int(existing.id) if existing is not None else None,
                        "importId": int(history.id),
                        "status": skip_status,
                        "message": message,
                        "validPriceRows": 0,
                        "matchedRows": 0,
                        "unmatchedRows": 0,
                        "rowsWritten": 0,
                        "preservedPreviousSnapshot": preserved,
                    }
                )
                continue
            matched_rows = []
            unmatched = 0
            seen: set[tuple[int, str, Decimal, Decimal | None]] = set()
            duplicates = 0
            for row in source.rows:
                product, matched_sku, sku_field = _match_pivot_row(
                    row, exact, folded, numeric, by_goods_id
                )
                if product is None:
                    unmatched += 1
                    continue
                identity = (int(product.id), matched_sku, row.price, row.stock)
                if identity in seen:
                    duplicates += 1
                    continue
                seen.add(identity)
                matched_rows.append((row, product, matched_sku, sku_field))
            current_rows = _published_rows_count(db, existing) if existing is not None else 0
            suspicious = bool(current_rows >= COMPLETENESS_MIN_CURRENT_ROWS and Decimal(len(matched_rows)) < Decimal(current_rows) * COMPLETENESS_WARN_RATIO)
            if suspicious and not allow_incomplete:
                results.append({"name": source.name, "stableKey": source.stable_key, "status": "blocked", "error": "completeness_override_required", "preservedPreviousSnapshot": True})
                continue
            if not matched_rows:
                results.append({"name": source.name, "stableKey": source.stable_key, "status": "error", "error": "no_internal_sku_matches", "preservedPreviousSnapshot": bool(existing)})
                continue
            try:
                with db.begin_nested():
                    price_list = existing
                    if price_list is None:
                        price_list = CompetitorPriceList(
                            price_format_id=price_format.id,
                            source_type=MULTI_SOURCE_TYPE,
                            source_key=source.stable_key,
                            coefficient=Decimal("1"),
                            price_coefficient=Decimal("1"),
                        )
                        db.add(price_list)
                        db.flush()
                    now = now_kz_naive()
                    price_list.display_name = f"{target_branch} — {source.name} — vidman-file"
                    price_list.supplier = source.name
                    price_list.competitor_name = source.name
                    price_list.region = target_branch
                    price_list.branch_id = target_branch
                    price_list.branch_code = target_branch
                    price_list.branch_name = target_branch
                    price_list.external_price_list_id = source.stable_key
                    price_list.sync_batch_id = batch_id
                    price_list.source_updated_at = now.isoformat()
                    price_list.last_checked_at = now
                    price_list.last_success_at = now
                    price_list.last_refresh_status = "updated"
                    price_list.last_refresh_message = "manual multi-source Vidman file"
                    price_list.price_date = date.today()
                    price_list.update_mode = "manual"
                    price_list.updated_at = now
                    db.execute(delete(CompetitorPriceListItem).where(CompetitorPriceListItem.price_list_id == price_list.id))
                    for row, product, matched_sku, sku_field in matched_rows:
                        db.add(
                            CompetitorPriceListItem(
                                price_list_id=price_list.id,
                                product_id=product.id,
                                provisor_goods_id=(
                                    int(row.fallback_sku)
                                    if row.fallback_sku and re.fullmatch(r"\d+", row.fallback_sku)
                                    else None
                                ),
                                name=row.name or product.name,
                                distributor_goods_name=row.name or product.name,
                                distributor_goods_id=row.distributor_sku,
                                distributor_price=row.price,
                                stock=row.stock,
                                expiry_date=row.expiry,
                                match_key=matched_sku,
                                match_type=f"manual_vidman_{sku_field.casefold()}",
                                match_score=Decimal("100"),
                                matched_sku=product.code,
                                raw_name=row.name,
                                raw_manufacturer=row.manufacturer,
                                raw_json=json.dumps(
                                    {"origin": "manual_vidman_multi_source", "batchId": batch_id, "primarySku": row.primary_sku, "fallbackSku": row.fallback_sku, "matchedBy": sku_field, "source": source.name},
                                    ensure_ascii=False,
                                    default=str,
                                ),
                            )
                        )
                    db.flush()
                    refresh_price_list_item_counters(db=db, price_list_ids=[int(price_list.id)])
                    assigned_ids = selected_price_format_ids_for_competitor_price_list(db=db, competitor_price_list_id=int(price_list.id))
                    affected_price_formats.update(assigned_ids)
                    history = ManualPriceListImport(
                        competitor_price_list_id=price_list.id,
                        source_key=source.stable_key,
                        original_filename=parsed.filename,
                        file_type=parsed.file_type,
                        file_checksum=parsed.checksum,
                        detected_sheet=parsed.sheet,
                        status="partial_success" if source.errors or unmatched else "success",
                        total_rows=parsed.total_rows,
                        empty_rows=parsed.empty_rows,
                        valid_rows=len(source.rows),
                        invalid_rows=len(source.errors),
                        duplicate_rows=duplicates,
                        matched_rows=len(matched_rows),
                        unmatched_rows=unmatched,
                        persisted_rows=len(matched_rows),
                        requested_by=requested_by,
                        started_at=now,
                        finished_at=now,
                        metadata_json=json.dumps({"origin": "manual_vidman_multi_source", "batchId": batch_id, "sourceName": source.name, "stableKey": source.stable_key}, ensure_ascii=False),
                    )
                    db.add(history)
                    db.flush()
                    results.append({"name": source.name, "stableKey": source.stable_key, "competitorPriceListId": int(price_list.id), "importId": int(history.id), "status": history.status, "matchedRows": len(matched_rows), "unmatchedRows": unmatched, "rowsWritten": len(matched_rows), "updateMode": "manual", "preservedAssignments": True})
            except Exception as exc:
                logger.exception("Failed one source in Vidman multi-source import: %s", source.name)
                results.append({"name": source.name, "stableKey": source.stable_key, "status": "error", "error": str(exc)[:1000], "preservedPreviousSnapshot": bool(existing)})

        now = now_kz_naive()
        batch_history = ManualPriceListImport(
            competitor_price_list_id=None,
            source_key=f"vidman-file-batch:{batch_id}",
            original_filename=parsed.filename,
            file_type=parsed.file_type,
            file_checksum=parsed.checksum,
            detected_sheet=parsed.sheet,
            status=(
                "success"
                if results
                and all(
                    row.get("status")
                    in {
                        "success",
                        "partial_success",
                        "skipped_empty",
                        "skipped_empty_existing_preserved",
                    }
                    for row in results
                )
                else "partial_success"
            ),
            total_rows=parsed.total_rows,
            empty_rows=parsed.empty_rows,
            valid_rows=sum(int(row.get("rowsWritten") or 0) for row in results),
            invalid_rows=sum(1 for row in results if row.get("status") in {"error", "blocked"}),
            matched_rows=sum(int(row.get("matchedRows") or 0) for row in results),
            unmatched_rows=sum(int(row.get("unmatchedRows") or 0) for row in results),
            persisted_rows=sum(int(row.get("rowsWritten") or 0) for row in results),
            requested_by=requested_by,
            started_at=now,
            finished_at=now,
            metadata_json=json.dumps(
                {
                    "origin": "manual_vidman_multi_source_batch",
                    "batchId": batch_id,
                    "priceFormatId": int(price_format.id),
                    "priceFormatCode": price_format.code,
                    "branch": target_branch,
                    "results": results,
                },
                ensure_ascii=False,
                default=str,
            ),
        )
        db.add(batch_history)
        db.commit()
        warnings: list[str] = []
        for price_format_id in affected_price_formats:
            try:
                sync_selected_competitor_configs(db=db, price_format_id=price_format_id)
                rebuild_competitor_prices_for_selected(db=db, price_format_id=price_format_id)
                db.commit()
                enqueue_percentile_preparation(db=db, price_format_id=price_format_id, reason="manual_vidman_multi_source_imported")
            except Exception:
                db.rollback()
                warnings.append(f"downstream_rebuild_failed:{price_format_id}")
                logger.exception("Downstream rebuild failed after committed multi-source Vidman import")
        return {"ok": True, "branch": target_branch, "batchId": batch_id, "batchImportId": int(batch_history.id), "filename": parsed.filename, "detectedSources": len(parsed.sources), "results": results, "warnings": warnings}
    finally:
        try:
            release_lock(db, name=lock_name, owner_token=owner_token)
        except Exception:
            db.rollback()
