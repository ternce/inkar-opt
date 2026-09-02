from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..deps import AppUser, user_can_access_branch
from ..models import CalculatedPrice, PriceFormat, PriceList, PricingWorkflowRun, Product


SAP_CATEGORIES = ("SuperVIP", "VIP", "1", "2")
SAP_HEADERS = [
    "Номер материала",
    "Категория/Прайс-лист",
    "Статус деблокирования",
    "Сумма/процентная ставка условия при отсутствии шкалы",
]
SAP_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class SapExportError(ValueError):
    def __init__(self, message: str, *, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class ResolvedSapVersion:
    price_format: PriceFormat
    price_list: PriceList
    workflow_run: PricingWorkflowRun


def normalize_sap_category(value: object) -> str | None:
    raw = str(value or "").strip()
    return raw if raw in SAP_CATEGORIES else None


def validate_sap_category_payload(value: object) -> str | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    if raw not in SAP_CATEGORIES:
        raise SapExportError("sap_category must be one of: SuperVIP, VIP, 1, 2")
    return raw


def _parse_activation_date(value: str | date) -> date:
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value or "").strip())
    except ValueError:
        raise SapExportError("activation_date must be YYYY-MM-DD")


def _branch_id_for_name(branch_name: str) -> str:
    from .references.types import BRANCHES

    branch_name = str(branch_name or "").strip()
    folded = branch_name.casefold()
    for row in BRANCHES:
        if str(row.get("name") or "").strip().casefold() == folded:
            return str(row.get("id") or "").strip()
    return branch_name


def _ensure_branch_access(branch_id: str, user: AppUser) -> None:
    if not user_can_access_branch(user, _branch_id_for_name(branch_id), branch_id):
        raise SapExportError("branch is not assigned to current user", status_code=403)


def _format_label(pf: PriceFormat) -> str:
    code = str(pf.code or "").strip()
    name = str(pf.name or "").strip()
    if code and name and code != name:
        return f"{code} {name}"
    return code or name or f"PriceFormat {pf.id}"


def _coerce_int_id(value: object, field_name: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        raise SapExportError(f"{field_name} must be an integer")


def _duplicate_values(values: Iterable[int]) -> list[int]:
    seen: set[int] = set()
    duplicates: list[int] = []
    for value in values:
        if value in seen and value not in duplicates:
            duplicates.append(value)
        seen.add(value)
    return duplicates


def _validate_unique_price_format_ids(price_format_ids: list[int]) -> None:
    duplicates = _duplicate_values(price_format_ids)
    if duplicates:
        raise SapExportError(
            "SAP file was not generated.\n"
            "Duplicate price_format_id values are not allowed:\n"
            + "\n".join(f"- {value}" for value in duplicates)
        )


def _load_price_formats(db: Session, price_format_ids: Iterable[int], branch_id: str, user: AppUser) -> list[PriceFormat]:
    ids = [_coerce_int_id(value, "price_format_id") for value in price_format_ids]
    if not ids:
        raise SapExportError("no selected formats")
    _validate_unique_price_format_ids(ids)
    _ensure_branch_access(branch_id, user)
    rows = db.execute(select(PriceFormat).where(PriceFormat.id.in_(ids))).scalars().all()
    by_id = {int(row.id): row for row in rows}
    errors: list[str] = []
    out: list[PriceFormat] = []
    for price_format_id in ids:
        pf = by_id.get(price_format_id)
        if pf is None:
            errors.append(f"PriceFormat {price_format_id} not found")
            continue
        if str(pf.branch or "").strip() != str(branch_id or "").strip():
            errors.append(f"{_format_label(pf)} belongs to another branch")
            continue
        if not user_can_access_branch(user, _branch_id_for_name(pf.branch), pf.branch):
            errors.append(f"{_format_label(pf)} is not available for current user")
            continue
        out.append(pf)
    if errors:
        raise SapExportError("SAP file was not generated.\n" + "\n".join(f"- {message}" for message in errors), status_code=403 if any("not available" in e for e in errors) else 400)
    return out


def _successful_versions_stmt(price_format_id: int, activation_date: date):
    return (
        select(PricingWorkflowRun, PriceList)
        .join(PriceList, PricingWorkflowRun.price_list_id == PriceList.id)
        .where(PricingWorkflowRun.price_format_id == price_format_id)
        .where(PricingWorkflowRun.status == "success")
        .where(PriceList.price_format_id == price_format_id)
        .where(PriceList.activation_date == activation_date)
        .order_by(PricingWorkflowRun.finished_at.desc().nulls_last(), PricingWorkflowRun.started_at.desc(), PricingWorkflowRun.id.desc())
    )


def _version_dict(run: PricingWorkflowRun, price_list: PriceList, is_latest: bool) -> dict:
    return {
        "workflow_run_id": int(run.id),
        "price_list_id": int(price_list.id),
        "price_list_number": price_list.number,
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        "is_latest": is_latest,
    }


def list_versions(db: Session, *, branch_id: str, activation_date: str | date, price_format_ids: list[int], user: AppUser) -> dict:
    effective_date = _parse_activation_date(activation_date)
    formats = _load_price_formats(db, price_format_ids, branch_id, user)
    payload = []
    for pf in formats:
        rows = db.execute(_successful_versions_stmt(int(pf.id), effective_date)).all()
        payload.append(
            {
                "price_format_id": int(pf.id),
                "code": pf.code,
                "name": pf.name,
                "sap_category": pf.sap_category,
                "versions": [_version_dict(run, price_list, index == 0) for index, (run, price_list) in enumerate(rows)],
            }
        )
    return {"branch_id": branch_id, "activation_date": effective_date.isoformat(), "formats": payload}


def resolve_latest_successful_versions(
    db: Session, *, branch_id: str, activation_date: str | date, price_format_ids: list[int], user: AppUser
) -> list[ResolvedSapVersion]:
    effective_date = _parse_activation_date(activation_date)
    formats = _load_price_formats(db, price_format_ids, branch_id, user)
    resolved: list[ResolvedSapVersion] = []
    missing: list[str] = []
    for pf in formats:
        row = db.execute(_successful_versions_stmt(int(pf.id), effective_date).limit(1)).first()
        if row is None:
            missing.append(_format_label(pf))
            continue
        run, price_list = row
        resolved.append(ResolvedSapVersion(price_format=pf, price_list=price_list, workflow_run=run))
    if missing:
        raise SapExportError(
            "SAP file was not generated.\n"
            f"No successful price list exists for {effective_date.isoformat()}:\n"
            + "\n".join(f"- {label}" for label in missing)
        )
    return resolved


def resolve_selected_versions(
    db: Session, *, branch_id: str, activation_date: str | date, items: list[dict], user: AppUser
) -> list[ResolvedSapVersion]:
    effective_date = _parse_activation_date(activation_date)
    if not items:
        raise SapExportError("no selected formats")
    parsed_items: list[dict] = []
    for item in items:
        selection_mode = str(item.get("selection_mode") or item.get("selectionMode") or "auto").strip().lower()
        if selection_mode not in {"auto", "manual"}:
            raise SapExportError("selection_mode must be auto or manual")
        parsed = {
            "price_format_id": _coerce_int_id(item.get("price_format_id") or item.get("priceFormatId"), "price_format_id"),
            "selection_mode": selection_mode,
        }
        if selection_mode == "manual":
            parsed["price_list_id"] = _coerce_int_id(item.get("price_list_id") or item.get("priceListId"), "price_list_id")
        parsed_items.append(parsed)

    price_format_ids = [item["price_format_id"] for item in parsed_items]
    formats = _load_price_formats(db, price_format_ids, branch_id, user)
    by_id = {int(pf.id): pf for pf in formats}
    resolved: list[ResolvedSapVersion] = []
    errors: list[str] = []
    for item in parsed_items:
        price_format_id = item["price_format_id"]
        pf = by_id.get(price_format_id)
        if pf is None:
            errors.append(f"PriceFormat {price_format_id} is invalid")
            continue
        if item["selection_mode"] == "auto":
            row = db.execute(_successful_versions_stmt(price_format_id, effective_date).limit(1)).first()
            if row is None:
                errors.append(f"{_format_label(pf)} has no successful price list for {effective_date.isoformat()}")
                continue
        else:
            price_list_id = item["price_list_id"]
            row = db.execute(
                select(PricingWorkflowRun, PriceList)
                .join(PriceList, PricingWorkflowRun.price_list_id == PriceList.id)
                .where(PricingWorkflowRun.price_format_id == price_format_id)
                .where(PricingWorkflowRun.price_list_id == price_list_id)
                .where(PricingWorkflowRun.status == "success")
                .where(PriceList.id == price_list_id)
                .where(PriceList.price_format_id == price_format_id)
                .where(PriceList.activation_date == effective_date)
                .order_by(PricingWorkflowRun.finished_at.desc().nulls_last(), PricingWorkflowRun.started_at.desc(), PricingWorkflowRun.id.desc())
                .limit(1)
            ).first()
            if row is None:
                errors.append(f"{_format_label(pf)} has invalid manual version {price_list_id}")
                continue
        run, price_list = row
        resolved.append(ResolvedSapVersion(price_format=pf, price_list=price_list, workflow_run=run))
    if errors:
        raise SapExportError("SAP file was not generated.\n" + "\n".join(f"- {message}" for message in errors))
    return resolved


def resolve_manual_versions(
    db: Session, *, branch_id: str, activation_date: str | date, items: list[dict], user: AppUser
) -> list[ResolvedSapVersion]:
    effective_date = _parse_activation_date(activation_date)
    if not items:
        raise SapExportError("no selected formats")
    parsed_items = [
        {
            "price_format_id": _coerce_int_id(item.get("price_format_id"), "price_format_id"),
            "price_list_id": _coerce_int_id(item.get("price_list_id"), "price_list_id"),
        }
        for item in items
    ]
    price_format_ids = [item["price_format_id"] for item in parsed_items]
    formats = _load_price_formats(db, price_format_ids, branch_id, user)
    by_id = {int(pf.id): pf for pf in formats}
    resolved: list[ResolvedSapVersion] = []
    errors: list[str] = []
    for item in parsed_items:
        price_format_id = item["price_format_id"]
        price_list_id = item["price_list_id"]
        pf = by_id.get(price_format_id)
        if pf is None:
            errors.append(f"PriceFormat {price_format_id} is invalid")
            continue
        row = db.execute(
            select(PricingWorkflowRun, PriceList)
            .join(PriceList, PricingWorkflowRun.price_list_id == PriceList.id)
            .where(PricingWorkflowRun.price_format_id == price_format_id)
            .where(PricingWorkflowRun.price_list_id == price_list_id)
            .where(PricingWorkflowRun.status == "success")
            .where(PriceList.id == price_list_id)
            .where(PriceList.price_format_id == price_format_id)
            .where(PriceList.activation_date == effective_date)
            .order_by(PricingWorkflowRun.finished_at.desc().nulls_last(), PricingWorkflowRun.started_at.desc(), PricingWorkflowRun.id.desc())
            .limit(1)
        ).first()
        if row is None:
            errors.append(f"{_format_label(pf)} has invalid manual version {price_list_id}")
            continue
        run, price_list = row
        resolved.append(ResolvedSapVersion(price_format=pf, price_list=price_list, workflow_run=run))
    if errors:
        raise SapExportError("SAP file was not generated.\n" + "\n".join(f"- {message}" for message in errors))
    return resolved


def _validate_resolved_price_lists_have_rows(db: Session, resolved_versions: list[ResolvedSapVersion]) -> None:
    price_list_ids = [int(item.price_list.id) for item in resolved_versions]
    counts = dict(
        db.execute(
            select(CalculatedPrice.price_list_id, func.count(CalculatedPrice.id))
            .where(CalculatedPrice.price_list_id.in_(price_list_ids))
            .group_by(CalculatedPrice.price_list_id)
        ).all()
    )
    empty = [item for item in resolved_versions if int(counts.get(int(item.price_list.id), 0)) == 0]
    if empty:
        raise SapExportError(
            "SAP file was not generated.\n"
            "The following generated price lists contain no calculated rows:\n"
            + "\n".join(f"- {item.price_list.number or _format_label(item.price_format)}" for item in empty)
        )


def _material_number(code: object) -> str:
    raw = str(code or "").strip()
    if not raw:
        raise SapExportError("SAP file was not generated.\n- Product.code is missing")
    return raw.lstrip("0") or "0"


def _money(value: object) -> Decimal:
    if value is None:
        raise SapExportError("SAP file was not generated.\n- CalculatedPrice.final_price is missing")
    return Decimal(str(value))


def build_sap_rows(db: Session, resolved_versions: list[ResolvedSapVersion]) -> list[dict]:
    if not resolved_versions:
        raise SapExportError("no selected formats")
    _validate_resolved_price_lists_have_rows(db, resolved_versions)
    by_price_list_id = {int(item.price_list.id): item for item in resolved_versions}
    format_order = {int(item.price_format.id): index for index, item in enumerate(resolved_versions)}
    rows = db.execute(
        select(CalculatedPrice, Product)
        .join(Product, Product.id == CalculatedPrice.product_id)
        .where(CalculatedPrice.price_list_id.in_(list(by_price_list_id)))
    ).all()
    out: list[dict] = []
    for cp, product in rows:
        resolved = by_price_list_id[int(cp.price_list_id)]
        material = _material_number(product.code)
        out.append(
            {
                "material": material,
                "category": resolved.price_format.name,
                "unlock_status": "",
                "price": _money(cp.final_price),
                "price_format_id": int(resolved.price_format.id),
                "format_order": format_order[int(resolved.price_format.id)],
            }
        )

    def sort_key(row: dict) -> tuple:
        material = str(row["material"])
        material_key = (0, int(material)) if material.isdigit() else (1, material)
        return material_key, int(row["format_order"])

    out.sort(key=sort_key)
    return out


def _excel_material(value: object) -> int | str:
    text = str(value or "").strip()
    if text.isdigit() and len(text) <= 15:
        return int(text)
    return text


def build_workbook(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(SAP_HEADERS)
    for row in rows:
        ws.append([_excel_material(row["material"]), row["category"], None, row["price"]])
    widths = [22, 32, 28, 72]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def build_export(
    db: Session,
    *,
    branch_id: str,
    activation_date: str | date,
    mode: str,
    price_format_ids: list[int] | None,
    items: list[dict] | None,
    user: AppUser,
) -> tuple[bytes, list[ResolvedSapVersion], int]:
    normalized_mode = str(mode or "").strip().lower()
    has_per_item_selection = any(
        isinstance(item, dict) and ("selection_mode" in item or "selectionMode" in item)
        for item in (items or [])
    )
    if items and has_per_item_selection:
        resolved = resolve_selected_versions(db, branch_id=branch_id, activation_date=activation_date, items=items, user=user)
    elif normalized_mode == "auto":
        resolved = resolve_latest_successful_versions(
            db,
            branch_id=branch_id,
            activation_date=activation_date,
            price_format_ids=price_format_ids or [],
            user=user,
        )
    elif normalized_mode == "manual":
        resolved = resolve_manual_versions(db, branch_id=branch_id, activation_date=activation_date, items=items or [], user=user)
    else:
        raise SapExportError("mode must be auto or manual")
    rows = build_sap_rows(db, resolved)
    return build_workbook(rows), resolved, len(rows)
