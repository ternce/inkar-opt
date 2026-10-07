from __future__ import annotations

import asyncio
import base64
import io
import json
import subprocess
import zlib
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import main
from backend.app.db import Base
from backend.app.models import (
    AppUser,
    CompetitorPriceList,
    CompetitorPriceListItem,
    ManualPriceListImport,
    CompetitorPrice,
    PriceFormat,
    PriceFormatCompetitorAssignment,
    Product,
    VidmanAccount,
    VidmanCompetitorPriceListSource,
    VidmanPriceList,
    VidmanRawItem,
)
from backend.app.services import vidman_file_import as service
from backend.app.services.vidman_competitor_price_lists import vidman_competitor_source_key
from backend.app.services.vidman_orchestrator import STATUS_MANUAL_MANAGED_SKIPPED, _refresh_single_plk


HEADERS = ["Account", "Price ID", "Name", "Manufacturer", "Expiry", "Price", "Qty/package", "Min order", "Stock"]
LEGACY_XLS_ZLIB_BASE64 = (
    "eNrtWE9IFFEY/73Znd111XV300ADWYTMzIjo0kVXI5MOK1pQEUKNs5MMq7MyrJAdyjKPQdCpiEDw0sXq0h8qqFuHwKhDEARax05BQQd1+t63M8taHlwosZjf8L75ft/798289773Zt4sJBZnHzQt4Rd0IYBVpwqhMpugVOWROCjfcaTq3SOUHB//FKoiNJAhFU9rX4flGMrxXoKC+8GXJIFPlIYwjv68ZaQ2EYfYB01IHzpJCtwmSwyN7FWSpc5yG8t7XPIZy262XGPZSWUXxWkspPvbD7qz+JTSwnkxyHYfcZ0PbNmPBrySs/jSdVEsq6LHNrXRrZnRHKzBHGjc+gzLsLXRRdTTAM7hu5MCvnkr9UXKt2+uXYDsP9baw+vYbyhBYArOWZ7gMzQh24LFRXjCzI5p1jKmUUNcJhrlHl3PT1gFCrMDtqkbqaOHqXS/NmZQfkazJs5pemHCNmxaw73nx017UnVLVgODhcl945qe00YMCuEZ00rl7axhU4njhbyeo9ZH8yOmtVej1jNG1tRNy1BlsznDjsq4z3EiviZO1PL6qSGZRR3rCV5FcdoJlu9+fZsZHkifYcsU7w3FHWSnfGg4uCxrUOUY5ygs5e7Szvoelle41R2sN7Gsp7lP99aBBlc5Ms1lrnJuK/VzgPEuvatMbyN95svg4+aZz+ndpM/3LV2on3+fnkUL9Zml+vKaRofoELduSjxJe3fhRpuPLBt/izwRJe767rjbZB1WEGU1wbLI5NsRJaa476rIAsQCJRYkFiwxlZhaYiFioRILEwuXWIRYxPVIrOORYI/kO76oRPmJE7iThJsnPYrCY9KjallVCRGTHj1HM0eeBHJxuF54Msm9dylJPOSB7C47NUThw4cPHz58+PDhY+tAuGfvAJ87wSfNUPFzg//rrFBa9X+T/Lc4hjxdBfow7YVFdxuTFc2f7VCF15bYYB3vf6HESerdRg7D7Eeu4vlL32ai/Hk2XDH+55ZQpf2vVuLnX+7/JzEY6n4="
)
PIVOT_XLS_ZLIB_BASE64 = (
    "eNrtWEtoU0EUvfMy+bZNk/QDbbHGglVrlSRtMVSa2PpDqvVTfxRBWxNUWhKJBdGFVmuXguBKcVPoxo2fjR9U0J0LoaILQRBaXbpRUHDRNp657+W10SAWtPh5N5k7c+69M3My35f3fNw/MXq7epK+kRjZaCbrJsccm0By54CP4M9mVTGXu5CylvxV4nZhIh12elDyzKnmUM33JGl0Sz6BJnqLdICOU1c6lQwuoHQwh16hOLRBC7oGi5eqmFWA9WHWZaxvcuRD1uvYcpF1G2InRA+Nx7saosYq3q/Vsc9Lqt27XOc1W8JUQU/VKj57SeixdmrPHOsd+DMdtbKYxgjztjmZSmZ6ByaoHBM4Rp+zQaJPuZ36OGjZF9YuCPYv+XZnAftlTRINUfYQL/ARLMgPUt+E3UeTycHwFHWKCuByJBtRd+cezPWRdDpxYkvCT6KbOmkPmmyWUdkuO+Q65GG5VXahHAMKSReJRcjWyAh0G9wtsoTEEvhislUGqQl5Mxwxrh+VNSQWszPKdZQzCkeLXItmf9SRg0Q9wtpQDFI7bglRy5EhjmhlY15MR6GYjvyY9YVi1hOGKBQKhbFfN4QxKttTSQ9RJBSJrgqF8bUTbevtT2YQlsKhpcIiCNt9Mu1R9xefd768866Ez4Fi6ASVctnPp4EPN9rU9Y8vtvXtiB9kyxDfcfpNuFRNHmXpnKqByl722FgXIzVwjZWsz3OrNVyuZl2OPYy8fkeFUdg0zDEX2FuPfppYXsaXzSkvR3nk/c57tSPv4itQvrF58nT5jVfxUarDWCVQX32GqVE0iqtXlNyP53JhnJpvWFd9d4K6NJ/BPWtc96U0TR4u+lnrSI2OMJFmjJWObEA2E0kgaSI7kN1EDiCHiZxAThO5gFwmcgO5TeQB8pioCKjI4C4KcBfMXc3GGc3DYzPbq2DuJSZS3L0mUtxLTaS4q9F5hFkVzL2BW3awXs3awzrC2svcFCutACuNY/wGK3295X6TxqwCJlKsylRVzQGkWOn2uYw0ZrSLWw2w3su6knUP62pmFNMCdEd1jztyVjxkiSWWWGKJJZb8TyL4GU5/bpX8lKY/EzuN9zrTSDPWa5J/VnZRGp9B/IvcSCnkGTo1r/VTSXaRa0v8ZJ3c+0Il+9B7hvqpj3n0z3v94j+NmPt7frqi79dtofn2PzMfnr+5/6+AaP8c"
)


def _session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vidman"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pivot_xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(
        [
            "SKU",
            "goodsId",
            "SKU дистрибьютора",
            "Название",
            "Срок годности",
            "Производитель дистрибьютора",
            "Цена Source A",
            "Остаток Source A",
            "Цена Source B",
            "Остаток Source B",
            "Цена Source C",
            "Остаток Source C",
        ]
    )
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _actual_client_pivot_xlsx(
    *, medicus_price: object = None, medservice_price: object = None
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Client report"
    sheet.append(["Анализ цен поставщиков за период"])
    sheet.append(["Отчет для Инкар УК"])
    sheet.append([])
    sheet.append([])
    sheet.append(
        [
            " \nКод ",
            "Наименование\u00a0",
            "  Производитель  ",
            " Эмити\n ",
            "СТОФАРМ\u00a0",
            "медикус",
            "медсервис",
            "инкар",
            "ПДЦ розн",
            "ПДЦ опт",
            "Произвольная колонка",
        ]
    )
    sheet.append(
        [
            "000000000001000017",
            "Medicine 17",
            "Maker",
            "1736,22",
            None,
            medicus_price,
            medservice_price,
            "2131,98",
            "999,01",
            "998,02",
            "997,03",
        ]
    )
    sheet.append(
        ["000000000001000097", "Medicine 97", "Maker", "451,36", "536,69", None, None, "330,75", "996,04", "995,05", "994,06"]
    )
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _seed_target(db, *, items_count: int = 0, update_mode: str = "auto"):
    pf = PriceFormat(code="FMT", name="Format")
    account = VidmanAccount(login="login-a", display_name="Account A")
    db.add_all([pf, account])
    db.flush()
    vidman_list = VidmanPriceList(account_id=account.id, main_id=1191, name="Amanat")
    price_list = CompetitorPriceList(
        price_format_id=pf.id,
        source_type="vidman",
        source_key=vidman_competitor_source_key(account.id, 1191),
        account_id=str(account.id),
        account_login=account.login,
        external_price_list_id="1191",
        display_name="Aktau - Amanat",
        items_count=items_count,
    )
    db.add_all([vidman_list, price_list])
    db.flush()
    source = VidmanCompetitorPriceListSource(
        account_id=account.id,
        main_id=1191,
        price_format_code=pf.code,
        is_active=True,
        update_mode=update_mode,
        region="Aktau",
        branch_name="Aktau",
        competitor_name="Amanat",
        competitor_price_list_id=price_list.id,
    )
    db.add(source)
    db.commit()
    return account, vidman_list, price_list, source


def test_xlsx_preview_reports_validation_duplicates_and_completeness_warning():
    db = _session()
    _account, _vidman_list, price_list, _source = _seed_target(db, items_count=200)
    content = _xlsx(
        [
            ["login-a", 1191, "Medicine A", "Maker", "01.01.2027", "12,50", 1, 1, 5],
            ["login-a", 1191, "Medicine A", "Maker", "01.01.2027", "12,50", 1, 1, 5],
            ["login-a", 1191, "Broken", "Maker", "", "zero", 1, 1, 0],
        ]
    )

    report = service.preview_vidman_file(
        db=db, competitor_price_list_id=price_list.id, content=content, filename="vidman.xlsx"
    )

    assert report["validRows"] == 1
    assert report["invalidRows"] == 1
    assert report["duplicateRows"] == 1
    assert report["requiresCompletenessOverride"] is True
    assert report["confirmationToken"] == report["checksum"]
    assert report["target"]["sourceKey"] == vidman_competitor_source_key(1, 1191)


def test_xls_is_supported(monkeypatch):
    content = zlib.decompress(base64.b64decode(LEGACY_XLS_ZLIB_BASE64))
    monkeypatch.setattr(
        service,
        "_convert_legacy_xls_to_xlsx",
        lambda _content: pytest.fail("valid legacy XLS must be parsed directly by xlrd"),
    )

    parsed = service.parse_vidman_file(content, "legacy.xls")

    assert parsed.file_type == "xls"
    assert len(parsed.valid_rows) == 1
    assert str(parsed.valid_rows[0].price) == "10.5000"


def _force_xlrd_corruption(monkeypatch):
    import xlrd

    def corrupt(*_args, **_kwargs):
        raise xlrd.biffh.XLRDError("Workbook corruption: seen[2] == 4")

    monkeypatch.setattr(xlrd, "open_workbook", corrupt)


def test_xls_corruption_falls_back_to_libreoffice_and_cleans_up(monkeypatch):
    _force_xlrd_corruption(monkeypatch)
    converted = _pivot_xlsx(
        [["SKU-1", None, "D-1", "Medicine", "", "Maker", 10, 1, 20, 2, 30, 3]]
    )
    observed: dict[str, object] = {}
    monkeypatch.setattr(service, "_libreoffice_executable", lambda: "libreoffice")

    def successful_conversion(command, **kwargs):
        observed["command"] = command
        observed["kwargs"] = kwargs
        input_path = service.Path(command[-1])
        observed["temp_dir"] = input_path.parent
        assert input_path.name == "upload.xls"
        assert input_path.read_bytes() == b"damaged legacy workbook"
        (input_path.parent / "upload.xlsx").write_bytes(converted)
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(service.subprocess, "run", successful_conversion)

    parsed = service.parse_vidman_pivot_file(
        b"damaged legacy workbook", "untrusted client name; rm -rf.xls", context="Aktau"
    )

    assert parsed.file_type == "xls"
    assert [source.name for source in parsed.sources] == ["Source A", "Source B", "Source C"]
    assert observed["command"][1:5] == ["--headless", "--convert-to", "xlsx", "--outdir"]
    assert observed["kwargs"]["shell"] is False
    assert observed["kwargs"]["timeout"] == service.LIBREOFFICE_CONVERSION_TIMEOUT_SECONDS
    assert not observed["temp_dir"].exists()


def test_xls_corruption_reports_libreoffice_unavailable(monkeypatch, caplog):
    _force_xlrd_corruption(monkeypatch)
    monkeypatch.setattr(service, "_libreoffice_executable", lambda: None)

    with pytest.raises(
        ValueError,
        match=r"^Legacy \.xls parsing failed and LibreOffice conversion is unavailable\.$",
    ):
        service.parse_vidman_file(b"damaged", "client.xls")

    assert "Workbook corruption: seen[2] == 4" in caplog.text


def test_xls_conversion_timeout_includes_both_failure_summaries(monkeypatch):
    _force_xlrd_corruption(monkeypatch)
    monkeypatch.setattr(service, "_libreoffice_executable", lambda: "libreoffice")

    def timeout(command, **_kwargs):
        raise subprocess.TimeoutExpired(command, service.LIBREOFFICE_CONVERSION_TIMEOUT_SECONDS)

    monkeypatch.setattr(service.subprocess, "run", timeout)

    with pytest.raises(ValueError) as error:
        service.parse_vidman_file(b"damaged", "client.xls")

    message = str(error.value)
    assert "XLRDError: Workbook corruption: seen[2] == 4" in message
    assert "conversion timed out after 45 seconds" in message


def test_xls_malformed_converted_xlsx_is_rejected_and_temp_files_are_cleaned(monkeypatch):
    _force_xlrd_corruption(monkeypatch)
    monkeypatch.setattr(service, "_libreoffice_executable", lambda: "libreoffice")
    observed: dict[str, object] = {}

    def malformed_conversion(command, **_kwargs):
        temp_dir = service.Path(command[-1]).parent
        observed["temp_dir"] = temp_dir
        (temp_dir / "upload.xlsx").write_bytes(b"not an xlsx")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(service.subprocess, "run", malformed_conversion)

    with pytest.raises(ValueError) as error:
        service.parse_vidman_file(b"damaged", "client.xls")

    message = str(error.value)
    assert "XLRDError: Workbook corruption: seen[2] == 4" in message
    assert "converted .xlsx could not be parsed" in message
    assert not observed["temp_dir"].exists()


def test_preview_rejects_mismatched_plk_identity():
    db = _session()
    _account, _vidman_list, price_list, _source = _seed_target(db)
    content = _xlsx([["login-a", 9999, "Medicine", "Maker", "", 10, 1, 1, 4]])

    with pytest.raises(ValueError, match="does not match selected target"):
        service.preview_vidman_file(
            db=db, competitor_price_list_id=price_list.id, content=content, filename="vidman.xlsx"
        )


def test_invalid_excel_and_missing_required_columns_are_rejected():
    with pytest.raises(ValueError, match="malformed .xlsx"):
        service.parse_vidman_file(b"not-an-excel-file", "broken.xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Manufacturer", "Stock"])
    sheet.append(["Maker", 4])
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match="required Vidman columns"):
        service.parse_vidman_file(stream.getvalue(), "missing.xlsx")


def test_confirm_requires_same_checksum_and_switches_source_to_manual(monkeypatch):
    db = _session()
    _account, _vidman_list, price_list, source = _seed_target(db)
    content = _xlsx([["login-a", 1191, "Medicine", "Maker", "", 10, 1, 1, 4]])
    checksum = service.parse_vidman_file(content, "vidman.xlsx").checksum

    with pytest.raises(ValueError, match="checksum"):
        service.import_vidman_file(
            db=db,
            competitor_price_list_id=price_list.id,
            content=content,
            filename="vidman.xlsx",
            expected_checksum="bad",
            allow_incomplete=False,
            requested_by="tester",
        )

    monkeypatch.setattr(service, "process_vidman_stage2", lambda *args, **kwargs: SimpleNamespace())
    monkeypatch.setattr(
        service,
        "build_vidman_competitor_price_list",
        lambda **kwargs: SimpleNamespace(
            skipped_reason="", trusted_match_rows=1, rows_written=1, unresolved_match_rows=0
        ),
    )
    result = service.import_vidman_file(
        db=db,
        competitor_price_list_id=price_list.id,
        content=content,
        filename="vidman.xlsx",
        expected_checksum=checksum,
        allow_incomplete=False,
        requested_by="tester",
    )

    assert result["updateMode"] == "manual"
    assert db.get(VidmanCompetitorPriceListSource, source.id).update_mode == "manual"
    assert db.execute(select(VidmanRawItem)).scalars().one().main_id == 1191
    history = db.execute(select(ManualPriceListImport)).scalars().one()
    assert history.original_filename == "vidman.xlsx"
    assert history.requested_by == "tester"


def test_auto_refresh_skips_manual_managed_plk_before_collector(monkeypatch):
    db = _session()
    account, vidman_list, price_list, _source = _seed_target(db, update_mode="manual")

    class CollectorMustNotRun:
        def __init__(self, *args, **kwargs):
            raise AssertionError("collector must not be created for a manual-managed PLK")

    monkeypatch.setattr("backend.app.services.vidman_orchestrator.VidmanRawCollector", CollectorMustNotRun)
    result = asyncio.run(
        _refresh_single_plk(
            db=db,
            credentials=SimpleNamespace(login="login-a", password="secret", config={}),
            account_id=account.id,
            account_label="Account A",
            price_list=vidman_list,
            apply=True,
        )
    )

    assert result.status == STATUS_MANUAL_MANAGED_SKIPPED
    assert result.skipped_reason == "manual_managed_source"
    assert result.competitor_price_list_id == price_list.id
    assert result.preserved_previous_snapshot is True


def test_explicit_switch_back_to_auto_restores_refresh_ownership():
    db = _session()
    _account, _vidman_list, price_list, source = _seed_target(db, update_mode="manual")

    result = service.set_vidman_update_mode(
        db=db, competitor_price_list_id=price_list.id, update_mode="auto"
    )

    assert result["updateMode"] == "auto"
    assert db.get(VidmanCompetitorPriceListSource, source.id).update_mode == "auto"


def test_empty_supplier_columns_are_reported_but_not_created_and_existing_is_preserved(monkeypatch):
    db = _session()
    pf = PriceFormat(code="UKK", name="Ust-Kamenogorsk", branch="Усть-Каменогорск")
    product_17 = Product(code="000000000001000017", name="Medicine 17", cost=1)
    product_97 = Product(code="000000000001000097", name="Medicine 97", cost=1)
    db.add_all([pf, product_17, product_97])
    db.commit()
    content = _actual_client_pivot_xlsx()

    parsed = service.parse_vidman_pivot_file(content, "actual-client.xlsx", context=pf.branch)

    assert parsed.header_row == 5
    assert parsed.format_variant == "supplier_columns"
    assert [source.name for source in parsed.sources] == ["Эмити", "Стофарм"]
    assert [len(source.rows) for source in parsed.sources] == [2, 1]
    assert parsed.sources[0].rows[0].primary_sku == "000000000001000017"
    assert str(parsed.sources[0].rows[0].price) == "1736.2200"
    assert str(parsed.sources[1].rows[0].price) == "536.6900"

    preview = service.preview_vidman_multi_source_file(
        db=db, price_format=pf, branch=pf.branch, content=content, filename="actual-client.xlsx"
    )
    assert preview["headerRow"] == 5
    assert preview["totalRows"] == 2
    assert preview["branch"] == pf.branch
    assert preview["detectedSources"] == 2
    assert preview["matchedProducts"] == 2
    assert preview["unmatchedCodes"] == 0
    assert preview["invalidPrices"] == 0
    assert preview["sources"][0]["skuMatchedRows"] == 2
    assert preview["sources"][1]["skuMatchedRows"] == 1
    assert {row["name"] for row in preview["sources"]} == {"Эмити", "Стофарм"}
    assert not {"Инкар", "ПДЦ розн", "ПДЦ опт", "Произвольная колонка"} & {row["name"] for row in preview["sources"]}

    preserved = CompetitorPriceList(
        price_format_id=pf.id,
        source_type=service.MULTI_SOURCE_TYPE,
        source_key=service._stable_file_source_key(context=pf.branch, source_name="медсервис"),
        display_name="медсервис",
        supplier="медсервис",
        competitor_name="медсервис",
        coefficient=1.37,
        price_coefficient=1.37,
        update_mode="auto",
    )
    db.add(preserved)
    db.flush()
    preserved_item = CompetitorPriceListItem(
        price_list_id=preserved.id,
        product_id=product_17.id,
        name="Preserved medicine",
        distributor_price=999,
        match_key=product_17.code,
        matched_sku=product_17.code,
    )
    assignment = PriceFormatCompetitorAssignment(
        price_format_id=pf.id,
        competitor_price_list_id=preserved.id,
        coefficient=1.19,
        is_active=True,
    )
    db.add_all([preserved_item, assignment])
    db.commit()
    preserved_id = int(preserved.id)
    preserved_item_id = int(preserved_item.id)
    assignment_id = int(assignment.id)

    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **_kwargs: {})
    first = service.import_vidman_multi_source_file(
        db=db,
        price_format=pf,
        branch=pf.branch,
        content=content,
        filename="actual-client.xlsx",
        expected_checksum=preview["confirmationToken"],
        allow_incomplete=False,
        requested_by="tester",
    )
    results = {row["name"]: row for row in first["results"]}
    assert first["ok"] is True
    assert set(results) == {"Эмити", "Стофарм"}
    non_empty_ids = {
        name: results[name]["competitorPriceListId"] for name in ["Эмити", "Стофарм"]
    }
    assert db.query(CompetitorPriceList).filter(
        CompetitorPriceList.source_type == service.MULTI_SOURCE_TYPE
    ).count() == 3
    assert service._existing_file_source(db, service._stable_file_source_key(context=pf.branch, source_name="Инкар")) is None
    assert service._existing_file_source(db, service._stable_file_source_key(context=pf.branch, source_name="ПДЦ розн")) is None
    assert service._existing_file_source(db, service._stable_file_source_key(context=pf.branch, source_name="ПДЦ опт")) is None
    preserved_after = db.get(CompetitorPriceList, preserved_id)
    assert str(preserved_after.price_coefficient) == "1.370000"
    assert preserved_after.update_mode == "auto"
    assert db.get(CompetitorPriceListItem, preserved_item_id).distributor_price == 999
    assignment_after = db.get(PriceFormatCompetitorAssignment, assignment_id)
    assert assignment_after.is_active is True
    assert str(assignment_after.coefficient) == "1.190000"
    item_counts = {
        row.display_name: db.query(CompetitorPriceListItem).filter(
            CompetitorPriceListItem.price_list_id == row.id
        ).count()
        for row in db.query(CompetitorPriceList).filter(
            CompetitorPriceList.source_type == service.MULTI_SOURCE_TYPE
        )
    }
    assert item_counts == {
        f"{pf.branch} — Эмити — vidman-file": 2,
        f"{pf.branch} — Стофарм — vidman-file": 1,
        "медсервис": 1,
    }

    skipped_history = db.query(ManualPriceListImport).filter(
        ManualPriceListImport.status.in_(
            ["skipped_empty", "skipped_empty_existing_preserved"]
        )
    ).all()
    assert skipped_history == []

    batch_history = db.get(ManualPriceListImport, first["batchImportId"])
    assert batch_history.status == "success"

    later_content = _actual_client_pivot_xlsx(medicus_price="700,25", medservice_price="800,50")
    later_checksum = service.parse_vidman_pivot_file(
        later_content, "actual-client-later.xlsx", context=pf.branch
    ).checksum
    later = service.import_vidman_multi_source_file(
        db=db,
        price_format=pf,
        branch=pf.branch,
        content=later_content,
        filename="actual-client-later.xlsx",
        expected_checksum=later_checksum,
        allow_incomplete=False,
        requested_by="tester-2",
    )
    later_results = {row["name"]: row for row in later["results"]}
    assert {
        name: later_results[name]["competitorPriceListId"]
        for name in ["Эмити", "Стофарм"]
    } == non_empty_ids
    assert db.query(CompetitorPriceList).filter(
        CompetitorPriceList.source_type == service.MULTI_SOURCE_TYPE
    ).count() == 3


def test_supplier_whitelist_zero_price_source_is_skipped_and_ignored_columns_leave_no_history(monkeypatch):
    emity = "\u042d\u043c\u0438\u0442\u0438"
    stoffarm = "\u0421\u0442\u043e\u0444\u0430\u0440\u043c"
    inkar = "\u0418\u043d\u043a\u0430\u0440"
    arbitrary = "Arbitrary supplier"
    branch = "\u0415\u0441\u0438\u043a"
    db = _session()
    pf = PriceFormat(code="ESIK", name="Esik", branch=branch)
    db.add_all([pf, Product(code="001", name="One", cost=1)])
    db.commit()

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "\u041a\u043e\u0434",
            "\u041d\u0430\u0438\u043c\u0435\u043d\u043e\u0432\u0430\u043d\u0438\u0435",
            "\u041f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c",
            f"  {emity.upper()}\n",
            f"{stoffarm}\u00a0",
            inkar,
            arbitrary,
        ]
    )
    sheet.append(["001", "One", "Maker", None, "10,50", "11,50", "12,50"])
    stream = io.BytesIO()
    workbook.save(stream)
    content = stream.getvalue()

    preview = service.preview_vidman_multi_source_file(
        db=db, price_format=pf, branch=branch, content=content, filename="zero-source.xlsx"
    )
    assert [(row["name"], row["status"]) for row in preview["sources"]] == [
        (emity, "skipped_empty"),
        (stoffarm, "ready"),
    ]
    assert preview["detectedSources"] == 2

    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **_kwargs: {})
    result = service.import_vidman_multi_source_file(
        db=db,
        price_format=pf,
        branch=branch,
        content=content,
        filename="zero-source.xlsx",
        expected_checksum=preview["confirmationToken"],
        allow_incomplete=False,
        requested_by="tester",
    )
    results = {row["name"]: row for row in result["results"]}
    assert results[emity]["status"] == "skipped_empty"
    assert results[emity]["competitorPriceListId"] is None
    assert results[stoffarm]["status"] == "success"
    assert db.query(CompetitorPriceList).filter(
        CompetitorPriceList.source_type == service.MULTI_SOURCE_TYPE
    ).count() == 1
    for ignored_name in (inkar, arbitrary):
        ignored_key = service._stable_file_source_key(
            context=branch, source_name=ignored_name
        )
        assert service._existing_file_source(db, ignored_key) is None
        assert db.query(ManualPriceListImport).filter(
            ManualPriceListImport.source_key == ignored_key
        ).count() == 0


def test_supplier_whitelist_identity_is_branch_plus_canonical_source(monkeypatch):
    emity = "\u042d\u043c\u0438\u0442\u0438"
    stoffarm = "\u0421\u0442\u043e\u0444\u0430\u0440\u043c"
    esik = "\u0415\u0441\u0438\u043a"
    shymkent = "\u0428\u044b\u043c\u043a\u0435\u043d\u0442"
    db = _session()
    pf = PriceFormat(code="ALM-FMT", name="Almaty format", branch="\u0410\u043b\u043c\u0430\u0442\u044b")
    db.add_all(
        [
            pf,
            Product(code="000000000001000017", name="Medicine 17", cost=1),
            Product(code="000000000001000097", name="Medicine 97", cost=1),
        ]
    )
    db.commit()
    content = _actual_client_pivot_xlsx()
    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **_kwargs: {})

    def run_import(branch: str, requested_by: str):
        preview = service.preview_vidman_multi_source_file(
            db=db, price_format=pf, branch=branch, content=content, filename="client.xlsx"
        )
        assert [row["name"] for row in preview["sources"]] == [emity, stoffarm]
        return service.import_vidman_multi_source_file(
            db=db,
            price_format=pf,
            branch=branch,
            content=content,
            filename="client.xlsx",
            expected_checksum=preview["confirmationToken"],
            allow_incomplete=False,
            requested_by=requested_by,
        )

    esik_first = run_import(esik, "tester-esik-1")
    esik_ids = {row["name"]: row["competitorPriceListId"] for row in esik_first["results"]}
    assert set(esik_ids) == {emity, stoffarm}
    assert {
        row["name"]: row["competitorPriceListId"]
        for row in run_import(esik, "tester-esik-2")["results"]
    } == esik_ids

    shymkent_result = run_import(shymkent, "tester-shymkent")
    shymkent_ids = {
        row["name"]: row["competitorPriceListId"]
        for row in shymkent_result["results"]
    }
    assert all(esik_ids[name] != shymkent_ids[name] for name in (emity, stoffarm))
    assert all(
        db.get(CompetitorPriceList, esik_ids[name]).branch_name == esik
        for name in (emity, stoffarm)
    )
    assert all(
        db.get(CompetitorPriceList, shymkent_ids[name]).branch_name == shymkent
        for name in (emity, stoffarm)
    )


def test_supplier_layout_numeric_excel_sku_matches_unique_zero_padded_product_code():
    db = _session()
    pf = PriceFormat(code="UKK", name="Ust-Kamenogorsk", branch="Усть-Каменогорск")
    db.add_all([pf, Product(code="000000000001000097", name="Medicine", cost=1)])
    db.commit()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Код", "Наименование", "Производитель", "эмити"])
    sheet.append([1000097, "Medicine", "Maker", 451.36])
    stream = io.BytesIO()
    workbook.save(stream)

    preview = service.preview_vidman_multi_source_file(
        db=db, price_format=pf, branch=pf.branch, content=stream.getvalue(), filename="numeric-code.xlsx"
    )

    assert preview["matchedProducts"] == 1
    assert preview["sources"][0]["skuMatchedRows"] == 1


def test_supplier_layout_skips_blank_prices_and_rejects_invalid_or_negative_prices():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Код", "Наименование", "Производитель", "эмити"])
    sheet.append(["001", "Valid", "Maker", "13,49"])
    sheet.append(["002", "Blank", "Maker", None])
    sheet.append(["003", "Negative", "Maker", "-1"])
    sheet.append(["004", "Malformed", "Maker", "not-a-price"])
    stream = io.BytesIO()
    workbook.save(stream)

    parsed = service.parse_vidman_pivot_file(stream.getvalue(), "prices.xlsx", context="Aktau")

    source = parsed.sources[0]
    assert len(source.rows) == 1
    assert str(source.rows[0].price) == "13.4900"
    assert [error.error_code for error in source.errors] == ["invalid_price", "invalid_price"]


def test_pivot_header_failure_includes_detection_diagnostics():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Report title"])
    sheet.append(["Код", "Наименование", "эмити"])
    stream = io.BytesIO()
    workbook.save(stream)

    with pytest.raises(ValueError) as error:
        service.parse_vidman_pivot_file(stream.getvalue(), "bad-layout.xlsx", context="Aktau")

    message = str(error.value)
    assert "header candidates" in message
    assert "detected fixed metadata columns" in message
    assert "supplier columns found" in message
    assert "Код/SKU" in message


def test_pivot_preview_detects_three_sources_and_matches_only_internal_skus():
    db = _session()
    pf = PriceFormat(code="AKT", name="Aktau", branch="Aktau")
    db.add_all(
        [
            pf,
            Product(code="001_A_2028", name="Primary", cost=1),
            Product(code="FALLBACK-CODE", name="Fallback", cost=1, provisor_goods_id=12345),
            Product(code="67890", name="Numeric", cost=1),
            Product(code="001234", name="Leading zero", cost=1),
            Product(code="NAME-ONLY", name="No fuzzy name fallback", cost=1),
        ]
    )
    db.commit()
    content = _pivot_xlsx(
        [
            ["001_A_2028", 999, "DIST-1", "Primary", "2028-01-01", "Maker", 10, 1, 11, 2, 12, 3],
            ["UNKNOWN", 12345.0, "DIST-2", "Fallback", "2028-01-01", "Maker", 20, 4, 21, 5, None, None],
            [67890.0, None, "DIST-3", "Numeric", "2028-01-01", "Maker", 30, 6, None, None, 32, 7],
            ["001234", None, "DIST-4", "Leading zero", "", "Maker", 40, 1, None, None, None, None],
            [1234.0, None, "DIST-5", "No fuzzy name fallback", "", "Maker", 50, 1, None, None, None, None],
        ]
    )

    report = service.preview_vidman_multi_source_file(
        db=db, price_format=pf, branch=pf.branch, content=content, filename="client.xlsx"
    )

    assert report["detectedSources"] == 3
    by_name = {row["name"]: row for row in report["sources"]}
    assert (by_name["Source A"]["validPriceRows"], by_name["Source A"]["skuMatchedRows"], by_name["Source A"]["skuUnmatchedRows"]) == (5, 4, 1)
    assert (by_name["Source B"]["validPriceRows"], by_name["Source B"]["skuMatchedRows"]) == (2, 2)
    assert (by_name["Source C"]["validPriceRows"], by_name["Source C"]["skuMatchedRows"]) == (2, 2)
    assert len({row["stableKey"] for row in report["sources"]}) == 3


def test_legacy_xls_pivot_detects_three_independent_sources():
    content = zlib.decompress(base64.b64decode(PIVOT_XLS_ZLIB_BASE64))

    parsed = service.parse_vidman_pivot_file(content, "client.xls", context="Aktau")

    assert parsed.file_type == "xls"
    assert [source.name for source in parsed.sources] == ["A", "B", "C"]
    assert [len(source.rows) for source in parsed.sources] == [2, 2, 2]


def test_multi_source_import_reuses_ids_and_preserves_assignments_and_coefficients(monkeypatch):
    db = _session()
    pf = PriceFormat(code="AKT", name="Aktau", branch="Aktau")
    db.add_all([pf, Product(code="0001", name="One", cost=1), Product(code="TWO", name="Two", cost=1, provisor_goods_id=2)])
    db.commit()
    content = _pivot_xlsx(
        [
            ["0001", None, "D-1", "One", "", "Maker", 10, 1, 11, 1, 12, 1],
            ["missing", 2.0, "D-2", "Two", "", "Maker", 20, 1, 21, 1, 22, 1],
        ]
    )
    checksum = service.parse_vidman_pivot_file(content, "client.xlsx", context="Aktau").checksum
    first = service.import_vidman_multi_source_file(
        db=db,
        price_format=pf,
        branch=pf.branch,
        content=content,
        filename="client.xlsx",
        expected_checksum=checksum,
        allow_incomplete=False,
        requested_by="tester",
    )
    ids = {row["name"]: row["competitorPriceListId"] for row in first["results"]}
    assert len(ids) == 3
    assert db.query(CompetitorPriceList).filter(CompetitorPriceList.source_type == service.MULTI_SOURCE_TYPE).count() == 3

    assignment = PriceFormatCompetitorAssignment(
        price_format_id=pf.id,
        competitor_price_list_id=ids["Source A"],
        coefficient=1.01,
        is_active=True,
    )
    db.add(assignment)
    db.get(CompetitorPriceList, ids["Source A"]).price_coefficient = 1.01
    db.commit()
    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **kwargs: {})

    second = service.import_vidman_multi_source_file(
        db=db,
        price_format=pf,
        branch=pf.branch,
        content=content,
        filename="client.xlsx",
        expected_checksum=checksum,
        allow_incomplete=False,
        requested_by="tester-2",
    )
    second_ids = {row["name"]: row["competitorPriceListId"] for row in second["results"]}

    assert second_ids == ids
    assert db.query(CompetitorPriceList).filter(CompetitorPriceList.source_type == service.MULTI_SOURCE_TYPE).count() == 3
    persisted_assignment = db.get(PriceFormatCompetitorAssignment, assignment.id)
    assert persisted_assignment.is_active is True
    assert str(persisted_assignment.coefficient) == "1.010000"
    derived = db.execute(select(CompetitorPrice).where(CompetitorPrice.price_format_id == pf.id)).scalars().all()
    assert derived
    assert {row.source_name for row in derived} == {
        f"{service.MULTI_SOURCE_TYPE}:{db.get(CompetitorPriceList, ids['Source A']).source_key}"
    }
    assert all(str(row.coefficient) == "1.010000" for row in derived)
    assert all(db.get(CompetitorPriceList, source_id).update_mode == "manual" for source_id in ids.values())
    batches = [row for row in db.query(ManualPriceListImport).all() if row.source_key.startswith("vidman-file-batch:")]
    assert len(batches) == 2
    assert all(len(json.loads(row.metadata_json)["results"]) == 3 for row in batches)


def test_multi_source_commits_each_source_before_loading_the_next(monkeypatch):
    db = _session()
    pf = PriceFormat(code="AKT-COMMITS", name="Aktau", branch="Aktau")
    db.add(Product(code="0001", name="One", cost=1))
    db.add(pf)
    db.commit()
    content = _pivot_xlsx(
        [["0001", None, "D-1", "One", "", "Maker", 10, 1, 11, 1, 12, 1]]
    )
    checksum = service.parse_vidman_pivot_file(content, "client.xlsx", context="Aktau").checksum
    commit_count = 0
    commits_seen_before_source: list[int] = []
    original_existing = service._existing_file_source

    def after_commit(_session):
        nonlocal commit_count
        commit_count += 1

    def tracked_existing(session, source_key):
        commits_seen_before_source.append(commit_count)
        return original_existing(session, source_key)

    event.listen(db, "after_commit", after_commit)
    monkeypatch.setattr(service, "_existing_file_source", tracked_existing)
    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **_kwargs: {})
    try:
        result = service.import_vidman_multi_source_file(
            db=db,
            price_format=pf,
            branch=pf.branch,
            content=content,
            filename="client.xlsx",
            expected_checksum=checksum,
            allow_incomplete=False,
            requested_by="tester",
        )
    finally:
        event.remove(db, "after_commit", after_commit)

    assert result["ok"] is True
    assert len(commits_seen_before_source) == 3
    assert commits_seen_before_source[1] > commits_seen_before_source[0]
    assert commits_seen_before_source[2] > commits_seen_before_source[1]


def test_multi_source_rolls_back_failed_source_and_preserves_completed_sources(monkeypatch):
    db = _session()
    pf = PriceFormat(code="AKT-PARTIAL", name="Aktau", branch="Aktau")
    db.add(Product(code="0001", name="One", cost=1))
    db.add(pf)
    db.commit()
    content = _pivot_xlsx(
        [["0001", None, "D-1", "One", "", "Maker", 10, 1, 11, 1, 12, 1]]
    )
    checksum = service.parse_vidman_pivot_file(content, "client.xlsx", context="Aktau").checksum
    original_counter = service.refresh_price_list_matched_item_counters
    counter_calls = 0
    rollback_count = 0

    def fail_second_counter(**kwargs):
        nonlocal counter_calls
        counter_calls += 1
        if counter_calls == 2:
            raise RuntimeError("forced source failure")
        return original_counter(**kwargs)

    def after_rollback(_session):
        nonlocal rollback_count
        rollback_count += 1

    event.listen(db, "after_rollback", after_rollback)
    monkeypatch.setattr(service, "refresh_price_list_matched_item_counters", fail_second_counter)
    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **_kwargs: {})
    try:
        result = service.import_vidman_multi_source_file(
            db=db,
            price_format=pf,
            branch=pf.branch,
            content=content,
            filename="client.xlsx",
            expected_checksum=checksum,
            allow_incomplete=False,
            requested_by="tester",
        )
    finally:
        event.remove(db, "after_rollback", after_rollback)

    by_name = {row["name"]: row for row in result["results"]}
    assert by_name["Source A"]["status"] == "success"
    assert by_name["Source B"]["status"] == "error"
    assert "forced source failure" in by_name["Source B"]["error"]
    assert by_name["Source C"]["status"] == "success"
    assert rollback_count >= 1
    assert db.get(CompetitorPriceList, by_name["Source A"]["competitorPriceListId"]) is not None


def test_multi_source_import_uses_explicit_branch_for_identity_and_display(monkeypatch):
    db = _session()
    pf = PriceFormat(code="ALM-FMT", name="Almaty format", branch="Алматы")
    db.add_all([pf, Product(code="0001", name="One", cost=1)])
    db.commit()
    content = _pivot_xlsx([["0001", None, "D-1", "One", "", "Maker", 10, 1, 11, 1, 12, 1]])
    monkeypatch.setattr(service, "enqueue_percentile_preparation", lambda **_kwargs: {})

    preview = service.preview_vidman_multi_source_file(
        db=db, price_format=pf, branch="Есик", content=content, filename="client.xlsx"
    )
    assert preview["branch"] == "Есик"
    assert preview["detectedSources"] == 3
    assert {row["stableKey"] for row in preview["sources"]} == {
        service._stable_file_source_key(context="Есик", source_name=name)
        for name in ["Source A", "Source B", "Source C"]
    }

    def run_import(branch: str, requested_by: str):
        return service.import_vidman_multi_source_file(
            db=db,
            price_format=pf,
            branch=branch,
            content=content,
            filename="client.xlsx",
            expected_checksum=preview["confirmationToken"],
            allow_incomplete=False,
            requested_by=requested_by,
        )

    first_esik = run_import("Есик", "tester-esik-1")
    first_esik_ids = {row["name"]: row["competitorPriceListId"] for row in first_esik["results"]}
    esik_sources = db.execute(
        select(CompetitorPriceList).where(CompetitorPriceList.id.in_(first_esik_ids.values()))
    ).scalars().all()
    assert first_esik["branch"] == "Есик"
    assert len(esik_sources) == 3
    assert all(row.branch_name == "Есик" and row.region == "Есик" for row in esik_sources)
    assert all(row.branch_id == "Есик" and row.branch_code == "Есик" for row in esik_sources)
    assert all(row.supplier == row.competitor_name for row in esik_sources)
    assert {row.display_name for row in esik_sources} == {
        "Есик — Source A — vidman-file",
        "Есик — Source B — vidman-file",
        "Есик — Source C — vidman-file",
    }
    assert all("Алматы" not in row.display_name for row in esik_sources)

    assignment = PriceFormatCompetitorAssignment(
        price_format_id=pf.id,
        competitor_price_list_id=first_esik_ids["Source A"],
        coefficient=1.17,
        is_active=True,
    )
    db.add(assignment)
    db.get(CompetitorPriceList, first_esik_ids["Source A"]).price_coefficient = 1.17
    db.commit()

    second_esik = run_import("Есик", "tester-esik-2")
    assert {row["name"]: row["competitorPriceListId"] for row in second_esik["results"]} == first_esik_ids
    persisted_assignment = db.get(PriceFormatCompetitorAssignment, assignment.id)
    assert persisted_assignment.is_active is True
    assert str(persisted_assignment.coefficient) == "1.170000"
    assert str(db.get(CompetitorPriceList, first_esik_ids["Source A"]).price_coefficient) == "1.170000"

    almaty = run_import("Алматы", "tester-almaty")
    shymkent = run_import("Шымкент", "tester-shymkent")
    almaty_ids = {row["name"]: row["competitorPriceListId"] for row in almaty["results"]}
    shymkent_ids = {row["name"]: row["competitorPriceListId"] for row in shymkent["results"]}
    assert first_esik_ids["Source A"] != almaty_ids["Source A"]
    assert first_esik_ids["Source A"] != shymkent_ids["Source A"]
    assert almaty_ids["Source A"] != shymkent_ids["Source A"]
    assert db.get(CompetitorPriceList, almaty_ids["Source A"]).branch_name == "Алматы"
    assert db.get(CompetitorPriceList, shymkent_ids["Source A"]).branch_name == "Шымкент"
    assert {
        row["name"]: row["competitorPriceListId"]
        for row in run_import("Алматы", "tester-almaty-2")["results"]
    } == almaty_ids
    assert {
        row["name"]: row["competitorPriceListId"]
        for row in run_import("Шымкент", "tester-shymkent-2")["results"]
    } == shymkent_ids


def test_multi_source_preview_requires_explicit_branch():
    db = _session()
    pf = PriceFormat(code="ALM-FMT", name="Almaty format", branch="Алматы")
    db.add(pf)
    db.commit()
    content = _pivot_xlsx([])

    with pytest.raises(ValueError, match="Не выбран филиал"):
        service.preview_vidman_multi_source_file(
            db=db, price_format=pf, branch="", content=content, filename="client.xlsx"
        )


def test_multi_source_preview_endpoint_returns_400_when_branch_is_missing():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    with Session() as db:
        db.add(PriceFormat(code="ALM-FMT", name="Almaty format", branch="Алматы"))
        db.commit()

    def override_db():
        with Session() as db:
            yield db

    user = AppUser(id=1, username="vidman-admin", role="admin", is_active=True)
    main.app.dependency_overrides[main.get_db] = override_db
    main.app.dependency_overrides[main.require_write_access] = lambda: user
    try:
        response = TestClient(main.app).post(
            "/api/price-formats/ALM-FMT/vidman-file/preview",
            files={"file": ("client.xlsx", _pivot_xlsx([]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    finally:
        main.app.dependency_overrides.pop(main.get_db, None)
        main.app.dependency_overrides.pop(main.require_write_access, None)

    assert response.status_code == 400
    assert response.json()["detail"] == "Не выбран филиал для импорта Vidman-файла."


def test_bad_source_block_preserves_its_snapshot_and_modes_are_independent():
    db = _session()
    pf = PriceFormat(code="AKT", name="Aktau", branch="Aktau")
    db.add_all([pf, Product(code="1", name="One", cost=1), Product(code="2", name="Two", cost=1)])
    db.commit()
    first_content = _pivot_xlsx([["1", None, "D", "One", "", "Maker", 10, 1, 20, 1, 30, 1]])
    first_checksum = service.parse_vidman_pivot_file(first_content, "first.xlsx", context="Aktau").checksum
    first = service.import_vidman_multi_source_file(
        db=db, price_format=pf, branch=pf.branch, content=first_content, filename="first.xlsx",
        expected_checksum=first_checksum, allow_incomplete=False, requested_by="tester",
    )
    ids = {row["name"]: row["competitorPriceListId"] for row in first["results"]}

    second_content = _pivot_xlsx([["missing", None, "D", "Unknown", "", "Maker", 11, 1, 99, 1, 31, 1], ["2", None, "D2", "Two", "", "Maker", 12, 1, None, None, 32, 1]])
    second_checksum = service.parse_vidman_pivot_file(second_content, "second.xlsx", context="Aktau").checksum
    second = service.import_vidman_multi_source_file(
        db=db, price_format=pf, branch=pf.branch, content=second_content, filename="second.xlsx",
        expected_checksum=second_checksum, allow_incomplete=False, requested_by="tester",
    )
    source_b_result = next(row for row in second["results"] if row["name"] == "Source B")
    assert source_b_result["status"] == "error"
    source_b_item = db.execute(
        select(CompetitorPriceListItem).where(CompetitorPriceListItem.price_list_id == ids["Source B"])
    ).scalar_one()
    assert str(source_b_item.distributor_price) == "20.0000"

    service.set_vidman_update_mode(db=db, competitor_price_list_id=ids["Source A"], update_mode="auto")
    assert db.get(CompetitorPriceList, ids["Source A"]).update_mode == "auto"
    assert db.get(CompetitorPriceList, ids["Source B"]).update_mode == "manual"
