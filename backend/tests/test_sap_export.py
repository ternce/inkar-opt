from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, func, select, update
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.db import Base
from backend.app.deps import ROLE_ADMIN, ROLE_PRICING_MANAGER, get_current_user, get_db
from backend.app.main import app
from backend.app.models import AppUser, CalculatedPrice, PriceFormat, PriceList, PricingWorkflowRun, Product, UserBranchAssignment
from backend.app.services.sap_export import (
    SAP_HEADERS,
    SapExportError,
    _money,
    build_export,
    build_sap_rows,
    build_workbook,
    list_versions,
    resolve_latest_successful_versions,
    resolve_manual_versions,
    resolve_selected_versions,
)


ACTIVATION_DATE = date(2026, 9, 2)


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    with Session() as session:
        yield session


def _admin() -> AppUser:
    return AppUser(id=1, username="admin", role=ROLE_ADMIN, is_active=True)


def _limited_user() -> AppUser:
    user = AppUser(id=2, username="manager", role=ROLE_PRICING_MANAGER, is_active=True)
    user.branches = [UserBranchAssignment(user_id=2, branch_id="1", branch_name="Алматы")]
    return user


def _format(db, code: str, category: str | None, *, branch: str = "Алматы", name: str | None = None) -> PriceFormat:
    row = PriceFormat(code=code, name=name or code, branch=branch, sap_category=category)
    db.add(row)
    db.flush()
    return row


def _product(db, code: str, *, name: str | None = None) -> Product:
    row = Product(code=code, name=name or code, cost=100)
    db.add(row)
    db.flush()
    return row


def _version(
    db,
    pf: PriceFormat,
    suffix: str,
    *,
    activation_date: date = ACTIVATION_DATE,
    status: str = "success",
    started_at: datetime = datetime(2026, 9, 1, 10, 0, 0),
    finished_at: datetime | None = datetime(2026, 9, 1, 10, 5, 0),
) -> tuple[PricingWorkflowRun, PriceList]:
    price_list = PriceList(number=f"{pf.code}_{activation_date.isoformat()}_{suffix}", price_format_id=pf.id, activation_date=activation_date)
    db.add(price_list)
    db.flush()
    run = PricingWorkflowRun(
        pricing_context_id=1,
        price_format_id=pf.id,
        price_list_id=price_list.id,
        price_list_number=price_list.number,
        status=status,
        started_at=started_at,
        finished_at=finished_at,
    )
    db.add(run)
    db.flush()
    return run, price_list


def _price(db, price_list: PriceList, product: Product, final_price: Decimal | str | int) -> CalculatedPrice:
    row = CalculatedPrice(
        price_list_id=price_list.id,
        product_id=product.id,
        cost=100,
        base_price=120,
        final_price=final_price,
        applied_reason="",
        zone="",
    )
    db.add(row)
    db.flush()
    return row


def _four_format_fixture(db, *, products_count: int = 2):
    formats = [
        _format(db, "SVIP", "SuperVIP", name="Есик_ИПЛ_SuperVIP"),
        _format(db, "VIP", "VIP", name="Есик_ИПЛ_VIP"),
        _format(db, "CAT1", "1", name="Есик_ИПЛ_1"),
        _format(db, "CAT2", "2", name="Есик_ИПЛ_2"),
    ]
    products = [_product(db, f"000000000001{idx:06d}") for idx in range(1, products_count + 1)]
    versions = []
    for index, pf in enumerate(formats):
        old_run, old_pl = _version(
            db,
            pf,
            "wf10",
            started_at=datetime(2026, 9, 1, 9, index, 0),
            finished_at=datetime(2026, 9, 1, 9, index, 30),
        )
        new_run, new_pl = _version(
            db,
            pf,
            "wf20",
            started_at=datetime(2026, 9, 1, 18, index, 0),
            finished_at=datetime(2026, 9, 1, 18, index, 30),
        )
        _version(
            db,
            pf,
            "wf30",
            status="running",
            started_at=datetime(2026, 9, 1, 19, index, 0),
            finished_at=None,
        )
        _version(
            db,
            pf,
            "wf40",
            status="error",
            started_at=datetime(2026, 9, 1, 20, index, 0),
            finished_at=datetime(2026, 9, 1, 20, index, 30),
        )
        other_run, other_pl = _version(db, pf, "wf50", activation_date=date(2026, 9, 3))
        for product in products:
            _price(db, old_pl, product, Decimal("4186.0517"))
            _price(db, new_pl, product, Decimal(100 + index))
            _price(db, other_pl, product, Decimal(900 + index))
        versions.append((pf, old_run, old_pl, new_run, new_pl, other_run, other_pl))
    db.commit()
    return formats, products, versions


def _sheet(content: bytes):
    return load_workbook(io.BytesIO(content), data_only=True).active


def test_auto_with_one_two_all_and_subset_formats(db):
    formats, _products, _versions = _four_format_fixture(db)
    for selected in ([formats[0].id], [formats[0].id, formats[1].id], [row.id for row in formats], [formats[0].id, formats[1].id, formats[3].id]):
        content, resolved, row_count = build_export(
            db,
            branch_id="Алматы",
            activation_date=ACTIVATION_DATE,
            mode="auto",
            price_format_ids=selected,
            items=None,
            user=_admin(),
        )
        assert len(resolved) == len(selected)
        assert row_count == len(selected) * 2
        assert _sheet(content).max_row == row_count + 1


def test_latest_successful_selected_and_newer_running_or_error_ignored(db):
    formats, _products, versions = _four_format_fixture(db)
    resolved = resolve_latest_successful_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, price_format_ids=[formats[0].id], user=_admin())
    assert resolved[0].workflow_run.id == versions[0][3].id
    assert resolved[0].price_list.id == versions[0][4].id


def test_multiple_successes_same_date_and_different_date_ignored(db):
    formats, _products, versions = _four_format_fixture(db)
    payload = list_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, price_format_ids=[formats[0].id], user=_admin())
    assert [version["workflow_run_id"] for version in payload["formats"][0]["versions"]] == [versions[0][3].id, versions[0][1].id]
    assert versions[0][5].id not in {version["workflow_run_id"] for version in payload["formats"][0]["versions"]}


def test_wrong_format_missing_success_and_unauthorized_branch_fail(db):
    pf = _format(db, "NO-SUCCESS", "VIP")
    _format(db, "NO-CAT", None)
    db.commit()
    with pytest.raises(SapExportError, match="No successful price list"):
        resolve_latest_successful_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, price_format_ids=[pf.id], user=_admin())
    with pytest.raises(SapExportError, match="another branch"):
        resolve_latest_successful_versions(db, branch_id="Астана", activation_date=ACTIVATION_DATE, price_format_ids=[pf.id], user=_admin())
    with pytest.raises(SapExportError, match="branch is not assigned"):
        resolve_latest_successful_versions(db, branch_id="Астана", activation_date=ACTIVATION_DATE, price_format_ids=[pf.id], user=_limited_user())


def test_export_works_when_sap_category_is_null_and_uses_price_format_code(db):
    pf = _format(db, "NULL-CAT", None, name="Есик_ИПЛ_SuperVIP")
    product = _product(db, "000000000001000015")
    _run, price_list = _version(db, pf, "wf1")
    _price(db, price_list, product, Decimal("4186.0517"))
    db.commit()

    content, resolved, row_count = build_export(
        db,
        branch_id="Алматы",
        activation_date=ACTIVATION_DATE,
        mode="auto",
        price_format_ids=[pf.id],
        items=None,
        user=_admin(),
    )

    sheet = _sheet(content)
    assert len(resolved) == 1
    assert row_count == 1
    assert sheet["B2"].value == "NULL-CAT"
    assert sheet["B2"].value != "SuperVIP"


def test_manual_latest_and_older_successful_version(db):
    formats, _products, versions = _four_format_fixture(db)
    latest = resolve_manual_versions(
        db,
        branch_id="Алматы",
        activation_date=ACTIVATION_DATE,
        items=[{"price_format_id": formats[0].id, "price_list_id": versions[0][4].id}],
        user=_admin(),
    )
    older = resolve_manual_versions(
        db,
        branch_id="Алматы",
        activation_date=ACTIVATION_DATE,
        items=[{"price_format_id": formats[0].id, "price_list_id": versions[0][2].id}],
        user=_admin(),
    )
    assert latest[0].workflow_run.id == versions[0][3].id
    assert older[0].workflow_run.id == versions[0][1].id


def test_manual_failed_wrong_date_wrong_format_and_wrong_branch_rejected(db):
    formats, _products, versions = _four_format_fixture(db)
    failed_pl = db.scalar(select(PriceList).where(PriceList.number.like("%wf40")).limit(1))
    with pytest.raises(SapExportError, match="invalid manual version"):
        resolve_manual_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, items=[{"price_format_id": formats[0].id, "price_list_id": failed_pl.id}], user=_admin())
    with pytest.raises(SapExportError, match="invalid manual version"):
        resolve_manual_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, items=[{"price_format_id": formats[0].id, "price_list_id": versions[0][6].id}], user=_admin())
    with pytest.raises(SapExportError, match="invalid manual version"):
        resolve_manual_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, items=[{"price_format_id": formats[1].id, "price_list_id": versions[0][4].id}], user=_admin())
    with pytest.raises(SapExportError, match="another branch"):
        resolve_manual_versions(db, branch_id="Астана", activation_date=ACTIVATION_DATE, items=[{"price_format_id": formats[0].id, "price_list_id": versions[0][4].id}], user=_admin())


def test_material_code_precision_blank_unlock_status_and_workbook_headers(db):
    pf = _format(db, "VIP", "VIP", name="Есик_ИПЛ_VIP")
    product = _product(db, "000000000000000000")
    _run, price_list = _version(db, pf, "wf1")
    _price(db, price_list, product, Decimal("4186.0517"))
    db.commit()
    resolved = resolve_latest_successful_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, price_format_ids=[pf.id], user=_admin())
    rows = build_sap_rows(db, resolved)
    assert rows[0]["material"] == "0"
    assert rows[0]["category"] == "VIP"
    assert rows[0]["unlock_status"] == ""
    assert rows[0]["price"] == Decimal("4186.0517")
    sheet = _sheet(build_workbook(rows))
    assert [cell.value for cell in sheet[1]] == SAP_HEADERS
    assert sheet.title == "Sheet1"
    assert sheet["A2"].value == 0
    assert sheet["A2"].data_type == "n"
    assert sheet["C2"].value is None
    assert sheet["D2"].value == 4186.0517
    assert sheet["D2"].number_format == "General"


def test_numeric_material_code_removes_leading_zeros(db):
    pf = _format(db, "VIP", None, name="Есик_ИПЛ_VIP")
    product = _product(db, "000000000001000015")
    _run, price_list = _version(db, pf, "wf1")
    _price(db, price_list, product, Decimal("100.1234"))
    db.commit()

    resolved = resolve_latest_successful_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, price_format_ids=[pf.id], user=_admin())
    sheet = _sheet(build_workbook(build_sap_rows(db, resolved)))

    assert sheet["A2"].value == 1000015
    assert sheet["A2"].data_type == "n"


def test_price_format_name_order_and_material_ascending(db):
    formats, products, versions = _four_format_fixture(db)
    extra = _product(db, "000000000000000009")
    for _pf, _old_run, _old_pl, _new_run, new_pl, *_rest in versions:
        _price(db, new_pl, extra, Decimal("200"))
    db.commit()
    content, _resolved, _row_count = build_export(
        db,
        branch_id="Алматы",
        activation_date=ACTIVATION_DATE,
        mode="auto",
        price_format_ids=[row.id for row in formats],
        items=None,
        user=_admin(),
    )
    rows = list(_sheet(content).iter_rows(min_row=2, values_only=True))
    assert rows[:4] == [
        (9, "SVIP", None, 200),
        (9, "VIP", None, 200),
        (9, "CAT1", None, 200),
        (9, "CAT2", None, 200),
    ]
    assert rows[4][0] == int(products[0].code)


def test_partial_empty_selected_price_list_rejects_entire_export(db):
    formats, _products, versions = _four_format_fixture(db)
    branch = formats[0].branch
    empty_price_list = versions[2][4]
    db.execute(delete(CalculatedPrice).where(CalculatedPrice.price_list_id == empty_price_list.id))
    db.commit()

    with pytest.raises(SapExportError) as exc_info:
        build_export(
            db,
            branch_id=branch,
            activation_date=ACTIVATION_DATE,
            mode="auto",
            price_format_ids=[row.id for row in formats],
            items=None,
            user=_admin(),
        )

    message = str(exc_info.value)
    assert "SAP file was not generated" in message
    assert "contain no calculated rows" in message
    assert empty_price_list.number in message


def test_multiple_empty_selected_price_lists_are_reported_together(db):
    formats, _products, versions = _four_format_fixture(db)
    branch = formats[0].branch
    empty_price_lists = [versions[1][4], versions[3][4]]
    db.execute(delete(CalculatedPrice).where(CalculatedPrice.price_list_id.in_([row.id for row in empty_price_lists])))
    db.commit()

    with pytest.raises(SapExportError) as exc_info:
        build_export(
            db,
            branch_id=branch,
            activation_date=ACTIVATION_DATE,
            mode="auto",
            price_format_ids=[row.id for row in formats],
            items=None,
            user=_admin(),
        )

    message = str(exc_info.value)
    assert empty_price_lists[0].number in message
    assert empty_price_lists[1].number in message


def test_auto_duplicate_format_id_rejected(db):
    formats, _products, _versions = _four_format_fixture(db)
    branch = formats[0].branch

    with pytest.raises(SapExportError, match="Duplicate price_format_id"):
        build_export(
            db,
            branch_id=branch,
            activation_date=ACTIVATION_DATE,
            mode="auto",
            price_format_ids=[formats[0].id, formats[0].id],
            items=None,
            user=_admin(),
        )


def test_manual_duplicate_format_entry_rejected(db):
    formats, _products, versions = _four_format_fixture(db)
    branch = formats[0].branch

    with pytest.raises(SapExportError, match="Duplicate price_format_id"):
        resolve_manual_versions(
            db,
            branch_id=branch,
            activation_date=ACTIVATION_DATE,
            items=[
                {"price_format_id": formats[0].id, "price_list_id": versions[0][4].id},
                {"price_format_id": formats[0].id, "price_list_id": versions[0][2].id},
            ],
            user=_admin(),
        )


def test_duplicate_sap_category_allowed_in_auto_and_manual_modes(db):
    formats, _products, versions = _four_format_fixture(db)
    branch = formats[0].branch
    formats[1].sap_category = "SuperVIP"
    db.commit()

    content, resolved, row_count = build_export(
        db,
        branch_id=branch,
        activation_date=ACTIVATION_DATE,
        mode="auto",
        price_format_ids=[formats[0].id, formats[1].id],
        items=None,
        user=_admin(),
    )
    assert len(resolved) == 2
    assert row_count == 4
    assert {row[1] for row in _sheet(content).iter_rows(min_row=2, values_only=True)} == {"SVIP", "VIP"}

    resolved = resolve_manual_versions(
        db,
        branch_id=branch,
        activation_date=ACTIVATION_DATE,
        items=[
            {"price_format_id": formats[0].id, "price_list_id": versions[0][4].id},
            {"price_format_id": formats[1].id, "price_list_id": versions[1][4].id},
        ],
        user=_admin(),
    )
    assert [item.price_format.id for item in resolved] == [formats[0].id, formats[1].id]


def test_mixed_auto_and_manual_selected_versions_work(db):
    formats, _products, versions = _four_format_fixture(db)

    content, resolved, row_count = build_export(
        db,
        branch_id=formats[0].branch,
        activation_date=ACTIVATION_DATE,
        mode="",
        price_format_ids=None,
        items=[
            {"price_format_id": formats[0].id, "selection_mode": "auto"},
            {"price_format_id": formats[1].id, "selection_mode": "manual", "price_list_id": versions[1][2].id},
            {"price_format_id": formats[2].id, "selection_mode": "auto"},
        ],
        user=_admin(),
    )

    assert [item.price_list.id for item in resolved] == [versions[0][4].id, versions[1][2].id, versions[2][4].id]
    assert row_count == 6
    rows = list(_sheet(content).iter_rows(min_row=2, values_only=True))
    assert "VIP" in {row[1] for row in rows}


def test_malformed_manual_ids_are_controlled_client_errors(db):
    with pytest.raises(SapExportError, match="price_format_id must be an integer"):
        resolve_manual_versions(
            db,
            branch_id="РђР»РјР°С‚С‹",
            activation_date=ACTIVATION_DATE,
            items=[{"price_format_id": "foo", "price_list_id": 1}],
            user=_admin(),
        )
    with pytest.raises(SapExportError, match="price_list_id must be an integer"):
        resolve_manual_versions(
            db,
            branch_id="РђР»РјР°С‚С‹",
            activation_date=ACTIVATION_DATE,
            items=[{"price_format_id": 1, "price_list_id": "bar"}],
            user=_admin(),
        )


def test_missing_product_code_and_null_final_price_rejected(db):
    pf = _format(db, "VIP", "VIP")
    product = _product(db, "000000000001000015")
    _run, price_list = _version(db, pf, "wf1")
    cp = _price(db, price_list, product, Decimal("100"))
    db.execute(update(Product).where(Product.id == product.id).values(code=""))
    db.commit()
    resolved = resolve_latest_successful_versions(db, branch_id="Алматы", activation_date=ACTIVATION_DATE, price_format_ids=[pf.id], user=_admin())
    with pytest.raises(SapExportError, match="Product.code"):
        build_sap_rows(db, resolved)

    with pytest.raises(SapExportError, match="final_price"):
        _money(None)


def test_large_export_and_no_price_mutation_or_generation(db):
    formats, _products, _versions = _four_format_fixture(db, products_count=5000)
    before_lists = db.scalar(select(func.count(PriceList.id)))
    before_prices = db.scalar(select(func.count(CalculatedPrice.id)))
    content, _resolved, row_count = build_export(
        db,
        branch_id="Алматы",
        activation_date=ACTIVATION_DATE,
        mode="auto",
        price_format_ids=[row.id for row in formats],
        items=None,
        user=_admin(),
    )
    assert row_count == 20000
    assert len(content) > 0
    assert db.scalar(select(func.count(PriceList.id))) == before_lists
    assert db.scalar(select(func.count(CalculatedPrice.id))) == before_prices


def test_versions_endpoint_contract(db):
    formats, _products, versions = _four_format_fixture(db)

    def override_db():
        yield db

    def override_user():
        return _admin()

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = override_user
    try:
        client = TestClient(app)
        response = client.get(
            f"/api/sap-export/versions?branch_id=Алматы&activation_date={ACTIVATION_DATE.isoformat()}&price_format_ids={formats[0].id}"
        )
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200
    payload = response.json()
    assert payload["branch_id"] == "Алматы"
    assert payload["activation_date"] == ACTIVATION_DATE.isoformat()
    assert payload["formats"][0]["versions"][0]["workflow_run_id"] == versions[0][3].id


def test_export_endpoint_accepts_mixed_items_payload(db):
    formats, _products, versions = _four_format_fixture(db)

    def override_db():
        yield db

    def override_user():
        return _admin()

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = override_user
    try:
        client = TestClient(app)
        response = client.post(
            "/api/sap-export",
            json={
                "branch_id": "Алматы",
                "activation_date": ACTIVATION_DATE.isoformat(),
                "items": [
                    {"price_format_id": formats[0].id, "selection_mode": "auto"},
                    {"price_format_id": formats[1].id, "selection_mode": "manual", "price_list_id": versions[1][2].id},
                ],
            },
        )
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)

    assert response.status_code == 200
    assert response.headers["X-SAP-Export-Rows"] == "4"
    assert response.headers["X-SAP-Export-Price-Lists"] == f"{versions[0][4].id},{versions[1][2].id}"
    sheet = _sheet(response.content)
    assert sheet.title == "Sheet1"
    assert {row[1] for row in sheet.iter_rows(min_row=2, values_only=True)} == {"SVIP", "VIP"}


def test_malformed_auto_id_endpoint_returns_controlled_client_error(db):
    def override_db():
        yield db

    def override_user():
        return _admin()

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = override_user
    try:
        client = TestClient(app)
        response = client.post(
            "/api/sap-export",
            json={
                "branch_id": "РђР»РјР°С‚С‹",
                "activation_date": ACTIVATION_DATE.isoformat(),
                "mode": "auto",
                "price_format_ids": ["abc"],
            },
        )
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 400
    assert "price_format_id must be an integer" in response.json()["detail"]
