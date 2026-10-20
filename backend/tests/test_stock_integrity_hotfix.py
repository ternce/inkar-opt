from __future__ import annotations

import asyncio
import csv
import io
import json
from datetime import date
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.db import Base
from backend.app.main import (
    _stock_quantities_for_price_list,
    export_price_list_csv,
    export_price_list_xlsx,
    get_generated_price_list_items,
)
from backend.app.models import (
    AppUser,
    BendRange,
    BranchCost,
    BranchStock,
    CalculatedPrice,
    MarkupRange,
    NoCompetitorMarkupRange,
    PriceFormat,
    PriceList,
    Product,
    ProductExtra,
    ReferenceUpdateStatus,
    StockSnapshot,
)
from backend.app.services.pricing import calculate_prices
from backend.app.services.references.imports import import_reference_excel


ESIK_STOCK = {
    "1005213": Decimal("32"),
    "1017769": Decimal("3"),
    "1003454": Decimal("140"),
    "1003450": Decimal("4"),
    "1005967": Decimal("31"),
    "1005754": Decimal("18"),
}


def _session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _csv_content(values: dict[str, Decimal]) -> bytes:
    lines = [
        "Материал;Остаток",
        *[f"{sku};{str(quantity).replace('.', ',')}" for sku, quantity in values.items()],
    ]
    return ("\n".join(lines) + "\n").encode("utf-8")


def _xlsx_content(values: dict[str, Decimal]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Материал", "Остаток"])
    for sku, quantity in values.items():
        sheet.append([sku, float(quantity)])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "content_factory"),
    [("Остаток_Есик.csv", _csv_content), ("Остаток_Есик.xlsx", _xlsx_content)],
)
def test_manual_stock_imports_publish_branch_snapshot_and_keep_product_extra_compatibility(filename, content_factory):
    db = _session()
    product = Product(code="1005213".zfill(18), name="Known", cost=1)
    db.add(product)
    db.commit()

    result = import_reference_excel(
        db=db,
        data_type="stock",
        branch_ids=["Есик"],
        content=content_factory({"1005213": Decimal("32.75")}),
        filename=filename,
        user_name="manual-test",
    )

    assert result.status == "success"
    status = db.scalar(
        select(ReferenceUpdateStatus)
        .where(ReferenceUpdateStatus.branch_id == "Есик")
        .where(ReferenceUpdateStatus.data_type == "stock")
    )
    assert status is not None
    assert status.active_source_type == "manual"
    assert status.active_stock_snapshot_id is not None
    assert db.scalar(select(BranchStock.stock).where(BranchStock.branch_id == "Есик")) == Decimal("32.7500")
    extra = db.scalar(select(ProductExtra).where(ProductExtra.product_id == product.id))
    assert extra is not None
    assert extra.stock == Decimal("32.7500")


def _make_price_list(db, *, number: str, branch_label: str, branch_id: str, snapshot_id: int, products: list[Product]):
    price_format = PriceFormat(code=number, name=number, branch=branch_label, reference_branch_id=branch_id)
    db.add(price_format)
    db.flush()
    price_list = PriceList(
        number=number,
        price_format_id=price_format.id,
        run_reference_versions_json=json.dumps(
            {"generationReconciliation": {"stock_snapshot_id": snapshot_id}}
        ),
    )
    db.add(price_list)
    db.flush()
    db.add_all(
        [
            CalculatedPrice(
                price_list_id=price_list.id,
                product_id=product.id,
                cost=Decimal("1"),
                base_price=Decimal("1"),
                final_price=Decimal("1"),
            )
            for product in products
        ]
    )
    db.commit()
    return price_list, price_format


async def _response_bytes(response) -> bytes:
    return b"".join([chunk async for chunk in response.body_iterator])


async def _response_text(response) -> str:
    chunks = [chunk async for chunk in response.body_iterator]
    return "".join(chunk.decode("utf-8") if isinstance(chunk, bytes) else chunk for chunk in chunks)


def test_esik_api_xlsx_csv_snapshot_parity_and_cross_branch_isolation():
    db = _session()
    products: list[Product] = []
    for sku in ESIK_STOCK:
        product = Product(code=sku.zfill(18), name=sku, cost=1)
        db.add(product)
        db.flush()
        db.add(ProductExtra(product_id=product.id, stock=Decimal("7777")))
        products.append(product)
    db.commit()

    first_import = import_reference_excel(
        db=db,
        data_type="stock",
        branch_ids=["Есик"],
        content=_csv_content(ESIK_STOCK),
        filename="Остаток_Есик.csv",
        user_name="manual-test",
    )
    assert first_import.status == "success"
    esik_status = db.scalar(
        select(ReferenceUpdateStatus)
        .where(ReferenceUpdateStatus.branch_id == "Есик")
        .where(ReferenceUpdateStatus.data_type == "stock")
    )
    first_snapshot_id = int(esik_status.active_stock_snapshot_id)
    price_list, price_format = _make_price_list(
        db,
        number="HOTFIX-ESIK-HISTORIC",
        branch_label="Есик",
        branch_id="Есик",
        snapshot_id=first_snapshot_id,
        products=products,
    )

    later_esik = {sku: Decimal("6000") + index for index, sku in enumerate(ESIK_STOCK)}
    second_import = import_reference_excel(
        db=db,
        data_type="stock",
        branch_ids=["Есик"],
        content=_xlsx_content(later_esik),
        filename="Остаток_Есик_2.xlsx",
        user_name="manual-test",
    )
    assert second_import.status == "success"
    assert db.get(StockSnapshot, first_snapshot_id).status == "superseded"

    admin = AppUser(username="audit", display_name="audit", role="admin", is_active=True)
    api_payload = get_generated_price_list_items(
        price_list.number,
        q=None,
        zone=None,
        top_filter="all",
        sort=None,
        page=1,
        page_size=100,
        db=db,
        current_user=admin,
    )
    api_values = {row["name"]: Decimal(str(row["stock"])) for row in api_payload["items"]}
    assert api_values == ESIK_STOCK
    assert all(value != Decimal("7777") for value in api_values.values())

    xlsx_response = export_price_list_xlsx(
        price_list.number, search=None, top_filter="all", top_sort="default", db=db
    )
    workbook = load_workbook(io.BytesIO(asyncio.run(_response_bytes(xlsx_response))), data_only=True)
    xlsx_values = {
        str(row[1].value): Decimal(str(row[3].value))
        for row in workbook.active.iter_rows(min_row=2)
    }
    workbook.close()
    assert xlsx_values == {sku.zfill(18): value for sku, value in ESIK_STOCK.items()}

    csv_response = export_price_list_csv(
        price_list.number, search=None, top_filter="all", top_sort="default", db=db
    )
    csv_rows = list(csv.reader(io.StringIO(asyncio.run(_response_text(csv_response)))))
    csv_values = {row[1]: Decimal(row[3]) for row in csv_rows[1:]}
    assert csv_values == {sku.zfill(18): value for sku, value in ESIK_STOCK.items()}

    other_stock = {sku: Decimal("999") for sku in ESIK_STOCK}
    other_import = import_reference_excel(
        db=db,
        data_type="stock",
        branch_ids=["1"],
        content=_csv_content(other_stock),
        filename="Остаток_Алматы.csv",
        user_name="manual-test",
    )
    assert other_import.status == "success"
    other_status = db.scalar(
        select(ReferenceUpdateStatus)
        .where(ReferenceUpdateStatus.branch_id == "1")
        .where(ReferenceUpdateStatus.data_type == "stock")
    )
    other_list, _ = _make_price_list(
        db,
        number="HOTFIX-OTHER",
        branch_label="Алматы",
        branch_id="1",
        snapshot_id=int(other_status.active_stock_snapshot_id),
        products=products,
    )
    other_payload = get_generated_price_list_items(
        other_list.number,
        q=None,
        zone=None,
        top_filter="all",
        sort=None,
        page=1,
        page_size=100,
        db=db,
        current_user=admin,
    )
    assert {row["name"]: Decimal(str(row["stock"])) for row in other_payload["items"]} == other_stock

    price_list.run_reference_versions_json = json.dumps(
        {"generationReconciliation": {"stock_snapshot_id": other_status.active_stock_snapshot_id}}
    )
    db.commit()
    with pytest.raises(HTTPException, match="does not match generated price-list branch") as exc_info:
        _stock_quantities_for_price_list(
            db,
            pl=price_list,
            pf=price_format,
            product_ids=[product.id for product in products],
        )
    assert exc_info.value.status_code == 409


def test_generation_records_active_manual_snapshot_provenance():
    db = _session()
    product = Product(code="1005213".zfill(18), name="Known", cost=Decimal("100"))
    db.add(product)
    db.commit()
    imported = import_reference_excel(
        db=db,
        data_type="stock",
        branch_ids=["Есик"],
        content=_csv_content({"1005213": Decimal("32")}),
        filename="Остаток_Есик.csv",
        user_name="manual-test",
    )
    assert imported.status == "success"
    stock_status = db.scalar(
        select(ReferenceUpdateStatus)
        .where(ReferenceUpdateStatus.branch_id == "Есик")
        .where(ReferenceUpdateStatus.data_type == "stock")
    )
    db.add(BranchCost(branch_id="Есик", product_id=product.id, sku=product.code, cost=Decimal("100")))
    db.add(
        ReferenceUpdateStatus(
            branch_id="Есик",
            data_type="cost",
            status="success",
            rows_count=1,
        )
    )
    price_format = PriceFormat(
        code="HOTFIX-GENERATION",
        name="HOTFIX-GENERATION",
        branch="Есик",
        reference_branch_id="Есик",
    )
    db.add(price_format)
    db.flush()
    db.add(MarkupRange(price_format_id=price_format.id, cost_from=0, cost_to=None, markup_percent=Decimal("0.15")))
    db.add(NoCompetitorMarkupRange(price_format_id=price_format.id, cost_from=0, cost_to=None, markup_percent=Decimal("0.03")))
    db.add(BendRange(price_format_id=price_format.id, price_from=0, bend_percent=Decimal("0.5")))
    db.commit()

    assert calculate_prices(
        db=db,
        price_format_code=price_format.code,
        price_list_number="HOTFIX-GENERATED",
        as_of=date(2026, 10, 21),
        activation_date=None,
        user="test",
        force_new_price_list=True,
    ) == 1
    price_list = db.scalar(select(PriceList).where(PriceList.number == "HOTFIX-GENERATED"))
    reconciliation = json.loads(price_list.run_reference_versions_json)["generationReconciliation"]
    assert int(reconciliation["stock_snapshot_id"]) == int(stock_status.active_stock_snapshot_id)
    assert reconciliation["stock_source_type"] == "manual"
    assert _stock_quantities_for_price_list(
        db, pl=price_list, pf=price_format, product_ids=[product.id]
    ) == {product.id: Decimal("32.0000")}
