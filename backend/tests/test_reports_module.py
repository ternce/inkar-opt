from __future__ import annotations

from datetime import datetime, timedelta
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import main
from backend.app.db import Base
from backend.app.deps import ROLE_PRICING_MANAGER, get_db
from backend.app.models import AppUser, CalculatedPrice, PriceFormat, PriceList, Product, ProductExtra, UserBranchAssignment


def _client_with_reports_data(*, user_branch: str = "Astana"):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)

    def override_db():
        with Session() as db:
            yield db

    user = AppUser(id=1, username="branch-user", role=ROLE_PRICING_MANAGER, is_active=True)
    user.branches = [UserBranchAssignment(user_id=1, branch_id="astana", branch_name=user_branch)]

    main.app.dependency_overrides[get_db] = override_db
    main.app.dependency_overrides[main.get_current_user] = lambda: user

    now = datetime(2026, 9, 8, 9, 0, 0)
    with Session() as db:
        pf_ast = PriceFormat(code="AST-PF", name="Astana PF", branch="Astana", sap_category="VIP", price_list_type="gpl")
        pf_alm = PriceFormat(code="ALM-PF", name="Almaty PF", branch="Almaty", sap_category="1")
        db.add_all([pf_ast, pf_alm])
        db.flush()
        prev_ast = PriceList(number="AST-PL-PREV", price_format_id=pf_ast.id, status="generated", created_at=now - timedelta(days=1))
        pl_ast = PriceList(number="AST-PL", price_format_id=pf_ast.id, status="generated", created_at=now)
        prev_alm = PriceList(number="ALM-PL-PREV", price_format_id=pf_alm.id, status="generated", created_at=now - timedelta(days=1))
        pl_alm = PriceList(number="ALM-PL", price_format_id=pf_alm.id, status="generated", created_at=now)
        db.add_all([prev_ast, pl_ast, prev_alm, pl_alm])
        db.flush()

        products = [
            Product(code="BOTH", name="Both Product", cost=10, top_rank=100),
            Product(code="RIGHT", name="Right Product", cost=10),
            Product(code="D045", name="Decrease 045", cost=10),
            Product(code="D010", name="Decrease 010", cost=10),
            Product(code="UNCH", name="Unchanged Product", cost=10),
            Product(code="UP", name="Increase Product", cost=10),
            Product(code="ZERO", name="Zero Old", cost=10),
            Product(code="ALM", name="Almaty Product", cost=10),
        ]
        db.add_all(products)
        db.flush()
        for product in products:
            db.add(ProductExtra(product_id=product.id, manufacturer=f"Maker {product.code}"))
        db.flush()

        previous_prices = {
            "BOTH": 1000,
            "RIGHT": 900,
            "D045": 1000,
            "D010": 1000,
            "UNCH": 1000,
            "UP": 1000,
            "ZERO": 0,
        }
        current_prices = {
            "BOTH": 990,
            "RIGHT": 910,
            "D045": 995.5,
            "D010": 999,
            "UNCH": 1000,
            "UP": 1001,
            "ZERO": 0,
        }
        by_code = {product.code: product for product in products}
        for code, old_price in previous_prices.items():
            product = by_code[code]
            db.add(
                CalculatedPrice(
                    price_list_id=prev_ast.id,
                    product_id=product.id,
                    cost=10,
                    base_price=111,
                    final_price=old_price,
                    competitor_price=995,
                    zone="right",
                    rating_global=10 if code == "BOTH" else None,
                )
            )
            db.add(
                CalculatedPrice(
                    price_list_id=pl_ast.id,
                    product_id=product.id,
                    cost=10,
                    base_price=222,
                    final_price=current_prices[code],
                    competitor_price=995,
                    lowest_competitor_price=995,
                    zone="left" if code in {"BOTH", "ZERO"} else "right",
                    rating_global=10 if code == "BOTH" else None,
                )
            )
        db.add(CalculatedPrice(price_list_id=prev_alm.id, product_id=by_code["ALM"].id, cost=10, base_price=100, final_price=1000, competitor_price=995, zone="right"))
        db.add(CalculatedPrice(price_list_id=pl_alm.id, product_id=by_code["ALM"].id, cost=10, base_price=100, final_price=990, competitor_price=995, zone="left"))
        db.commit()

    return TestClient(main.app)


def _clear_overrides():
    main.app.dependency_overrides.pop(get_db, None)
    main.app.dependency_overrides.pop(main.get_current_user, None)


def test_rank_1_report_uses_left_zone_and_reference_columns():
    client = _client_with_reports_data()
    try:
        response = client.get("/api/reports/rank-1?price_list_id=AST-PL&limit=20")
        assert response.status_code == 200
        payload = response.json()
        by_material = {row["material"]: row for row in payload["items"]}

        assert set(by_material) == {"BOTH", "ZERO"}
        assert "oldPrice" not in by_material["BOTH"]
        assert by_material["BOTH"]["newPrice"] == 990.0
        assert by_material["BOTH"]["rank"] == 1
        assert by_material["BOTH"]["top1500"] == 10
        assert by_material["ZERO"]["top1500"] == 0
        assert by_material["BOTH"]["customerCategory"] == "VIP"
        assert payload["summary"]["totalRank1"] == 2
    finally:
        _clear_overrides()


def test_decreases_use_previous_final_price_and_include_mild_decreases():
    client = _client_with_reports_data()
    try:
        response = client.get("/api/reports/decreases?price_list_id=AST-PL&limit=20")
        assert response.status_code == 200
        payload = response.json()
        by_material = {row["material"]: row for row in payload["items"]}

        assert set(by_material) == {"BOTH", "D045", "D010"}
        assert by_material["BOTH"]["oldPrice"] == 1000.0
        assert by_material["BOTH"]["oldPrice"] != 222.0
        assert by_material["BOTH"]["newPrice"] == 990.0
        assert by_material["BOTH"]["decreaseKzt"] == -10.0
        assert by_material["BOTH"]["decreasePercent"] == -0.01
        assert by_material["D045"]["decreasePercent"] == -0.0045
        assert by_material["D010"]["decreasePercent"] == -0.001
        assert "UNCH" not in by_material
        assert "UP" not in by_material
        assert "ZERO" not in by_material
        assert by_material["BOTH"]["region"] == "Astana"
        assert by_material["BOTH"]["customerCategory"] == "VIP"
        assert payload["context"]["previousPriceListNumber"] == "AST-PL-PREV"

        searched = client.get("/api/reports/decreases?price_list_id=AST-PL&q=Maker D045&limit=20").json()
        assert searched["total"] == 1
        assert searched["items"][0]["material"] == "D045"
    finally:
        _clear_overrides()


def test_reports_respect_access_and_reference_excel_formatting():
    client = _client_with_reports_data()
    try:
        denied = client.get("/api/reports/rank-1?price_list_id=ALM-PL")
        assert denied.status_code == 403

        rank_response = client.get("/api/reports/rank-1/export.xlsx?price_list_id=AST-PL")
        assert rank_response.status_code == 200
        rank_sheet = load_workbook(io.BytesIO(rank_response.content), data_only=True).active
        assert rank_sheet.title == "КАТ ВИП"
        assert list(next(rank_sheet.iter_rows(values_only=True))) == [
            "Категория клиента",
            "Материал",
            "Краткий Tекст Материала",
            "Производитель",
            "ТОП-1500",
            "Новая цена",
            "Ранг 1",
        ]
        assert rank_sheet.freeze_panes == "A2"
        assert rank_sheet.auto_filter.ref == f"A1:G{rank_sheet.max_row}"
        assert rank_sheet["F2"].number_format.find("₸") >= 0
        assert rank_sheet["E2"].value == 10

        decrease_response = client.get("/api/reports/decreases/export.xlsx?price_list_id=AST-PL")
        assert decrease_response.status_code == 200
        decrease_sheet = load_workbook(io.BytesIO(decrease_response.content), data_only=True).active
        assert decrease_sheet.title == "ВИП"
        assert list(next(decrease_sheet.iter_rows(values_only=True))) == [
            "Регион",
            "Категория клиента",
            "Материал",
            "Краткий Tекст Материала",
            "ТОП - 1500",
            "нов цена с НДС",
            "Старая цена",
            "Снижение в тг",
            "Снижение в %",
            "Производитель",
        ]
        assert decrease_sheet.freeze_panes == "A2"
        assert decrease_sheet.auto_filter.ref == f"A1:J{decrease_sheet.max_row}"
        assert decrease_sheet["F2"].number_format.find("₸") >= 0
        assert decrease_sheet["G2"].number_format.find("₸") >= 0
        assert decrease_sheet["H2"].number_format.find("₸") >= 0
        assert decrease_sheet["I2"].number_format == "0.0%"
        rows = list(decrease_sheet.iter_rows(values_only=True))
        exported = {row[2]: row for row in rows[1:]}
        assert exported["BOTH"][0] == "Astana"
        assert exported["BOTH"][1] == "VIP"
        assert exported["BOTH"][6] == 1000
        assert exported["BOTH"][7] == -10
        assert exported["BOTH"][8] == -0.01
    finally:
        _clear_overrides()
