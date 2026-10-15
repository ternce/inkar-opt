from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from decimal import Decimal
import io
import os
from pathlib import Path
import threading
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, event, func, inspect, insert, select, text
from sqlalchemy.exc import IntegrityError, OperationalError, TimeoutError as SQLAlchemyTimeoutError
from sqlalchemy.orm import Session
from sqlalchemy.pool import NullPool

from backend.app.db import Base
from backend.app.models import (
    CompetitorMatchDependencyRevision,
    CompetitorMatchJob,
    CompetitorMatchSnapshot,
    CompetitorMatchSnapshotItem,
    CompetitorCodeMapping,
    CompetitorPrice,
    CompetitorPriceList,
    CompetitorPriceListItem,
    CompetitorSourcePublication,
    CompetitorSourcePublicationItem,
    CompetitorSourceVersion,
    CompetitorSourceVersionItem,
    PriceFormat,
    PriceFormatCompetitorAssignment,
    PriceFormatCompetitorMaterialization,
    PriceList,
    PricingContext,
    PricingWorkflowRun,
    Product,
    CalculatedPrice,
    BendRange,
    MarkupRange,
    NoCompetitorMarkupRange,
    SourceGoodsMatch,
)
from backend.app.services.matching_cache.capture import capture_and_publish_legacy_source_version
from backend.app.services.matching_cache.config import matching_cache_enabled
from backend.app.services.matching_cache.dependencies import (
    MANUAL_CODE_MAPPINGS,
    PRODUCT_CATALOG,
    build_dependency_vector,
    bump_dependency_revision,
    ensure_dependency_revision,
)
from backend.app.services.matching_cache.fingerprints import (
    LEGACY_FIELD_ALIASES,
    SOURCE_CONTENT_FIELDS,
    canonical_source_content,
)
from backend.app.services.matching_cache.dependency_tracking import DEPENDENCY_MODEL_COVERAGE
from backend.app.services.matching_cache.versions import (
    create_or_reuse_pending_source_version,
    get_source_version_metadata,
    mark_source_version_matching,
    mark_source_version_ready,
    publish_source_version,
    record_source_version_failure,
)
from backend.app.services.matching_cache.jobs import (
    claim_next_matching_job,
    enqueue_matching_job,
    fail_matching_job,
    recover_stale_matching_jobs,
)
from backend.app.services.matching_cache.shadow import (
    ShadowSnapshotConflict,
    build_shadow_snapshot,
    get_or_create_shadow_snapshot,
)
from backend.app.services.matching_cache import shadow as shadow_service
from backend.app.services.matching_cache import capture as capture_service
from backend.app.services.matching_cache.materialization import (
    MatchingCacheConflict,
    MatchingCacheUnavailable,
)
from backend.app.services.matching_cache.differential import compare_legacy_and_shadow_matches, compare_records
from backend.app.timezone import now_kz_naive
from backend.app.services.pricing import resolve_competitor_price
from backend.app.services.competitor_matching import rematch_price_list_items_by_product
from backend.app.services import competitor_matching as competitor_matching_service
from backend.app.services.competitor_matching import rebuild_competitor_prices_for_selected
from backend.app.services.competitor_price_lists import upsert_provisor_price_list, sync_selected_competitor_configs
from backend.app.services.manual_price_list_import import import_manual_price_list
from backend.app.services.pricing import calculate_price_for_product
from backend.app.services.pricing_workflow.analytics import build_workflow_analytics
from backend.app.services.pricing_workflow.exports import export_workflow_run
from backend.app.services.pricing_workflow.locks import acquire_transaction_write_locks


REPO_ROOT = Path(__file__).resolve().parents[2]
UP_MIGRATION = REPO_ROOT / "backend/scripts/migrations/20261014_phase2b1_matching_cache_foundation.up.sql"
DOWN_MIGRATION = REPO_ROOT / "backend/scripts/migrations/20261014_phase2b1_matching_cache_foundation.down.sql"
PHASE2B2_UP_MIGRATION = REPO_ROOT / "backend/scripts/migrations/20261014_phase2b2_shadow_jobs.up.sql"
PHASE2B2_DOWN_MIGRATION = REPO_ROOT / "backend/scripts/migrations/20261014_phase2b2_shadow_jobs.down.sql"
LINEAGE_UP_MIGRATION = REPO_ROOT / "backend/scripts/migrations/20261015_phase2b4_stable_source_lineage.up.sql"
LINEAGE_DOWN_MIGRATION = REPO_ROOT / "backend/scripts/migrations/20261015_phase2b4_stable_source_lineage.down.sql"
NEW_TABLES = {
    "competitor_source_versions",
    "competitor_source_version_items",
    "competitor_match_dependency_revisions",
    "competitor_match_snapshots",
    "competitor_match_snapshot_items",
    "competitor_match_jobs",
    "price_format_competitor_materializations",
    "competitor_source_publications",
    "competitor_source_publication_items",
}


def _postgres_url() -> str:
    url = os.getenv("TEST_POSTGRES_URL", "").strip()
    if not url:
        pytest.skip("TEST_POSTGRES_URL is required; Phase 2B.1 must be tested on PostgreSQL 16")
    if url.startswith("postgres://"):
        url = "postgresql+psycopg://" + url[len("postgres://") :]
    elif url.startswith("postgresql://") and "+psycopg" not in url:
        url = "postgresql+psycopg://" + url[len("postgresql://") :]
    return url


def _execute_script(connection, path: Path) -> None:
    connection.exec_driver_sql(path.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def pg_engine():
    url = _postgres_url()
    admin = create_engine(url, poolclass=NullPool)
    schema = f"phase2b1_{uuid.uuid4().hex}"
    try:
        with admin.begin() as connection:
            version = int(connection.execute(text("SHOW server_version_num")).scalar_one())
            if not (160000 <= version < 170000):
                pytest.skip(f"PostgreSQL 16 required, found server_version_num={version}")
            connection.exec_driver_sql(f'CREATE SCHEMA "{schema}"')
    except Exception as exc:
        admin.dispose()
        pytest.skip(f"TEST_POSTGRES_URL is unavailable: {exc}")

    engine = create_engine(
        url,
        pool_size=5,
        max_overflow=0,
        connect_args={"options": f"-csearch_path={schema},public"},
    )
    Base.metadata.create_all(engine)
    yield engine
    engine.dispose()
    with admin.begin() as connection:
        connection.exec_driver_sql(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
    admin.dispose()


@pytest.fixture()
def source_id(pg_engine) -> int:
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        price_format = PriceFormat(code=f"P2B1-{token}", name="Phase 2B.1", branch="test")
        db.add(price_format)
        db.flush()
        source = CompetitorPriceList(
            price_format_id=price_format.id,
            source_type="manual",
            source_key=f"manual:{token}",
            display_name="Phase 2B.1 source",
        )
        db.add(source)
        db.commit()
        return int(source.id)


def _rows(price: str = "10.00") -> list[dict]:
    return [
        {
            "source_ordinal": 0,
            "source_row_key": "sku:A",
            "distributor_goods_id": "A",
            "name": "Medicine A",
            "raw_manufacturer": "Factory",
            "distributor_price": price,
            "raw_json": {"goods": {"id": "A", "price": price}},
        },
        {
            "source_ordinal": 1,
            "source_row_key": "sku:A",
            "distributor_goods_id": "A",
            "name": "Medicine A",
            "raw_manufacturer": "Factory",
            "distributor_price": price,
            "raw_json": {"goods": {"price": price, "id": "A"}},
        },
    ]


def test_models_are_created_on_postgresql_16(pg_engine):
    assert NEW_TABLES <= set(inspect(pg_engine).get_table_names())


def test_migration_upgrade_and_downgrade_on_postgresql_16():
    url = _postgres_url()
    admin = create_engine(url, poolclass=NullPool)
    schema = f"phase2b1_migration_{uuid.uuid4().hex}"
    try:
        with admin.begin() as connection:
            connection.exec_driver_sql(f'CREATE SCHEMA "{schema}"')
        engine = create_engine(
            url,
            poolclass=NullPool,
            connect_args={"options": f"-csearch_path={schema},public"},
        )
        with engine.begin() as connection:
            connection.exec_driver_sql("CREATE TABLE price_formats (id SERIAL PRIMARY KEY)")
            connection.exec_driver_sql("CREATE TABLE products (id SERIAL PRIMARY KEY)")
            connection.exec_driver_sql("CREATE TABLE competitor_price_lists (id SERIAL PRIMARY KEY)")
            connection.exec_driver_sql(
                "CREATE TABLE competitor_price_list_items (id SERIAL PRIMARY KEY)"
            )
            _execute_script(connection, UP_MIGRATION)
            _execute_script(connection, PHASE2B2_UP_MIGRATION)
            _execute_script(connection, LINEAGE_UP_MIGRATION)
        migration_inspector = inspect(engine)
        assert NEW_TABLES <= set(migration_inspector.get_table_names())
        for table_name in NEW_TABLES:
            migration_columns = {column["name"] for column in migration_inspector.get_columns(table_name)}
            orm_columns = set(Base.metadata.tables[table_name].columns.keys())
            assert migration_columns == orm_columns
        with engine.begin() as connection:
            _execute_script(connection, LINEAGE_DOWN_MIGRATION)
            _execute_script(connection, PHASE2B2_DOWN_MIGRATION)
            _execute_script(connection, DOWN_MIGRATION)
        assert not (NEW_TABLES & set(inspect(engine).get_table_names()))
        assert {"price_formats", "products", "competitor_price_lists"} <= set(inspect(engine).get_table_names())
        engine.dispose()
    finally:
        with admin.begin() as connection:
            connection.exec_driver_sql(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
        admin.dispose()


def test_source_version_reuses_identical_content_and_changes_generation(pg_engine, source_id):
    with Session(pg_engine) as db:
        first = create_or_reuse_pending_source_version(
            db=db, competitor_price_list_id=source_id, rows=_rows(), source_metadata={"providerTimestamp": "ignored"}
        )
        db.commit()
        assert first.reused is False
        assert int(first.version.generation) == 1

    with Session(pg_engine) as db:
        reused = create_or_reuse_pending_source_version(
            db=db, competitor_price_list_id=source_id, rows=_rows(), source_metadata={"providerTimestamp": "different"}
        )
        db.commit()
        assert reused.reused is True
        assert int(reused.version.id) == int(first.version.id)

    with Session(pg_engine) as db:
        changed = create_or_reuse_pending_source_version(
            db=db, competitor_price_list_id=source_id, rows=_rows("9.99")
        )
        db.commit()
        assert changed.reused is False
        assert int(changed.version.generation) == 2
        assert changed.version.content_sha256 != first.version.content_sha256


def test_duplicate_rows_are_retained_and_affect_fingerprint(pg_engine, source_id):
    duplicated = canonical_source_content(_rows())
    reordered = canonical_source_content(list(reversed(_rows())))
    single = canonical_source_content(_rows()[:1])
    assert duplicated.row_count == 2
    assert [row.duplicate_ordinal for row in duplicated.rows] == [0, 1]
    assert duplicated.content_sha256 == reordered.content_sha256
    assert duplicated.content_sha256 != single.content_sha256

    with Session(pg_engine) as db:
        created = create_or_reuse_pending_source_version(
            db=db, competitor_price_list_id=source_id, rows=_rows()
        )
        db.commit()
        persisted = db.execute(
            select(CompetitorSourceVersionItem)
            .where(CompetitorSourceVersionItem.source_version_id == created.version.id)
            .order_by(CompetitorSourceVersionItem.source_ordinal)
        ).scalars().all()
        assert len(persisted) == 2
        assert [row.duplicate_ordinal for row in persisted] == [0, 1]


def test_surrogate_ids_do_not_change_content_but_legacy_match_fields_do():
    first = _rows()
    second = _rows()
    for ordinal, row in enumerate(first, start=100):
        row["legacy_source_item_id"] = ordinal
        row["product_id"] = 10
        row["match_key"] = "manual:ABC"
        row["match_type"] = "manual_code"
        row["match_score"] = "1.0"
        row["matched_sku"] = "ABC"
    for ordinal, row in enumerate(second, start=200):
        row["legacy_source_item_id"] = ordinal
        row["product_id"] = 10
        row["match_key"] = "manual:ABC"
        row["match_type"] = "manual_code"
        row["match_score"] = "1"
        row["matched_sku"] = "ABC"

    assert canonical_source_content(first).content_sha256 == canonical_source_content(second).content_sha256
    second[0]["product_id"] = 11
    assert canonical_source_content(first).content_sha256 != canonical_source_content(second).content_sha256


def test_dependency_revision_changes_vector(pg_engine):
    scope = f"catalog-{uuid.uuid4().hex}"
    with Session(pg_engine) as db:
        initial = ensure_dependency_revision(db=db, scope_type=PRODUCT_CATALOG, scope_key=scope)
        db.commit()
        assert int(initial.revision) == 1
        before = build_dependency_vector(db=db, scopes={PRODUCT_CATALOG: scope})
        changed = bump_dependency_revision(
            db=db,
            scope_type=PRODUCT_CATALOG,
            scope_key=scope,
            fingerprint="a" * 64,
            metadata={"reason": "test"},
        )
        db.commit()
        after = build_dependency_vector(db=db, scopes={PRODUCT_CATALOG: scope})
        assert int(changed.revision) == 2
        assert before.fingerprint != after.fingerprint
        assert after.revisions[f"{PRODUCT_CATALOG}:{scope}"] == 2


def test_concurrent_identical_version_creation_returns_one_generation(pg_engine, source_id):
    barrier = threading.Barrier(2)

    def create() -> tuple[int, bool]:
        with Session(pg_engine) as db:
            barrier.wait(timeout=10)
            result = create_or_reuse_pending_source_version(
                db=db,
                competitor_price_list_id=source_id,
                rows=_rows("8.88"),
            )
            version_id = int(result.version.id)
            reused = result.reused
            db.commit()
            return version_id, reused

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: create(), range(2)))
    assert len({item[0] for item in results}) == 1
    assert sorted(item[1] for item in results) == [False, True]
    with Session(pg_engine) as db:
        count = db.execute(
            select(func.count(CompetitorSourceVersion.id))
            .where(CompetitorSourceVersion.competitor_price_list_id == source_id)
            .where(CompetitorSourceVersion.content_sha256 == canonical_source_content(_rows("8.88")).content_sha256)
        ).scalar_one()
        assert int(count) == 1


def test_failed_publication_transaction_rolls_back(pg_engine, source_id):
    with Session(pg_engine) as db:
        original = create_or_reuse_pending_source_version(
            db=db, competitor_price_list_id=source_id, rows=_rows("7.76")
        )
        original_id = int(original.version.id)
        mark_source_version_matching(db=db, version_id=original_id)
        mark_source_version_ready(db=db, version_id=original_id)
        publish_source_version(db=db, version_id=original_id)
        replacement = create_or_reuse_pending_source_version(
            db=db, competitor_price_list_id=source_id, rows=_rows("7.77")
        )
        replacement_id = int(replacement.version.id)
        mark_source_version_matching(db=db, version_id=replacement_id)
        mark_source_version_ready(db=db, version_id=replacement_id)
        db.commit()

    with pytest.raises(RuntimeError, match="publication failed"):
        with Session(pg_engine) as db:
            try:
                publish_source_version(db=db, version_id=replacement_id)
                raise RuntimeError("publication failed")
            except Exception:
                db.rollback()
                raise

    with Session(pg_engine) as db:
        old_metadata = get_source_version_metadata(db=db, version_id=original_id)
        replacement_metadata = get_source_version_metadata(db=db, version_id=replacement_id)
        assert old_metadata is not None
        assert replacement_metadata is not None
        assert old_metadata.status == "published"
        assert replacement_metadata.status == "ready"
        failed = record_source_version_failure(
            db=db, version_id=replacement_id, error="publication failed"
        )
        db.commit()
        assert failed.status == "failed"


def test_legacy_pricing_is_unchanged_when_cache_is_disabled(pg_engine, monkeypatch):
    monkeypatch.delenv("MATCH_CACHE_ENABLED", raising=False)
    assert matching_cache_enabled() is False
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        price_format = PriceFormat(code=f"LEGACY-{token}", name="Legacy", branch="test")
        product = Product(code=f"SKU-{token}", name="Legacy product", cost=1)
        db.add_all([price_format, product])
        db.flush()
        source_name = f"manual:legacy:{token}"
        db.add_all(
            [
                CompetitorPrice(
                    price_format_id=price_format.id,
                    product_id=None,
                    source_name=source_name,
                    coefficient=Decimal("1.25"),
                ),
                CompetitorPrice(
                    price_format_id=price_format.id,
                    product_id=product.id,
                    source_name=source_name,
                    source_price=Decimal("80"),
                ),
            ]
        )
        db.commit()
        resolved = resolve_competitor_price(db, int(price_format.id), int(product.id))
        assert resolved.competitor_price == Decimal("100")
        assert resolved.applied_source == source_name


def _enable_dual_write(monkeypatch, *, shadow: bool = False) -> None:
    monkeypatch.setenv("MATCH_CACHE_ENABLED", "true")
    monkeypatch.setenv("MATCH_CACHE_DUAL_WRITE", "true")
    monkeypatch.setenv("MATCH_CACHE_SHADOW_ENABLED", "true" if shadow else "false")
    monkeypatch.setenv("MATCH_CACHE_READ", "false")
    monkeypatch.setenv("MATCH_CACHE_WORKER_ENABLED", "false")


def _add_legacy_item(db: Session, source_id: int, *, code: str, price: str = "10") -> int:
    row = CompetitorPriceListItem(
        price_list_id=source_id,
        product_id=None,
        name=f"Medicine {code}",
        distributor_goods_name=f"Medicine {code}",
        distributor_goods_id=code,
        distributor_price=Decimal(price),
        match_type="unmatched",
        matched_sku="",
        raw_name=f"Medicine {code}",
        raw_json="{}",
    )
    db.add(row)
    db.flush()
    return int(row.id)


def test_streaming_import_capture_reuses_and_republishes_changed_content(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        _add_legacy_item(db, source_id, code="CACHE-A")
        db.commit()
    with Session(pg_engine) as db:
        first = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id, source_metadata={"providerTimestamp": "one"}
        )
        db.commit()
        assert first.succeeded and not first.reused and first.row_count == 1
    with Session(pg_engine) as db:
        reused = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id, source_metadata={"providerTimestamp": "two"}
        )
        db.commit()
        assert reused.succeeded and reused.reused
        assert reused.version_id == first.version_id
        item = db.scalar(
            select(CompetitorPriceListItem).where(CompetitorPriceListItem.price_list_id == source_id)
        )
        old_item_id = int(item.id)
        db.delete(item)
        db.flush()
        new_item_id = _add_legacy_item(db, source_id, code="CACHE-A")
        assert new_item_id != old_item_id
        same_content_new_surrogate = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        assert same_content_new_surrogate.reused
        assert same_content_new_surrogate.version_id == first.version_id
        item = db.get(CompetitorPriceListItem, new_item_id)
        item.distributor_price = Decimal("11")
        changed = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        db.commit()
        assert changed.succeeded and not changed.reused
    with Session(pg_engine) as db:
        statuses = db.execute(
            select(CompetitorSourceVersion.status)
            .where(CompetitorSourceVersion.competitor_price_list_id == source_id)
            .order_by(CompetitorSourceVersion.generation)
        ).scalars().all()
        assert statuses == ["superseded", "published"]
        item = db.scalar(
            select(CompetitorPriceListItem).where(CompetitorPriceListItem.price_list_id == source_id)
        )
        item.distributor_price = Decimal("10")
        reverted = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        db.commit()
        assert reverted.reused and reverted.version_id == first.version_id
        statuses = db.execute(
            select(CompetitorSourceVersion.status)
            .where(CompetitorSourceVersion.competitor_price_list_id == source_id)
            .order_by(CompetitorSourceVersion.generation)
        ).scalars().all()
        assert statuses == ["published", "superseded"]


def test_capture_failure_savepoint_preserves_legacy_import(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        item_id = _add_legacy_item(db, source_id, code=f"SAFE-{uuid.uuid4().hex}")

        def fail_capture(**_kwargs):
            raise RuntimeError("simulated cache persistence failure")

        monkeypatch.setattr(capture_service, "create_or_reuse_legacy_source_version", fail_capture)
        outcome = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        assert outcome.enabled and not outcome.succeeded
        db.commit()
    with Session(pg_engine) as db:
        assert db.get(CompetitorPriceListItem, item_id) is not None
        assert db.scalar(
            select(func.count(CompetitorSourceVersion.id)).where(
                CompetitorSourceVersion.competitor_price_list_id == source_id
            )
        ) == 0


def test_provisor_and_manual_import_publication_paths_dual_write(
    pg_engine, monkeypatch
):
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex
    format_code = f"IMPORT-{token}"
    with Session(pg_engine) as db:
        provisor = upsert_provisor_price_list(
            db=db,
            price_format_code=format_code,
            filial_id=991001,
            items=[
                {
                    "id": 1,
                    "goodsId": 2,
                    "filialId": 991001,
                    "distributorGoodsId": "PV-1",
                    "distributorGoodsName": "Provisor item",
                    "goodsPrice": "12.50",
                    "stored": "3",
                }
            ],
            run_matching=False,
        )
        provisor_id = int(provisor.id)
        assert db.scalar(
            select(func.count(CompetitorSourceVersion.id)).where(
                CompetitorSourceVersion.competitor_price_list_id == provisor_id,
                CompetitorSourceVersion.status == "published",
            )
        ) == 1
        manual = import_manual_price_list(
            db=db,
            price_format_code=format_code,
            content=b"SKU,Name,Price\nMANUAL-1,Manual item,9.99\n",
            filename="manual.csv",
            display_name="Manual test",
        )
        manual_id = int(manual["id"])
        assert db.scalar(
            select(func.count(CompetitorSourceVersion.id)).where(
                CompetitorSourceVersion.competitor_price_list_id == manual_id,
                CompetitorSourceVersion.status == "published",
            )
        ) == 1


def test_dependency_tracking_is_atomic_and_auto_learn_does_not_loop(
    pg_engine, monkeypatch
):
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        product = Product(code=f"DEP-{token}", name="Before", cost=1)
        db.add(product)
        db.commit()
        product_id = int(product.id)
    with Session(pg_engine) as db:
        baseline = db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == PRODUCT_CATALOG,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        )
        product = db.get(Product, product_id)
        product.name = "Rolled back"
        db.flush()
        db.rollback()
    with Session(pg_engine) as db:
        after_rollback = db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == PRODUCT_CATALOG,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        )
        assert after_rollback == baseline
        product = db.get(Product, product_id)
        product.name = "Committed"
        db.commit()
    with Session(pg_engine) as db:
        after_commit = db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == PRODUCT_CATALOG,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        )
        assert int(after_commit) == int(baseline) + 1
        auto = SourceGoodsMatch(
            source_type="provisor",
            distributor_goods_id=f"AUTO-{token}",
            product_id=product_id,
            match_method="fuzzy",
        )
        db.add(auto)
        db.commit()
        assert db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == MANUAL_CODE_MAPPINGS,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        ) is None
        auto.match_method = "manual_approved"
        auto.goods_id = 123
        db.commit()
        assert db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == MANUAL_CODE_MAPPINGS,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        ) == 1
        db.delete(auto)
        db.commit()
        assert db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == MANUAL_CODE_MAPPINGS,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        ) == 2


def test_matching_dependency_writer_model_coverage_is_explicit():
    assert {
        "Product",
        "ProductExtra",
        "InternalProductNormalized",
        "CompetitorCodeMapping",
        "ProvisorGoodsMap",
        "ProductSubstituteMatch",
        "VidmanProductMatch",
        "SourceGoodsMatch",
    } <= set(DEPENDENCY_MODEL_COVERAGE)


def test_concurrent_dependency_updates_are_not_lost(pg_engine, monkeypatch):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        before = db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == PRODUCT_CATALOG,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        ) or 0
    barrier = threading.Barrier(2)

    def add_product(index: int) -> None:
        with Session(pg_engine) as db:
            db.add(Product(code=f"CONCURRENT-{uuid.uuid4().hex}", name=f"Concurrent {index}", cost=1))
            barrier.wait(timeout=10)
            db.commit()

    with ThreadPoolExecutor(max_workers=2) as executor:
        list(executor.map(add_product, range(2)))
    with Session(pg_engine) as db:
        after = db.scalar(
            select(CompetitorMatchDependencyRevision.revision).where(
                CompetitorMatchDependencyRevision.scope_type == PRODUCT_CATALOG,
                CompetitorMatchDependencyRevision.scope_key == "global",
            )
        )
        assert int(after) == int(before) + 2


def test_shadow_match_uses_legacy_semantics_without_active_side_effects(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    code = f"SHADOW-{uuid.uuid4().hex}"
    with Session(pg_engine) as db:
        product = Product(code=code, name="Shadow Medicine", cost=1)
        db.add(product)
        db.flush()
        product_id = int(product.id)
        item_id = _add_legacy_item(db, source_id, code=code)
        db.commit()
    with Session(pg_engine) as db:
        nested = db.begin_nested()
        price_list = db.get(CompetitorPriceList, source_id)
        rematch_price_list_items_by_product(db=db, price_list=price_list)
        expected_item = db.get(CompetitorPriceListItem, item_id)
        expected = (
            expected_item.product_id,
            expected_item.match_type,
            expected_item.matched_sku,
        )
        nested.rollback()
        db.expire_all()
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        snapshot_id = int(snapshot.id)
        db.commit()
    with Session(pg_engine) as db:
        result = build_shadow_snapshot(
            db=db,
            snapshot_id=snapshot_id,
            expected_dependency_fingerprint=vector.fingerprint,
        )
        db.commit()
        assert result.row_count == 1
        decision = db.scalar(
            select(CompetitorMatchSnapshotItem).where(
                CompetitorMatchSnapshotItem.match_snapshot_id == snapshot_id
            )
        )
        assert (decision.product_id, decision.match_type, decision.matched_sku) == expected
    with Session(pg_engine) as db:
        active = db.get(CompetitorPriceListItem, item_id)
        assert active.product_id is None
        assert active.match_type == "unmatched"
        assert active.matched_sku == ""

        nested = db.begin_nested()
        price_list = db.get(CompetitorPriceList, source_id)
        rematch_price_list_items_by_product(db=db, price_list=price_list, flush=True)
        parity = compare_legacy_and_shadow_matches(
            db=db,
            competitor_price_list_id=source_id,
            snapshot_id=snapshot_id,
        )
        parity.assert_exact()
        nested.rollback()
        db.expire_all()


def test_lightweight_matcher_preserves_ties_manual_priority_and_item_reuse(
    pg_engine, monkeypatch
):
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        price_format = PriceFormat(code=f"ORDER-{token}", name="Ordering", branch="test")
        tie_product = Product(code=f"TIE-{token}", name="Tie Medicine 10 mg", cost=1)
        second_tie_product = Product(
            code=f"TIE-SECOND-{token}", name="Tie Medicine 10 mg", cost=1
        )
        mapped_product = Product(code=f"MAPPED-{token}", name="Mapped Product", cost=1)
        db.add_all([price_format, tie_product, second_tie_product, mapped_product])
        db.flush()
        source = CompetitorPriceList(
            price_format_id=price_format.id,
            source_type="vidman",
            source_key=f"ordering:{token}",
            display_name="Ordering source",
        )
        db.add(source)
        db.flush()
        first_tie = CompetitorPriceListItem(
            price_list_id=source.id,
            name=tie_product.name,
            raw_name=tie_product.name,
            distributor_goods_name=tie_product.name,
            distributor_goods_id=f"TIE-A-{token}",
            distributor_price=Decimal("10"),
            raw_json="{}",
        )
        second_tie = CompetitorPriceListItem(
            price_list_id=source.id,
            name=tie_product.name,
            raw_name=tie_product.name,
            distributor_goods_name=tie_product.name,
            distributor_goods_id=f"TIE-B-{token}",
            distributor_price=Decimal("10"),
            raw_json="{}",
        )
        unused_duplicate = CompetitorPriceListItem(
            price_list_id=source.id,
            name=tie_product.name,
            raw_name=tie_product.name,
            distributor_goods_name=tie_product.name,
            distributor_goods_id=f"TIE-C-{token}",
            distributor_price=Decimal("10"),
            raw_json="{}",
        )
        mapped_item = CompetitorPriceListItem(
            price_list_id=source.id,
            name=tie_product.name,
            raw_name=tie_product.name,
            distributor_goods_name=tie_product.name,
            distributor_goods_id=f"MAP-{token}",
            distributor_price=Decimal("1"),
            raw_json="{}",
        )
        rejected_item = CompetitorPriceListItem(
            price_list_id=source.id,
            name="Rejected Medicine 20 mg",
            raw_name="Rejected Medicine 20 mg",
            distributor_goods_name="Rejected Medicine 20 mg",
            distributor_goods_id=f"REJECT-{token}",
            distributor_price=Decimal("0.50"),
            raw_json='{"padding":"' + ("x" * 65536) + '"}',
        )
        db.add_all([
            first_tie,
            second_tie,
            unused_duplicate,
            mapped_item,
            rejected_item,
        ])
        db.flush()
        db.add_all([
            CompetitorCodeMapping(
                platform="vidman",
                source_external_key=mapped_item.distributor_goods_id,
                source_match_key=f"vidman:{mapped_item.distributor_goods_id}",
                source_name=mapped_item.name,
                our_product_id=mapped_product.id,
                our_sku=mapped_product.code,
                status="mapped",
                confidence=Decimal("100"),
            ),
            CompetitorCodeMapping(
                platform="vidman",
                source_external_key=rejected_item.distributor_goods_id,
                source_match_key=f"vidman:{rejected_item.distributor_goods_id}",
                source_name=rejected_item.name,
                status="rejected",
            ),
        ])
        db.commit()

        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=int(source.id)
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        build_shadow_snapshot(
            db=db,
            snapshot_id=int(snapshot.id),
            expected_dependency_fingerprint=vector.fingerprint,
        )

        item_identity_count_before = sum(
            isinstance(instance, CompetitorPriceListItem)
            for instance in db.identity_map.values()
        )
        rematch_price_list_items_by_product(db=db, price_list=source)
        item_identity_count_after = sum(
            isinstance(instance, CompetitorPriceListItem)
            for instance in db.identity_map.values()
        )

        # Equal-price ties retain source ID order.  Each product consumes a
        # different item and the extra duplicate remains unmatched.
        assert first_tie.product_id == tie_product.id
        assert first_tie.match_type == "exact_composite"
        assert second_tie.product_id == second_tie_product.id
        assert second_tie.match_type == "exact_composite"
        assert unused_duplicate.product_id is None
        assert unused_duplicate.match_type == "unmatched"
        # Manual mapping is applied before automatic candidate selection even
        # though this row is cheaper and has the same automatic match name.
        assert mapped_item.product_id == mapped_product.id
        assert mapped_item.match_type == "manual_code_mapping"
        assert rejected_item.product_id is None
        assert rejected_item.match_type == "manual_rejected"
        parity = compare_legacy_and_shadow_matches(
            db=db,
            competitor_price_list_id=int(source.id),
            snapshot_id=int(snapshot.id),
        )
        assert parity.legacy_count == 5
        assert parity.shadow_count == 5
        parity.assert_exact()
        assert item_identity_count_after == item_identity_count_before
        db.rollback()


def test_bounded_connection_pool_exhaustion_recovers(pg_engine):
    with pg_engine.connect() as connection:
        schema = connection.execute(text("SELECT current_schema()" )).scalar_one()
    bounded = create_engine(
        pg_engine.url,
        pool_size=2,
        max_overflow=0,
        pool_timeout=0.2,
        connect_args={"options": f"-csearch_path={schema},public"},
    )
    first = bounded.connect()
    second = bounded.connect()
    try:
        with pytest.raises(SQLAlchemyTimeoutError):
            bounded.connect()
        first.close()
        first = None
        with bounded.connect() as recovered:
            assert recovered.execute(text("SELECT 1")).scalar_one() == 1
    finally:
        if first is not None:
            first.close()
        second.close()
        bounded.dispose()


def _cache_materialization_fixture(pg_engine, monkeypatch, *, label: str) -> dict:
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        price_format = PriceFormat(code=f"CACHE-{label}-{token}", name="Cache", branch="test")
        first_product = Product(code=f"CACHE-A-{token}", name=f"Cache Medicine A {token}", cost=1)
        second_product = Product(code=f"CACHE-B-{token}", name=f"Cache Medicine B {token}", cost=1)
        db.add_all([price_format, first_product, second_product])
        db.flush()
        source = CompetitorPriceList(
            price_format_id=price_format.id,
            source_type="manual",
            source_key=f"cache:{label}:{token}",
            display_name="Cache source",
            supplier="Cache supplier",
            price_coefficient=Decimal("1.10"),
            items_count=2,
        )
        db.add(source)
        db.flush()
        db.add_all([
            CompetitorPriceListItem(
                price_list_id=source.id,
                name=first_product.name,
                raw_name=first_product.name,
                distributor_goods_name=first_product.name,
                distributor_goods_id=f"A-{token}",
                distributor_price=Decimal("10"),
                raw_json="{}",
            ),
            CompetitorPriceListItem(
                price_list_id=source.id,
                name=second_product.name,
                raw_name=second_product.name,
                distributor_goods_name=second_product.name,
                distributor_goods_id=f"B-{token}",
                distributor_price=Decimal("20"),
                raw_json="{}",
            ),
        ])
        db.add(PriceFormatCompetitorAssignment(
            price_format_id=price_format.id,
            competitor_price_list_id=source.id,
            is_active=True,
            coefficient=Decimal("1.10"),
        ))
        db.flush()
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=int(source.id)
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        build_shadow_snapshot(
            db=db,
            snapshot_id=int(snapshot.id),
            expected_dependency_fingerprint=vector.fingerprint,
        )
        db.commit()
        return {
            "price_format_id": int(price_format.id),
            "source_id": int(source.id),
            "source_version_id": int(capture.version_id),
            "snapshot_id": int(snapshot.id),
        }


def _physical_competitor_records(db: Session, price_format_id: int) -> list[dict]:
    return [
        {
            "productId": row.product_id,
            "sourceName": row.source_name,
            "supplier": row.supplier,
            "priceDate": row.price_date,
            "coefficient": row.coefficient,
            "sourcePrice": row.source_price,
            "matchType": row.match_type,
            "sourceItemId": row.source_item_id,
            "sourceGoodsId": row.source_goods_id,
            "sourceDistributorGoodsId": row.source_distributor_goods_id,
            "sourceManufacturer": row.source_manufacturer,
        }
        for row in db.scalars(
            select(CompetitorPrice)
            .where(CompetitorPrice.price_format_id == price_format_id)
            .where(CompetitorPrice.product_id.is_not(None))
            .order_by(CompetitorPrice.id.asc())
        )
    ]


def _delete_and_reinsert_source_rows(db: Session, source_id: int) -> tuple[list[int], list[int]]:
    table = CompetitorPriceListItem.__table__
    columns = [column for column in table.columns if column.name != "id"]
    saved = db.execute(
        select(*columns)
        .where(CompetitorPriceListItem.price_list_id == source_id)
        .order_by(CompetitorPriceListItem.id.asc())
    ).mappings().all()
    old_ids = list(db.scalars(
        select(CompetitorPriceListItem.id)
        .where(CompetitorPriceListItem.price_list_id == source_id)
        .order_by(CompetitorPriceListItem.id.asc())
    ))
    db.execute(
        delete(CompetitorPriceListItem).where(
            CompetitorPriceListItem.price_list_id == source_id
        )
    )
    new_ids = list(db.scalars(
        insert(table).returning(table.c.id),
        [dict(row) for row in saved],
    ))
    db.flush()
    return [int(value) for value in old_ids], [int(value) for value in new_ids]


def test_cache_read_set_materialization_matches_legacy_and_warm_reuses(
    pg_engine, monkeypatch
):
    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="hit")
    with Session(pg_engine) as db:
        legacy_summary = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=fixture["price_format_id"]
        )
        legacy_rows = _physical_competitor_records(db, fixture["price_format_id"])
        assert legacy_summary["_matchingCache"]["fullRematchInvocationCount"] == 1
        db.rollback()

    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
    matcher_calls = 0

    def forbidden_matcher(**_kwargs):
        nonlocal matcher_calls
        matcher_calls += 1
        raise AssertionError("full matcher must not run on a cache hit")

    monkeypatch.setattr(
        competitor_matching_service,
        "rematch_price_list_items_by_product",
        forbidden_matcher,
    )
    raw_source_selects: list[str] = []

    def observe_sql(_conn, _cursor, statement, _parameters, _context, _executemany):
        normalized = str(statement).casefold()
        if normalized.lstrip().startswith("select") and "competitor_price_list_items" in normalized:
            raw_source_selects.append(normalized)

    event.listen(pg_engine, "before_cursor_execute", observe_sql)
    try:
        with Session(pg_engine) as db:
            cold = rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )
            cold_rows = _physical_competitor_records(db, fixture["price_format_id"])
            compare_records(
                artifact="cache_competitors_prices",
                legacy=legacy_rows,
                shadow=cold_rows,
                identity_fields=("productId", "sourceName", "sourceItemId"),
            ).assert_exact()
            cache = cold["_matchingCache"]
            assert cache["hit"] is True
            assert cache["materializationReused"] is False
            assert cache["sourceVersionIds"] == [fixture["source_version_id"]]
            assert cache["matchingSnapshotIds"] == [fixture["snapshot_id"]]
            assert cache["fullRematchInvocationCount"] == 0
            assert cache["rawSourceRowsRead"] == 0
            first_fingerprint = cache["materializationFingerprint"]
            db.commit()

        with Session(pg_engine) as db:
            before_ids = list(db.scalars(
                select(CompetitorPrice.id)
                .where(CompetitorPrice.price_format_id == fixture["price_format_id"])
                .where(CompetitorPrice.product_id.is_not(None))
                .order_by(CompetitorPrice.id.asc())
            ))
            warm = rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )
            after_ids = list(db.scalars(
                select(CompetitorPrice.id)
                .where(CompetitorPrice.price_format_id == fixture["price_format_id"])
                .where(CompetitorPrice.product_id.is_not(None))
                .order_by(CompetitorPrice.id.asc())
            ))
            assert warm["_matchingCache"]["materializationReused"] is True
            assert warm["_matchingCache"]["reason"] == "materialization_unchanged"
            assert before_ids == after_ids
            db.rollback()

        with Session(pg_engine) as db:
            source = db.get(CompetitorPriceList, fixture["source_id"])
            source.price_coefficient = Decimal("1.25")
            assignment = db.scalar(
                select(PriceFormatCompetitorAssignment).where(
                    PriceFormatCompetitorAssignment.price_format_id
                    == fixture["price_format_id"],
                    PriceFormatCompetitorAssignment.competitor_price_list_id
                    == fixture["source_id"],
                )
            )
            assignment.coefficient = Decimal("1.25")
            db.commit()
        with Session(pg_engine) as db:
            changed = rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )
            changed_cache = changed["_matchingCache"]
            assert changed_cache["materializationReused"] is False
            assert changed_cache["materializationFingerprint"] != first_fingerprint
            assert {
                str(value)
                for value in db.scalars(
                    select(CompetitorPrice.coefficient).where(
                        CompetitorPrice.price_format_id == fixture["price_format_id"],
                        CompetitorPrice.product_id.is_not(None),
                    )
                )
            } == {"1.250000"}
            db.rollback()
    finally:
        event.remove(pg_engine, "before_cursor_execute", observe_sql)
    assert matcher_calls == 0
    assert raw_source_selects == []


def test_identical_delete_reinsert_reuses_snapshot_but_remaps_physical_ids(
    pg_engine, monkeypatch
):
    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="reinsert")
    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")

    with Session(pg_engine) as db:
        first = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=fixture["price_format_id"]
        )
        first_rows = _physical_competitor_records(db, fixture["price_format_id"])
        first_publication_id = int(db.scalar(
            select(CompetitorSourcePublication.id).where(
                CompetitorSourcePublication.competitor_price_list_id
                == fixture["source_id"],
                CompetitorSourcePublication.status == "published",
            )
        ))
        db.commit()
    assert first["_matchingCache"]["materializationReused"] is False

    with Session(pg_engine) as db:
        old_ids, new_ids = _delete_and_reinsert_source_rows(db, fixture["source_id"])
        assert set(old_ids).isdisjoint(new_ids)
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=fixture["source_id"]
        )
        assert capture.succeeded and capture.reused
        assert capture.version_id == fixture["source_version_id"]
        assert capture.publication_id != first_publication_id
        db.commit()

    matcher_calls = 0

    def forbidden_matcher(**_kwargs):
        nonlocal matcher_calls
        matcher_calls += 1
        raise AssertionError("identical reimport must reuse the ready snapshot")

    monkeypatch.setattr(
        competitor_matching_service,
        "rematch_price_list_items_by_product",
        forbidden_matcher,
    )
    with Session(pg_engine) as db:
        rebuilt = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=fixture["price_format_id"]
        )
        rebuilt_rows = _physical_competitor_records(db, fixture["price_format_id"])
        assert rebuilt["_matchingCache"]["materializationReused"] is False
        assert rebuilt["_matchingCache"]["matchingSnapshotIds"] == [
            fixture["snapshot_id"]
        ]
        assert {row["sourceItemId"] for row in rebuilt_rows} == set(new_ids)
        assert not ({row["sourceItemId"] for row in rebuilt_rows} & set(old_ids))
        without_lineage = lambda rows: [
            {key: value for key, value in row.items() if key != "sourceItemId"}
            for row in rows
        ]
        compare_records(
            artifact="identical_reimport_business_rows",
            legacy=without_lineage(first_rows),
            shadow=without_lineage(rebuilt_rows),
            identity_fields=("productId", "sourceName", "sourceDistributorGoodsId"),
        ).assert_exact()
        db.rollback()
    assert matcher_calls == 0


def test_duplicate_physical_rows_have_deterministic_one_to_one_lineage(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        for _ in range(2):
            _add_legacy_item(db, source_id, code=f"DUP-{token}", price="7.77")
        first = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        first_version = int(first.version_id)
        old_ids, new_ids = _delete_and_reinsert_source_rows(db, source_id)
        second = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        assert second.reused and int(second.version_id) == first_version
        mappings = db.execute(
            select(
                CompetitorSourceVersionItem.source_ordinal,
                CompetitorSourceVersionItem.duplicate_ordinal,
                CompetitorSourcePublicationItem.physical_source_item_id,
            )
            .join(
                CompetitorSourcePublicationItem,
                CompetitorSourcePublicationItem.source_version_item_id
                == CompetitorSourceVersionItem.id,
            )
            .where(
                CompetitorSourcePublicationItem.source_publication_id
                == int(second.publication_id)
            )
            .order_by(CompetitorSourceVersionItem.source_ordinal.asc())
        ).all()
        assert mappings == [(0, 0, new_ids[0]), (1, 1, new_ids[1])]
        assert set(old_ids).isdisjoint(new_ids)

        db.add(CompetitorSourcePublicationItem(
            source_publication_id=int(second.publication_id),
            source_version_item_id=db.scalar(
                select(CompetitorSourceVersionItem.id).where(
                    CompetitorSourceVersionItem.source_version_id == first_version,
                    CompetitorSourceVersionItem.source_ordinal == 0,
                )
            ),
            physical_source_item_id=new_ids[1],
        ))
        with pytest.raises(IntegrityError):
            db.flush()
        db.rollback()


def test_missing_physical_lineage_fails_closed(pg_engine, monkeypatch):
    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="missing-lineage")
    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
    with Session(pg_engine) as db:
        publication_id = int(db.scalar(
            select(CompetitorSourcePublication.id).where(
                CompetitorSourcePublication.competitor_price_list_id
                == fixture["source_id"],
                CompetitorSourcePublication.status == "published",
            )
        ))
        lineage_id = int(db.scalar(
            select(CompetitorSourcePublicationItem.id)
            .where(
                CompetitorSourcePublicationItem.source_publication_id
                == publication_id
            )
            .limit(1)
        ))
        db.execute(
            delete(CompetitorSourcePublicationItem).where(
                CompetitorSourcePublicationItem.id == lineage_id
            )
        )
        db.commit()
    with Session(pg_engine) as db:
        with pytest.raises(MatchingCacheConflict, match="physical publication is incomplete"):
            rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )


def test_concurrent_import_waits_for_cache_publication_then_forces_rematerialization(
    pg_engine, monkeypatch
):
    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="concurrent-reimport")
    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
    cache_entered = threading.Event()
    release_cache = threading.Event()
    import_waiting = threading.Event()

    def pause_cache(_conn, _cursor, statement, _parameters, _context, _many):
        if str(statement).casefold().lstrip().startswith("insert into competitors_prices"):
            cache_entered.set()
            assert release_cache.wait(timeout=10)

    def publish_cache() -> None:
        with Session(pg_engine) as db:
            rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )
            db.commit()

    def replace_source() -> tuple[list[int], list[int], int]:
        with Session(pg_engine) as db:
            import_waiting.set()
            acquire_transaction_write_locks(
                db, source_ids=[fixture["source_id"]], retry_timeout_seconds=10.0
            )
            old_ids, new_ids = _delete_and_reinsert_source_rows(
                db, fixture["source_id"]
            )
            capture = capture_and_publish_legacy_source_version(
                db=db, competitor_price_list_id=fixture["source_id"]
            )
            assert capture.succeeded and capture.reused
            db.commit()
            return old_ids, new_ids, int(capture.publication_id)

    event.listen(pg_engine, "before_cursor_execute", pause_cache)
    try:
        with ThreadPoolExecutor(max_workers=2) as executor:
            cache_future = executor.submit(publish_cache)
            assert cache_entered.wait(timeout=10)
            import_future = executor.submit(replace_source)
            assert import_waiting.wait(timeout=10)
            assert not import_future.done()
            release_cache.set()
            cache_future.result(timeout=15)
            old_ids, new_ids, publication_id = import_future.result(timeout=15)
    finally:
        event.remove(pg_engine, "before_cursor_execute", pause_cache)

    assert set(old_ids).isdisjoint(new_ids)
    with Session(pg_engine) as db:
        rebuilt = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=fixture["price_format_id"]
        )
        assert rebuilt["_matchingCache"]["materializationReused"] is False
        assert publication_id == int(db.scalar(
            select(CompetitorSourcePublication.id).where(
                CompetitorSourcePublication.competitor_price_list_id
                == fixture["source_id"],
                CompetitorSourcePublication.status == "published",
            )
        ))
        assert {
            row["sourceItemId"]
            for row in _physical_competitor_records(db, fixture["price_format_id"])
        } == set(new_ids)


def test_cache_materialization_failure_rolls_back_to_previous_publication(
    pg_engine, monkeypatch
):
    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="rollback")
    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")

    with Session(pg_engine) as db:
        rebuild_competitor_prices_for_selected(
            db=db, price_format_id=fixture["price_format_id"]
        )
        db.commit()
    with Session(pg_engine) as db:
        previous_rows = _physical_competitor_records(db, fixture["price_format_id"])
        previous = db.scalar(
            select(PriceFormatCompetitorMaterialization).where(
                PriceFormatCompetitorMaterialization.price_format_id
                == fixture["price_format_id"],
                PriceFormatCompetitorMaterialization.is_current.is_(True),
            )
        )
        previous_id = int(previous.id)
        previous_key = previous.materialization_key
        source = db.get(CompetitorPriceList, fixture["source_id"])
        source.price_coefficient = Decimal("1.30")
        db.commit()

    def fail_publication(_conn, _cursor, statement, _parameters, _context, _executemany):
        normalized = str(statement).casefold().lstrip()
        if normalized.startswith("insert into competitors_prices"):
            raise RuntimeError("simulated materialization resource exhaustion")

    event.listen(pg_engine, "before_cursor_execute", fail_publication)
    try:
        with Session(pg_engine) as db:
            with pytest.raises(RuntimeError, match="resource exhaustion"):
                rebuild_competitor_prices_for_selected(
                    db=db, price_format_id=fixture["price_format_id"]
                )
            db.rollback()
    finally:
        event.remove(pg_engine, "before_cursor_execute", fail_publication)

    with Session(pg_engine) as db:
        assert _physical_competitor_records(db, fixture["price_format_id"]) == previous_rows
        current = db.scalar(
            select(PriceFormatCompetitorMaterialization).where(
                PriceFormatCompetitorMaterialization.price_format_id
                == fixture["price_format_id"],
                PriceFormatCompetitorMaterialization.is_current.is_(True),
            )
        )
        assert int(current.id) == previous_id
        assert current.materialization_key == previous_key
        assert current.status == "published"
        assert db.scalar(
            select(func.count(PriceFormatCompetitorMaterialization.id)).where(
                PriceFormatCompetitorMaterialization.price_format_id
                == fixture["price_format_id"]
            )
        ) == 1


def test_cache_read_fallback_required_and_stale_contracts(pg_engine, monkeypatch):
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex
    with Session(pg_engine) as db:
        price_format = PriceFormat(code=f"CACHE-MISS-{token}", name="Miss", branch="test")
        product = Product(code=f"MISS-{token}", name=f"Miss Medicine {token}", cost=1)
        db.add_all([price_format, product])
        db.flush()
        source = CompetitorPriceList(
            price_format_id=price_format.id,
            source_type="manual",
            source_key=f"miss:{token}",
            display_name="Missing cache",
            items_count=1,
        )
        db.add(source)
        db.flush()
        db.add(CompetitorPriceListItem(
            price_list_id=source.id,
            name=product.name,
            raw_name=product.name,
            distributor_goods_name=product.name,
            distributor_goods_id=f"MISS-{token}",
            distributor_price=Decimal("10"),
            raw_json="{}",
        ))
        db.add(PriceFormatCompetitorAssignment(
            price_format_id=price_format.id,
            competitor_price_list_id=source.id,
            is_active=True,
        ))
        db.commit()
        price_format_id = int(price_format.id)

    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
    with Session(pg_engine) as db:
        with pytest.raises(MatchingCacheUnavailable, match="snapshot_absent"):
            rebuild_competitor_prices_for_selected(db=db, price_format_id=price_format_id)
        db.rollback()

    with Session(pg_engine) as db:
        publication_rebuild = rebuild_competitor_prices_for_selected(
            db=db,
            price_format_id=price_format_id,
            allow_matching_cache=False,
        )
        assert publication_rebuild["_matchingCache"]["hit"] is False
        assert publication_rebuild["_matchingCache"]["reason"] == (
            "cache_bypassed_during_source_publication"
        )
        assert publication_rebuild["_matchingCache"]["fullRematchInvocationCount"] == 1
        db.rollback()

    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "fallback")
    with Session(pg_engine) as db:
        fallback = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=price_format_id
        )
        assert fallback["_matchingCache"]["hit"] is False
        assert fallback["_matchingCache"]["reason"].startswith("snapshot_absent")
        assert fallback["_matchingCache"]["fullRematchInvocationCount"] == 1
        db.rollback()

    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="stale")
    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
    with Session(pg_engine) as db:
        bump_dependency_revision(db=db, scope_type=PRODUCT_CATALOG)
        db.commit()
    with Session(pg_engine) as db:
        with pytest.raises(MatchingCacheConflict, match="dependency revision mismatch"):
            rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )
        assert db.scalar(
            select(func.count(CompetitorPrice.id)).where(
                CompetitorPrice.price_format_id == fixture["price_format_id"],
                CompetitorPrice.product_id.is_not(None),
            )
        ) == 0
        db.rollback()


def test_new_publication_never_reuses_older_snapshot(pg_engine, monkeypatch):
    fixture = _cache_materialization_fixture(pg_engine, monkeypatch, label="republish")
    with Session(pg_engine) as db:
        item = db.scalar(
            select(CompetitorPriceListItem)
            .where(CompetitorPriceListItem.price_list_id == fixture["source_id"])
            .order_by(CompetitorPriceListItem.id.asc())
            .limit(1)
        )
        item.distributor_price = Decimal("11")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=fixture["source_id"]
        )
        assert capture.reused is False
        assert int(capture.version_id) != fixture["source_version_id"]
        db.commit()

    monkeypatch.setenv("MATCH_CACHE_READ", "true")
    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
    with Session(pg_engine) as db:
        with pytest.raises(MatchingCacheUnavailable, match="snapshot_absent"):
            rebuild_competitor_prices_for_selected(
                db=db, price_format_id=fixture["price_format_id"]
            )
        db.rollback()

    monkeypatch.setenv("MATCH_CACHE_READ_MODE", "fallback")
    with Session(pg_engine) as db:
        fallback = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=fixture["price_format_id"]
        )
        assert fallback["_matchingCache"]["hit"] is False
        assert fallback["_matchingCache"]["reason"].startswith("snapshot_absent")
        assert fallback["_matchingCache"]["fullRematchInvocationCount"] == 1
        db.rollback()


def test_shadow_snapshot_rejects_stale_dependencies(pg_engine, source_id, monkeypatch):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        _add_legacy_item(db, source_id, code=f"STALE-{uuid.uuid4().hex}")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        snapshot_id = int(snapshot.id)
        bump_dependency_revision(db=db, scope_type=PRODUCT_CATALOG)
        db.commit()
    with Session(pg_engine) as db:
        with pytest.raises(ShadowSnapshotConflict, match="stale dependency"):
            build_shadow_snapshot(
                db=db,
                snapshot_id=snapshot_id,
                expected_dependency_fingerprint=vector.fingerprint,
            )
        db.rollback()


def test_catalog_revision_change_during_matching_rejects_publication(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        _add_legacy_item(db, source_id, code=f"REVISION-RACE-{uuid.uuid4().hex}")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        snapshot_id = int(snapshot.id)
        db.commit()

    matcher_entered = threading.Event()
    release_matcher = threading.Event()
    original_matcher = shadow_service.rematch_price_list_items_by_product

    def paused_matcher(**kwargs):
        matcher_entered.set()
        assert release_matcher.wait(timeout=10)
        return original_matcher(**kwargs)

    monkeypatch.setattr(
        shadow_service, "rematch_price_list_items_by_product", paused_matcher
    )

    def run_shadow() -> str:
        with Session(pg_engine) as db:
            try:
                build_shadow_snapshot(
                    db=db,
                    snapshot_id=snapshot_id,
                    expected_dependency_fingerprint=vector.fingerprint,
                )
            except ShadowSnapshotConflict as exc:
                db.rollback()
                return str(exc)
            raise AssertionError("shadow publication unexpectedly succeeded")

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(run_shadow)
        assert matcher_entered.wait(timeout=10)
        with Session(pg_engine) as db:
            bump_dependency_revision(db=db, scope_type=PRODUCT_CATALOG)
            db.commit()
        release_matcher.set()
        assert future.result(timeout=15) == "dependencies changed during shadow matching"

    with Session(pg_engine) as db:
        assert db.get(CompetitorMatchSnapshot, snapshot_id).status == "pending"
        assert db.scalar(select(func.count(CompetitorMatchSnapshotItem.id)).where(
            CompetitorMatchSnapshotItem.match_snapshot_id == snapshot_id
        )) == 0


def test_crash_during_shadow_construction_leaves_no_partial_snapshot(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        item_id = _add_legacy_item(db, source_id, code=f"CRASH-{uuid.uuid4().hex}")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        snapshot_id = int(snapshot.id)
        db.commit()

    def crash(**_kwargs):
        item = _kwargs["db"].get(CompetitorPriceListItem, item_id)
        item.match_type = "partial_shadow_write"
        _kwargs["db"].flush()
        raise RuntimeError("simulated matcher crash")

    spool_paths: list[Path] = []
    original_named_temporary_file = shadow_service.tempfile.NamedTemporaryFile

    def track_spool(*args, **kwargs):
        handle = original_named_temporary_file(*args, **kwargs)
        spool_paths.append(Path(handle.name))
        return handle

    monkeypatch.setattr(shadow_service.tempfile, "NamedTemporaryFile", track_spool)
    monkeypatch.setattr(shadow_service, "rematch_price_list_items_by_product", crash)
    with Session(pg_engine) as db:
        with pytest.raises(RuntimeError, match="simulated matcher crash"):
            build_shadow_snapshot(
                db=db,
                snapshot_id=snapshot_id,
                expected_dependency_fingerprint=vector.fingerprint,
            )
        db.rollback()
    with Session(pg_engine) as db:
        active = db.get(CompetitorPriceListItem, item_id)
        assert active.product_id is None
        assert active.match_type == "unmatched"
        assert db.scalar(
            select(func.count(CompetitorMatchSnapshotItem.id)).where(
                CompetitorMatchSnapshotItem.match_snapshot_id == snapshot_id
            )
        ) == 0
    assert spool_paths
    assert all(not path.exists() for path in spool_paths)


def test_job_claim_retry_and_crash_recovery(pg_engine, source_id, monkeypatch):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        _add_legacy_item(db, source_id, code=f"JOB-{uuid.uuid4().hex}")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        job, reused = enqueue_matching_job(
            db=db,
            source_version_id=int(capture.version_id),
            dependency_fingerprint="f" * 64,
            max_attempts=2,
        )
        job_id = int(job.id)
        assert not reused
        db.commit()
    with Session(pg_engine) as db:
        claimed = claim_next_matching_job(db=db, worker_id="worker-a", lease_seconds=10)
        assert claimed is not None and int(claimed.id) == job_id
        db.commit()
    with Session(pg_engine) as db:
        status = fail_matching_job(
            db=db,
            job_id=job_id,
            worker_id="worker-a",
            error="temporary",
            retryable=True,
            base_retry_seconds=1,
        )
        assert status == "pending"
        job = db.get(CompetitorMatchJob, job_id)
        job.available_at = now_kz_naive() - timedelta(seconds=1)
        db.commit()
    with Session(pg_engine) as db:
        claimed = claim_next_matching_job(db=db, worker_id="worker-b", lease_seconds=10)
        claimed.lease_expires_at = now_kz_naive() - timedelta(seconds=1)
        db.commit()
    with Session(pg_engine) as db:
        assert recover_stale_matching_jobs(db=db) == 1
        db.commit()
    with Session(pg_engine) as db:
        recovered = db.get(CompetitorMatchJob, job_id)
        assert recovered.status == "failed"
        assert recovered.attempt_count == 2


def test_connection_loss_during_match_rolls_back_legacy_side_effects(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        item_id = _add_legacy_item(db, source_id, code=f"CONNECTION-{uuid.uuid4().hex}")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        snapshot_id = int(snapshot.id)
        db.commit()

    def lose_connection(**kwargs):
        item = kwargs["db"].get(CompetitorPriceListItem, item_id)
        item.match_type = "connection_was_lost"
        kwargs["db"].flush()
        raise OperationalError("shadow matcher", {}, ConnectionError("simulated disconnect"))

    monkeypatch.setattr(shadow_service, "rematch_price_list_items_by_product", lose_connection)
    with Session(pg_engine) as db:
        with pytest.raises(OperationalError):
            build_shadow_snapshot(
                db=db,
                snapshot_id=snapshot_id,
                expected_dependency_fingerprint=vector.fingerprint,
            )
        db.rollback()
    with Session(pg_engine) as db:
        assert db.get(CompetitorPriceListItem, item_id).match_type == "unmatched"
        assert db.scalar(select(func.count(CompetitorMatchSnapshotItem.id)).where(
            CompetitorMatchSnapshotItem.match_snapshot_id == snapshot_id
        )) == 0


def test_real_postgres_backend_termination_leaves_no_partial_shadow(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    with Session(pg_engine) as db:
        item_id = _add_legacy_item(db, source_id, code=f"TERMINATE-{uuid.uuid4().hex}")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=vector
        )
        snapshot_id = int(snapshot.id)
        db.commit()

    admin = create_engine(_postgres_url(), poolclass=NullPool)

    def terminate_backend(**kwargs):
        session = kwargs["db"]
        item = session.get(CompetitorPriceListItem, item_id)
        item.match_type = "must_rollback_after_backend_termination"
        session.flush()
        backend_pid = int(session.scalar(text("SELECT pg_backend_pid()")))
        with admin.begin() as connection:
            assert connection.execute(
                text("SELECT pg_terminate_backend(:pid)"), {"pid": backend_pid}
            ).scalar_one() is True
        session.execute(text("SELECT 1"))

    monkeypatch.setattr(
        shadow_service, "rematch_price_list_items_by_product", terminate_backend
    )
    try:
        with Session(pg_engine) as db:
            with pytest.raises(OperationalError):
                build_shadow_snapshot(
                    db=db,
                    snapshot_id=snapshot_id,
                    expected_dependency_fingerprint=vector.fingerprint,
                )
            db.rollback()
    finally:
        admin.dispose()

    with Session(pg_engine) as db:
        assert db.get(CompetitorPriceListItem, item_id).match_type == "unmatched"
        assert db.scalar(
            select(func.count(CompetitorMatchSnapshotItem.id)).where(
                CompetitorMatchSnapshotItem.match_snapshot_id == snapshot_id
            )
        ) == 0


def test_partial_snapshot_insert_failure_rolls_back_and_keeps_previous_ready_snapshot(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    code = f"PARTIAL-{uuid.uuid4().hex}"
    with Session(pg_engine) as db:
        product = Product(code=code, name=f"Medicine {code}", cost=1)
        db.add(product)
        db.flush()
        _add_legacy_item(db, source_id, code=code)
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        first_vector = build_dependency_vector(db=db)
        first, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=first_vector
        )
        first_id = int(first.id)
        db.commit()
    with Session(pg_engine) as db:
        build_shadow_snapshot(
            db=db,
            snapshot_id=first_id,
            expected_dependency_fingerprint=first_vector.fingerprint,
        )
        bump_dependency_revision(db=db, scope_type=PRODUCT_CATALOG)
        second_vector = build_dependency_vector(db=db)
        second, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=int(capture.version_id), vector=second_vector
        )
        second_id = int(second.id)
        db.commit()

    with Session(pg_engine) as db:
        original_bulk_insert = db.bulk_insert_mappings

        def fail_after_insert(mapper, mappings, *args, **kwargs):
            result = original_bulk_insert(mapper, mappings, *args, **kwargs)
            if mapper is CompetitorMatchSnapshotItem:
                db.flush()
                raise OperationalError(
                    "snapshot insert", {}, ConnectionError("simulated post-write disconnect")
                )
            return result

        monkeypatch.setattr(db, "bulk_insert_mappings", fail_after_insert)
        with pytest.raises(OperationalError):
            build_shadow_snapshot(
                db=db,
                snapshot_id=second_id,
                expected_dependency_fingerprint=second_vector.fingerprint,
            )
        db.rollback()

    with Session(pg_engine) as db:
        assert db.get(CompetitorMatchSnapshot, first_id).status == "ready"
        assert db.get(CompetitorMatchSnapshot, second_id).status == "pending"
        assert db.scalar(select(func.count(CompetitorMatchSnapshotItem.id)).where(
            CompetitorMatchSnapshotItem.match_snapshot_id == second_id
        )) == 0


def test_end_to_end_materialization_calculation_analytics_and_excel_parity(
    pg_engine, monkeypatch
):
    _enable_dual_write(monkeypatch)
    token = uuid.uuid4().hex

    def competitor_rows(db: Session, price_format_id: int) -> list[dict]:
        product_codes = dict(db.execute(select(Product.id, Product.code)).all())
        return [
            {
                "productCode": product_codes.get(row.product_id),
                "sourceName": row.source_name,
                "supplier": row.supplier,
                "priceDate": row.price_date,
                "coefficient": row.coefficient,
                "sourcePrice": row.source_price,
                "matchType": row.match_type,
                "sourceGoodsId": row.source_goods_id,
                "sourceDistributorGoodsId": row.source_distributor_goods_id,
                "sourceManufacturer": row.source_manufacturer,
            }
            for row in db.scalars(select(CompetitorPrice).where(
                CompetitorPrice.price_format_id == price_format_id,
                CompetitorPrice.product_id.is_not(None),
            ))
        ]

    def calculated_record(db: Session, product: Product, price_format: PriceFormat) -> dict:
        price, debug = calculate_price_for_product(
            db=db, product=product, price_format=price_format, as_of=now_kz_naive().date()
        )
        return {
            "productCode": product.code,
            "cost": debug["cost"],
            "basePrice": debug["base_price"],
            "competitorPrice": debug["competitor_price"],
            "priceFromCompetitor": debug["price_from_competitor"],
            "lowestCompetitorPrice": debug["lowest_competitor_price"],
            "chosenCompetitorPrice": debug["chosen_competitor_price"],
            "bendPercentUsed": debug["bend_percent_used"],
            "markupPercentUsed": debug["markup_percent_used"],
            "finalPrice": price,
            "appliedSource": debug.get("applied_source_name") or "",
            "reason": debug.get("reason") or "",
            "zone": debug.get("zone") or "",
            "usedSubstitute": bool(debug.get("used_substitute")),
            "usedPercentile": bool(debug.get("used_percentile")),
        }

    def persist_workflow_artifacts(
        db: Session, *, label: str, price_format: PriceFormat, product: Product,
        pricing_context_id: int, record: dict
    ) -> tuple[dict, list[list[object]]]:
        price_list = PriceList(
            number=f"{label}-{token}",
            price_format_id=price_format.id,
            run_snapshot_json="{}",
            run_sources_json="{}",
            run_lists_json="[]",
        )
        db.add(price_list)
        db.flush()
        db.add(CalculatedPrice(
            price_list_id=price_list.id,
            product_id=product.id,
            cost=record["cost"],
            base_price=record["basePrice"],
            competitor_price=record["competitorPrice"],
            price_from_competitor=record["priceFromCompetitor"],
            lowest_competitor_price=record["lowestCompetitorPrice"],
            chosen_competitor_price=record["chosenCompetitorPrice"],
            bend_percent_used=record["bendPercentUsed"],
            markup_percent_used=record["markupPercentUsed"],
            final_price=record["finalPrice"],
            applied_reason=record["reason"],
            applied_source_name=record["appliedSource"],
            used_substitute=record["usedSubstitute"],
            used_percentile=record["usedPercentile"],
            zone=record["zone"],
        ))
        run = PricingWorkflowRun(
            pricing_context_id=pricing_context_id,
            price_format_id=price_format.id,
            price_list_id=price_list.id,
            price_list_number=price_list.number,
            status="success",
        )
        db.add(run)
        db.flush()
        analytics = build_workflow_analytics(db=db, price_list_id=int(price_list.id))
        _name, content, _mime = export_workflow_run(
            db=db, run_id=int(run.id), fmt="xlsx", include="price"
        )
        sheet = load_workbook(io.BytesIO(content), data_only=True).active
        excel_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        return analytics, excel_rows

    with Session(pg_engine) as db:
        price_format = PriceFormat(
            code=f"E2E-{token}", name="Phase 2B.3 E2E", branch="test", progib=0
        )
        product = Product(
            code=f"E2E-SKU-{token}",
            name=f"Exact Medicine {token}",
            cost=Decimal("100"),
            provisor_goods_id=700001,
        )
        db.add_all([price_format, product])
        db.flush()
        pricing_context = PricingContext(
            branch_id=f"test-{token}", region="test", sales_channel="retail", name=f"ctx-{token}"
        )
        db.add(pricing_context)
        db.flush()
        db.add_all([
            MarkupRange(price_format_id=price_format.id, cost_from=0, cost_to=None, markup_percent=Decimal("0.15")),
            NoCompetitorMarkupRange(price_format_id=price_format.id, cost_from=0, cost_to=None, markup_percent=Decimal("0.03")),
            BendRange(price_format_id=price_format.id, price_from=0, bend_percent=Decimal("0")),
        ])

        source_specs = [
            ("manual", f"manual:{token}", Decimal("110"), Decimal("1.10"), ""),
            ("provisor", f"provisor:{token}", Decimal("105"), Decimal("1.00"), ""),
            ("vidman", f"account:test:main:{token}", Decimal("108"), Decimal("0.98"), ""),
            ("manual_vidman", f"authoritative:{token}", Decimal("107"), Decimal("1.02"), ""),
            ("emit", f"emit:{token}", Decimal("50"), Decimal("1.00"), "multi_price_per_sku"),
        ]
        sources: list[CompetitorPriceList] = []
        original_items: dict[int, dict] = {}
        snapshot_ids: list[int] = []
        for source_type, source_key, price, coefficient, percentile_mode in source_specs:
            source = CompetitorPriceList(
                price_format_id=price_format.id,
                source_type=source_type,
                source_key=source_key,
                display_name=source_key,
                supplier=source_key,
                competitor_name=source_key,
                price_coefficient=coefficient,
            )
            db.add(source)
            db.flush()
            authoritative = source_type == "manual_vidman"
            item = CompetitorPriceListItem(
                price_list_id=source.id,
                product_id=product.id if authoritative else None,
                name=product.name,
                raw_name=product.name,
                distributor_goods_name=product.name,
                distributor_goods_id=product.code,
                provisor_goods_id=700001 if source_type == "provisor" else None,
                distributor_price=price,
                matched_sku=product.code if authoritative else "",
                match_key=product.code if authoritative else "",
                match_type="manual_vidman_sku" if authoritative else "unmatched",
                match_score=Decimal("100") if authoritative else None,
                raw_json="{}",
            )
            db.add(item)
            db.add(PriceFormatCompetitorAssignment(
                price_format_id=price_format.id,
                competitor_price_list_id=source.id,
                is_active=True,
                coefficient=coefficient,
                percentile_mode=percentile_mode,
            ))
            db.flush()
            original_items[int(item.id)] = {
                LEGACY_FIELD_ALIASES.get(field, field): getattr(
                    item, LEGACY_FIELD_ALIASES.get(field, field)
                )
                for field in SOURCE_CONTENT_FIELDS
            }
            capture = capture_and_publish_legacy_source_version(
                db=db, competitor_price_list_id=int(source.id)
            )
            vector = build_dependency_vector(db=db)
            snapshot, _ = get_or_create_shadow_snapshot(
                db=db, source_version_id=int(capture.version_id), vector=vector
            )
            build_shadow_snapshot(
                db=db,
                snapshot_id=int(snapshot.id),
                expected_dependency_fingerprint=vector.fingerprint,
            )
            snapshot_ids.append(int(snapshot.id))
            sources.append(source)

        # Release source-version transaction locks before the independent
        # workflow materialization lease is acquired.
        db.commit()
        reimported_items: dict[int, dict] = {}
        for source, snapshot_id in zip(sources, snapshot_ids):
            old_ids, new_ids = _delete_and_reinsert_source_rows(
                db, int(source.id)
            )
            assert set(old_ids).isdisjoint(new_ids)
            republished = capture_and_publish_legacy_source_version(
                db=db, competitor_price_list_id=int(source.id)
            )
            assert republished.succeeded and republished.reused
            snapshot = db.get(CompetitorMatchSnapshot, snapshot_id)
            assert int(snapshot.source_version_id) == int(republished.version_id)
            assert len(old_ids) == len(new_ids) == 1
            reimported_items[new_ids[0]] = original_items[old_ids[0]]
        original_items = reimported_items
        db.commit()
        # Establish the legacy decision set explicitly.  The selected-source
        # rebuild may legitimately skip raw matching for source modes whose
        # materialization policy changes independently of the matcher.
        for source in sources[:-1]:
            rematch_price_list_items_by_product(db=db, price_list=source)
        db.flush()
        for source, snapshot_id in zip(sources[:-1], snapshot_ids[:-1]):
            compare_legacy_and_shadow_matches(
                db=db,
                competitor_price_list_id=int(source.id),
                snapshot_id=snapshot_id,
            ).assert_exact()
        # Direct matcher calls use transaction-scoped source locks.
        db.commit()
        sync_selected_competitor_configs(db=db, price_format_id=int(price_format.id))
        rebuild_competitor_prices_for_selected(
            db=db, price_format_id=int(price_format.id)
        )
        legacy_competitors = competitor_rows(db, int(price_format.id))
        legacy_calculated = calculated_record(db, product, price_format)
        legacy_analytics, legacy_excel = persist_workflow_artifacts(
            db=db,
            label="LEGACY",
            price_format=price_format,
            product=product,
            pricing_context_id=int(pricing_context.id),
            record=legacy_calculated,
        )

        for item_id, values in original_items.items():
            item = db.get(CompetitorPriceListItem, item_id)
            for field, value in values.items():
                setattr(item, field, value)
        db.flush()
        # Phase 3B invalidates physical lineage on matching-relevant row
        # UPDATEs.  The rows now exactly match the original immutable
        # generation again, so republish its one-to-one physical lineage
        # before exercising the required-cache workflow.
        for source in sources:
            republished = capture_and_publish_legacy_source_version(
                db=db, competitor_price_list_id=int(source.id)
            )
            assert republished.succeeded and republished.reused
        db.commit()
        db.execute(delete(CompetitorPrice).where(
            CompetitorPrice.price_format_id == price_format.id,
            CompetitorPrice.product_id.is_not(None),
        ))
        db.flush()
        monkeypatch.setenv("MATCH_CACHE_READ", "true")
        monkeypatch.setenv("MATCH_CACHE_READ_MODE", "required")
        monkeypatch.setattr(
            competitor_matching_service,
            "rematch_price_list_items_by_product",
            lambda **_kwargs: (_ for _ in ()).throw(
                AssertionError("cache E2E must not invoke full-source matching")
            ),
        )
        cache_summary = rebuild_competitor_prices_for_selected(
            db=db, price_format_id=int(price_format.id)
        )
        assert cache_summary["_matchingCache"]["hit"] is True
        assert cache_summary["_matchingCache"]["fullRematchInvocationCount"] == 0
        assert cache_summary["_matchingCache"]["rawSourceRowsRead"] == 0
        assert set(cache_summary["_matchingCache"]["matchingSnapshotIds"]) == set(snapshot_ids[:-1])
        shadow_competitors = competitor_rows(db, int(price_format.id))
        shadow_calculated = calculated_record(db, product, price_format)
        shadow_analytics, shadow_excel = persist_workflow_artifacts(
            db=db,
            label="SHADOW",
            price_format=price_format,
            product=product,
            pricing_context_id=int(pricing_context.id),
            record=shadow_calculated,
        )

        competitor_report = compare_records(
            artifact="competitors_prices",
            legacy=legacy_competitors,
            shadow=shadow_competitors,
            identity_fields=("productCode", "sourceName", "sourceDistributorGoodsId"),
        )
        calculated_report = compare_records(
            artifact="calculated_prices",
            legacy=[legacy_calculated],
            shadow=[shadow_calculated],
            identity_fields=("productCode",),
        )
        competitor_report.assert_exact()
        calculated_report.assert_exact()
        assert legacy_analytics == shadow_analytics
        assert legacy_excel == shadow_excel
        assert all(not row["sourceName"].startswith("emit:") for row in legacy_competitors)
        db.rollback()


def test_concurrent_import_waits_for_shadow_and_publishes_new_source_version(
    pg_engine, source_id, monkeypatch
):
    _enable_dual_write(monkeypatch)
    code = f"CONCURRENT-SHADOW-{uuid.uuid4().hex}"
    with Session(pg_engine) as db:
        db.add(Product(code=code, name=f"Medicine {code}", cost=1))
        db.flush()
        item_id = _add_legacy_item(db, source_id, code=code, price="10")
        capture = capture_and_publish_legacy_source_version(
            db=db, competitor_price_list_id=source_id
        )
        old_version_id = int(capture.version_id)
        vector = build_dependency_vector(db=db)
        snapshot, _ = get_or_create_shadow_snapshot(
            db=db, source_version_id=old_version_id, vector=vector
        )
        snapshot_id = int(snapshot.id)
        db.commit()

    matcher_entered = threading.Event()
    release_matcher = threading.Event()
    import_waiting = threading.Event()
    original_matcher = shadow_service.rematch_price_list_items_by_product

    def paused_matcher(**kwargs):
        matcher_entered.set()
        assert release_matcher.wait(timeout=10)
        return original_matcher(**kwargs)

    monkeypatch.setattr(shadow_service, "rematch_price_list_items_by_product", paused_matcher)

    def run_shadow() -> None:
        with Session(pg_engine) as db:
            build_shadow_snapshot(
                db=db,
                snapshot_id=snapshot_id,
                expected_dependency_fingerprint=vector.fingerprint,
            )
            db.commit()

    def run_import() -> int:
        with Session(pg_engine) as db:
            import_waiting.set()
            acquire_transaction_write_locks(
                db,
                source_ids=[source_id],
                retry_timeout_seconds=10.0,
            )
            item = db.get(CompetitorPriceListItem, item_id)
            item.distributor_price = Decimal("11")
            db.flush()
            changed = capture_and_publish_legacy_source_version(
                db=db, competitor_price_list_id=source_id
            )
            assert changed.succeeded
            db.commit()
            return int(changed.version_id)

    with ThreadPoolExecutor(max_workers=2) as executor:
        shadow_future = executor.submit(run_shadow)
        assert matcher_entered.wait(timeout=10)
        # A normal read-only request uses another bounded-pool connection while
        # the shadow worker owns its transaction-scoped source lock.
        with Session(pg_engine) as query_db:
            assert query_db.scalar(
                select(func.count(CompetitorPriceListItem.id)).where(
                    CompetitorPriceListItem.price_list_id == source_id
                )
            ) == 1
        import_future = executor.submit(run_import)
        assert import_waiting.wait(timeout=10)
        assert not import_future.done()
        release_matcher.set()
        shadow_future.result(timeout=15)
        new_version_id = import_future.result(timeout=15)

    assert new_version_id != old_version_id
    with Session(pg_engine) as db:
        assert db.get(CompetitorMatchSnapshot, snapshot_id).status == "ready"
        assert db.get(CompetitorSourceVersion, old_version_id).status == "superseded"
        assert db.get(CompetitorSourceVersion, new_version_id).status == "published"
