from __future__ import annotations

import io
import json
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.db import Base
from backend.app.deps import ROLE_ADMIN, get_current_user, get_db
from backend.app.main import app
from backend.app.models import (
    AppUser,
    CalculatedPrice,
    CompetitorPrice,
    CompetitorPriceList,
    CompetitorPriceListItem,
    PriceFormat,
    PriceFormatCompetitorAssignment,
    PriceList,
    Product,
)
from backend.app.services.competitor_assignments import repair_manual_vidman_assignment_modes
from backend.app.services.competitor_matching import rebuild_competitor_prices_for_selected
from backend.app.services.competitor_source_config import (
    MULTI_PRICE_PERCENTILE_MODE,
    default_percentile_mode_for_source,
    effective_percentile_mode,
)

EMITI_RU = "\u042d\u043c\u0438\u0442\u0438"


def _session_factory():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return engine, sessionmaker(bind=engine)


def test_manual_emiti_is_physical_but_genuine_emit_stays_aggregated():
    manual_emiti = CompetitorPriceList(
        source_type="manual_vidman",
        source_key="emiti",
        display_name=EMITI_RU,
        supplier="Emit International",
    )
    genuine_emit = CompetitorPriceList(
        source_type="emit",
        source_key="emit:1108",
        display_name="Emit International Kostanay",
    )

    assert default_percentile_mode_for_source(manual_emiti) == ""
    assert effective_percentile_mode(manual_emiti, MULTI_PRICE_PERCENTILE_MODE) == ""
    assert default_percentile_mode_for_source(genuine_emit) == MULTI_PRICE_PERCENTILE_MODE
    assert effective_percentile_mode(genuine_emit, "") == MULTI_PRICE_PERCENTILE_MODE


def test_repair_only_normalizes_stale_manual_vidman_assignments():
    engine, SessionLocal = _session_factory()
    db = SessionLocal()
    try:
        price_format = PriceFormat(code="MV-REPAIR", name="Manual Vidman repair", branch="1")
        manual_emiti = CompetitorPriceList(
            source_type="manual_vidman",
            source_key="emiti",
            display_name=EMITI_RU,
        )
        genuine_emit = CompetitorPriceList(
            source_type="emit",
            source_key="emit:1108",
            display_name="Emit International Kostanay",
        )
        db.add_all([price_format, manual_emiti, genuine_emit])
        db.flush()
        manual_assignment = PriceFormatCompetitorAssignment(
            price_format_id=price_format.id,
            competitor_price_list_id=manual_emiti.id,
            percentile_mode=MULTI_PRICE_PERCENTILE_MODE,
        )
        emit_assignment = PriceFormatCompetitorAssignment(
            price_format_id=price_format.id,
            competitor_price_list_id=genuine_emit.id,
            percentile_mode=MULTI_PRICE_PERCENTILE_MODE,
        )
        db.add_all([manual_assignment, emit_assignment])
        db.flush()

        repaired_ids = repair_manual_vidman_assignment_modes(db=db, price_format_id=price_format.id)

        assert repaired_ids == [manual_assignment.id]
        assert manual_assignment.percentile_mode == ""
        assert emit_assignment.percentile_mode == MULTI_PRICE_PERCENTILE_MODE
    finally:
        db.close()
        engine.dispose()


def test_existing_manual_vidman_rows_rebuild_and_export_without_reimport():
    engine, SessionLocal = _session_factory()
    db = SessionLocal()
    try:
        product = Product(code="SKU-100", name="Authoritative product", cost=Decimal("100"))
        price_format = PriceFormat(code="MV-EXISTING", name="Manual Vidman existing rows", branch="1")
        db.add_all([product, price_format])
        db.flush()

        emiti = CompetitorPriceList(
            source_type="manual_vidman",
            source_key="emiti",
            display_name=EMITI_RU,
            supplier=EMITI_RU,
            items_count=1,
        )
        stofarm = CompetitorPriceList(
            source_type="manual_vidman",
            source_key="stofarm",
            display_name="Stofarm",
            supplier="Stofarm",
            items_count=1,
        )
        db.add_all([emiti, stofarm])
        db.flush()

        # These are valid rows already persisted by the importer. Their raw
        # names cannot match the Product, proving that rebuild preserves the
        # importer-established Product.code mapping.
        emiti_item = CompetitorPriceListItem(
            price_list_id=emiti.id,
            product_id=product.id,
            name="Unrelated Emit row name",
            raw_name="Unrelated Emit row name",
            distributor_price=Decimal("87.50"),
            matched_sku=product.code,
            match_key=product.code,
            match_type="manual_vidman_sku",
            match_score=Decimal("100"),
        )
        stofarm_item = CompetitorPriceListItem(
            price_list_id=stofarm.id,
            product_id=product.id,
            name="Unrelated Stofarm row name",
            raw_name="Unrelated Stofarm row name",
            distributor_price=Decimal("91.25"),
            matched_sku=product.code,
            match_key=product.code,
            match_type="manual_vidman_sku",
            match_score=Decimal("100"),
        )
        db.add_all([emiti_item, stofarm_item])
        db.add_all(
            [
                PriceFormatCompetitorAssignment(
                    price_format_id=price_format.id,
                    competitor_price_list_id=emiti.id,
                    percentile_mode=MULTI_PRICE_PERCENTILE_MODE,
                ),
                PriceFormatCompetitorAssignment(
                    price_format_id=price_format.id,
                    competitor_price_list_id=stofarm.id,
                    percentile_mode="",
                ),
            ]
        )
        db.flush()

        rebuild_competitor_prices_for_selected(db=db, price_format_id=price_format.id)
        db.flush()

        repaired_assignment = db.execute(
            select(PriceFormatCompetitorAssignment).where(
                PriceFormatCompetitorAssignment.competitor_price_list_id == emiti.id
            )
        ).scalar_one()
        assert repaired_assignment.percentile_mode == ""
        assert emiti_item.product_id == product.id
        assert emiti_item.matched_sku == product.code
        assert emiti_item.match_type == "manual_vidman_sku"
        prices = db.execute(
            select(CompetitorPrice)
            .where(CompetitorPrice.price_format_id == price_format.id)
            .order_by(CompetitorPrice.source_name.asc())
        ).scalars().all()
        assert {row.source_name: float(row.source_price) for row in prices} == {
            "manual_vidman:emiti": 87.5,
            "manual_vidman:stofarm": 91.25,
        }

        generated = PriceList(
            number="MV-EXISTING-PL",
            price_format_id=price_format.id,
            run_sources_json=json.dumps(
                {
                    "selectedCompetitorSources": [
                        {
                            "id": emiti.id,
                            "sourceType": emiti.source_type,
                            "sourceKey": emiti.source_key,
                            "displayName": emiti.display_name,
                            "coefficient": 1,
                        },
                        {
                            "id": stofarm.id,
                            "sourceType": stofarm.source_type,
                            "sourceKey": stofarm.source_key,
                            "displayName": stofarm.display_name,
                            "coefficient": 1,
                        },
                    ]
                },
                ensure_ascii=False,
            ),
        )
        db.add(generated)
        db.flush()
        db.add(
            CalculatedPrice(
                price_list_id=generated.id,
                product_id=product.id,
                cost=Decimal("100"),
                base_price=Decimal("120"),
                final_price=Decimal("119"),
            )
        )
        db.commit()

        def override_db():
            session = SessionLocal()
            try:
                yield session
            finally:
                session.close()

        app.dependency_overrides[get_db] = override_db
        app.dependency_overrides[get_current_user] = lambda: AppUser(
            id=1,
            username="admin",
            role=ROLE_ADMIN,
            is_active=True,
        )
        client = TestClient(app)
        response = client.get("/api/generated-price-lists/MV-EXISTING-PL/export.xlsx")
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content), data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        headers = list(rows[0])
        values = list(rows[1])
        assert values[headers.index(EMITI_RU)] == 87.5
        assert values[headers.index("Stofarm")] == 91.25
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        db.close()
        engine.dispose()
