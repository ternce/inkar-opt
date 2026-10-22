from __future__ import annotations

import json
import asyncio
import base64
import csv
import io
from datetime import date
from decimal import Decimal

import pytest
import httpx
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.app.main as main
from backend.app.config import get_settings
from backend.app.db import Base
from backend.app.deps import require_admin_access, require_write_access
from backend.app.main import (
    _stock_quantities_for_price_list,
    export_price_list_csv,
    export_price_list_xlsx,
    get_generated_price_list_items,
)
from backend.app.models import (
    AppUser,
    BranchStock,
    CalculatedPrice,
    PriceFormat,
    PriceList,
    Product,
    ReferenceImportJob,
    ReferenceUpdateStatus,
    RefreshJob,
    RefreshLock,
    SapBalanceScheduledRun,
    StockSnapshot,
)
from backend.app.services.references.sap_balance import (
    SAP_PLANTS,
    SapBalanceItem,
    SapBalancePayload,
    branch_id_for_plant,
    fetch_balance,
    parse_balance_response,
)
from backend.app.services.references.stock_snapshots import (
    CanonicalStockRow,
    canonical_rows_from_sap,
    import_sap_balance_payload,
    publish_stock_snapshot,
)
from backend.app.services.references.imports import import_reference_excel
from backend.app.services.references.stock_diagnostics import compare_stock_quantities
from backend.app.services.sku import normalize_sku
from backend.scripts.diagnose_sap_balance_json import mapping_report, write_reconciliation_csv
from backend.app.services.refresh_queue import (
    claim_next_queued_refresh_job,
    create_queued_provisor_auto_refresh_job,
    create_queued_sap_balance_job,
)


def _session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _response(plant: str, items: list[dict]) -> dict:
    return {"MT_BalanceResponse": {"status": "SUCCESS", "rowCount": len(items), "items": items}}


def _item(material: str, quantity: str, *, plant: str = "1004", lgort: str = "1000", charge: str = "") -> dict:
    return {
        "material": material,
        "plant": plant,
        "lgort": lgort,
        "charge": charge,
        "endingStockQuantity": quantity,
        "endingStockValue": "999999.99",
    }


def _job(db, branch_id: str, source_type: str = "excel") -> ReferenceImportJob:
    row = ReferenceImportJob(
        data_type="stock",
        branch_ids_json=json.dumps([branch_id]),
        filename="fixture.xlsx",
        source_type=source_type,
        status="running",
    )
    db.add(row)
    db.flush()
    return row


def test_authoritative_mapping_has_all_12_explicit_plants(monkeypatch):
    monkeypatch.setenv("SAP_BALANCE_BRANCH_MAP_JSON", "{}")
    get_settings.cache_clear()
    settings = get_settings()
    try:
        assert [row.plant for row in SAP_PLANTS] == [
            "1001", "1002", "1003", "1004", "1005", "1006",
            "1007", "1008", "1009", "1010", "1011", "1013",
        ]
        assert {row.plant: branch_id_for_plant(settings, row.plant) for row in SAP_PLANTS} == {
            "1001": "1", "1002": "2", "1003": "5", "1004": "Есик",
            "1005": "6", "1006": "8", "1007": "Семей", "1008": "16",
            "1009": "10", "1010": "3", "1011": "4", "1013": "17",
        }
        with pytest.raises(ValueError, match="1012"):
            branch_id_for_plant(settings, "1012")
    finally:
        get_settings.cache_clear()


def test_sap_parser_preserves_leading_zeroes_and_decimal_quantity():
    payload = parse_balance_response(
        _response("1004", [_item("0000000001005213", "32.7500")]),
        requested_plant="1004",
        effective_date=date(2026, 10, 14),
    )
    assert payload.items[0].material == "0000000001005213"
    assert payload.items[0].quantity == Decimal("32.7500")


def test_fetch_balance_uses_basic_auth_and_only_requested_plant_lgort(monkeypatch):
    monkeypatch.setenv("SAP_BALANCE_ENABLED", "true")
    monkeypatch.setenv("SAP_BALANCE_ALLOW_PIQ_REQUESTS", "true")
    monkeypatch.setenv("SAP_BALANCE_USERNAME", "local-user")
    monkeypatch.setenv("SAP_BALANCE_PASSWORD", "local-pass")
    monkeypatch.setenv("SAP_BALANCE_COMPANY_CODE", "LOCAL")
    get_settings.cache_clear()
    captured = {}

    def handler(request: httpx.Request) -> httpx.Response:
        captured["body"] = json.loads(request.content)
        captured["authorization"] = request.headers.get("Authorization")
        return httpx.Response(200, json=_response("1004", [_item("0000000001005213", "32")]))

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await fetch_balance(
                get_settings(),
                plant="1004",
                date_from=date(2026, 10, 14),
                date_to=date(2026, 10, 14),
                client=client,
            )

    try:
        result = asyncio.run(run())
        assert result.plant == "1004"
        assert captured["body"]["plants"] == ["1004"]
        assert captured["body"]["lgort"] == ["1000"]
        assert captured["authorization"] == "Basic " + base64.b64encode(b"local-user:local-pass").decode()
    finally:
        get_settings.cache_clear()


def test_sap_mutation_access_requires_admin():
    non_admin = AppUser(username="pricing", role="pricing_manager", is_active=True)
    with pytest.raises(HTTPException) as exc_info:
        require_admin_access(non_admin)
    assert exc_info.value.status_code == 403


def test_sap_submission_routes_are_admin_only_while_manual_imports_keep_write_access():
    dependency_calls = {
        route.path: {dependency.call for dependency in route.dependant.dependencies}
        for route in main.app.routes
        if hasattr(route, "dependant")
    }
    for path in {
        "/api/references/sap/refresh",
        "/api/references/sap/catalog-preview",
        "/api/references/sap/import-payload",
    }:
        assert require_admin_access in dependency_calls[path]
    for path in {"/api/references/import", "/api/references/import/batch"}:
        assert require_write_access in dependency_calls[path]
        assert require_admin_access not in dependency_calls[path]


def test_offline_sap_endpoints_are_disabled_independently_by_default(monkeypatch):
    monkeypatch.setattr(main.settings, "sap_offline_import_enabled", False)
    monkeypatch.setattr(main.settings, "sap_balance_enabled", True)
    admin = AppUser(username="admin", role="admin", is_active=True)

    with pytest.raises(HTTPException) as preview_error:
        asyncio.run(
            main.post_reference_sap_catalog_preview(
                plant="1004", effective_date=None, file=None, db=_session(), current_user=admin
            )
        )
    with pytest.raises(HTTPException) as publication_error:
        asyncio.run(
            main.post_reference_sap_import_payload(
                plant="1004",
                effective_date=None,
                approve_catalog_registration=True,
                expected_sha256="0" * 64,
                force_manual_override=False,
                file=None,
                db=_session(),
                current_user=admin,
            )
        )
    assert preview_error.value.status_code == 503
    assert publication_error.value.status_code == 503


def test_piq_endpoint_requires_separate_explicit_authorization():
    settings = get_settings().model_copy(
        update={
            "sap_balance_enabled": True,
            "sap_balance_target_environment": "piq",
            "sap_balance_allow_piq_requests": False,
            "sap_balance_base_url": "https://sap-piq.inkar.kz:53443/RESTAdapter/getBalance",
            "sap_balance_username": "fixture-user",
            "sap_balance_password": "fixture-password",
            "sap_balance_company_code": "LOCAL",
        }
    )
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        return httpx.Response(500, request=request)

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await fetch_balance(
                settings,
                plant="1004",
                date_from=date(2026, 10, 14),
                date_to=date(2026, 10, 14),
                client=client,
            )

    with pytest.raises(ValueError, match="SAP_BALANCE_ALLOW_PIQ_REQUESTS=true"):
        asyncio.run(run())
    assert calls == 0


def test_pip_target_rejects_piq_endpoint_before_network_io():
    settings = get_settings().model_copy(
        update={
            "sap_balance_enabled": True,
            "sap_balance_target_environment": "pip",
            "sap_balance_allow_piq_requests": True,
            "sap_balance_base_url": "https://sap-piq.inkar.kz:53443/RESTAdapter/getBalance",
            "sap_balance_username": "fixture-user",
            "sap_balance_password": "fixture-password",
            "sap_balance_company_code": "LOCAL",
        }
    )

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: httpx.Response(500, request=request))) as client:
            return await fetch_balance(
                settings,
                plant="1004",
                date_from=date(2026, 10, 14),
                date_to=date(2026, 10, 14),
                client=client,
            )

    with pytest.raises(ValueError, match="not a PIP endpoint"):
        asyncio.run(run())


def test_sap_queue_keeps_explicit_plants_and_effective_date():
    db = _session()
    job = create_queued_sap_balance_job(
        db,
        plants=["1001", "1013", "1001"],
        effective_date="2026-10-16",
        requested_by="phase5-test",
    )
    metadata = json.loads(job.metadata_json)
    assert job.source_type == "sap"
    assert job.mode == "stock"
    assert job.status == "queued"
    assert metadata["plants"] == ["1001", "1013"]
    assert metadata["effective_date"] == "2026-10-16"
    assert metadata["force_manual_override"] is False


def test_scheduled_sap_enqueue_is_durable_and_deduplicated_for_all_approved_branches(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    monkeypatch.setattr(main, "SessionLocal", Session)
    monkeypatch.setattr(
        main,
        "settings",
        main.settings.model_copy(
            update={
                "sap_balance_enabled": True,
                "sap_balance_schedule_enabled": True,
                "sap_balance_branch_map_json": "{}",
            }
        ),
    )
    asyncio.run(main._queue_scheduled_sap_balance())
    asyncio.run(main._queue_scheduled_sap_balance())
    with Session() as db:
        jobs = db.scalars(select(RefreshJob).where(RefreshJob.source_type == "sap")).all()
        reservations = db.scalars(select(SapBalanceScheduledRun)).all()
        assert len(jobs) == 1
        assert len(reservations) == 12
        assert {row.plant for row in reservations} == {row.plant for row in SAP_PLANTS}
        assert {row.refresh_job_id for row in reservations} == {jobs[0].id}


def test_sap_worker_completion_releases_shared_mutex_for_provisor(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    material = "1005213".zfill(18)
    with Session() as db:
        db.add(Product(code=normalize_sku(material), name="Esik worker fixture", cost=1))
        db.commit()
        queued = create_queued_sap_balance_job(
            db,
            plants=["1004"],
            effective_date="2026-10-15",
            requested_by="test",
        )

    async def fake_fetch_balance(_settings, *, plant, date_from, date_to):
        return parse_balance_response(
            _response(plant, [_item(material, "32", plant=plant)]),
            requested_plant=plant,
            effective_date=date_to,
        )

    monkeypatch.setattr(main, "SessionLocal", Session)
    monkeypatch.setattr(main, "fetch_balance", fake_fetch_balance)
    monkeypatch.setattr(
        main,
        "settings",
        main.settings.model_copy(
            update={
                "sap_balance_enabled": True,
                "sap_manual_precedence_hours": 24,
            }
        ),
    )

    with Session() as db:
        claimed, owner_token = claim_next_queued_refresh_job(db, worker_id="worker-sap")
        assert claimed is not None and claimed.id == queued.id and owner_token
    asyncio.run(main._execute_claimed_refresh_job(int(queued.id), owner_token=owner_token))

    with Session() as db:
        completed = db.get(RefreshJob, queued.id)
        assert completed is not None and completed.status == "success"
        assert db.scalars(select(RefreshLock)).all() == []
        status = db.scalar(
            select(ReferenceUpdateStatus)
            .where(ReferenceUpdateStatus.branch_id == "Есик")
            .where(ReferenceUpdateStatus.data_type == "stock")
        )
        assert status is not None and status.active_source_type == "sap"
        provisor = create_queued_provisor_auto_refresh_job(db, mode="selected", requested_by="test")

    with Session() as db:
        next_job, next_token = claim_next_queued_refresh_job(db, worker_id="worker-provisor")
        assert next_job is not None and next_job.id == provisor.id and next_token


@pytest.mark.parametrize(
    ("body", "message"),
    [
        ({"MT_BalanceResponse": {"status": "ERROR", "rowCount": 0, "items": []}}, "business error"),
        ({"MT_BalanceResponse": {"status": "SUCCESS", "rowCount": 2, "items": []}}, "incomplete"),
        (_response("1004", [_item("1", "1", plant="1005")]), "plant mismatch"),
        (_response("1004", [_item("1", "1", lgort="4000")]), "LGORT mismatch"),
    ],
)
def test_sap_parser_rejects_unsafe_responses(body, message):
    with pytest.raises(ValueError, match=message):
        parse_balance_response(body, requested_plant="1004", effective_date=date(2026, 10, 14))


def test_sap_repeat_visible_records_are_preserved_and_aggregated_exactly():
    db = _session()
    product = Product(code="1".zfill(18), name="Repeated batch fixture", cost=1)
    db.add(product)
    db.flush()
    body = _response("1004", [_item("0001", "2.1250"), _item("0001", "3.3750")])
    payload = parse_balance_response(body, requested_plant="1004", effective_date=date(2026, 10, 14))
    rows = canonical_rows_from_sap(db, payload)
    assert len(payload.items) == 2
    assert len(rows) == 1
    assert rows[0].quantity == Decimal("5.5000")
    assert rows[0].material == "0001"
    assert rows[0].raw["batch_row_count"] == 2
    assert [row["source_row_number"] for row in rows[0].raw["sap_rows"]] == [1, 2]
    assert [row["quantity"] for row in rows[0].raw["sap_rows"]] == ["2.1250", "3.3750"]


def test_sap_multiple_batches_for_one_material_publish_one_product_row():
    db = _session()
    material = "0000000001005213"
    product = Product(code=material, name="Batch aggregation fixture", cost=1)
    db.add(product)
    db.flush()
    payload = parse_balance_response(
        _response(
            "1004",
            [
                _item(material, "0.1001", charge="batch-a"),
                _item(material, "0.2002", charge="batch-b"),
                _item(material, "0.3003", charge="batch-c"),
            ],
        ),
        requested_plant="1004",
        effective_date=date(2026, 10, 14),
    )
    rows = canonical_rows_from_sap(db, payload)
    assert [(row.product_id, row.quantity) for row in rows] == [(product.id, Decimal("0.6006"))]
    assert [row["charge"] for row in rows[0].raw["sap_rows"]] == ["batch-a", "batch-b", "batch-c"]


def test_catalog_reconciliation_classifies_only_verified_identifiers(tmp_path):
    db = _session()
    db.add_all(
        [
            Product(code="000000000000000101", name="Exact product", cost=1),
            Product(code="000000000000000999", name="Potential duplicate", cost=1),
        ]
    )
    db.commit()
    payload = parse_balance_response(
        _response(
            "1004",
            [
                {**_item("000000000000000101", "1"), "materialText": "Exact product"},
                {**_item("000000000000000202", "2"), "materialText": "Potential duplicate"},
                {**_item("000000000000000303", "3"), "materialText": "Absent product"},
            ],
        ),
        requested_plant="1004",
        effective_date=date(2026, 10, 14),
    )
    report, unresolved = mapping_report(db, payload)
    assert report["classificationCounts"] == {
        "exactVerifiedMatches": 1,
        "authoritativeIdentifierMatches": 0,
        "absentFromCurrentCatalogIdentifiers": 1,
        "potentialDuplicatesRequiringManualVerification": 1,
        "unresolvedOrAmbiguous": 0,
    }
    assert report["publicationSafe"] is False
    assert {row["sap_material_id"] for row in unresolved} == {
        "000000000000000202",
        "000000000000000303",
    }
    potential = next(row for row in unresolved if row["sap_material_id"].endswith("202"))
    assert potential["potential_matching_catalog_identifier"] == "000000000000000999"

    output = tmp_path / "reconciliation.csv"
    write_reconciliation_csv(output, unresolved)
    assert output.read_text(encoding="utf-8-sig").count("\n") == 3


def test_sap_canonicalization_rejects_ambiguous_internal_code_normalization():
    db = _session()
    db.add_all(
        [
            Product(code="000000000000000123", name="First", cost=1),
            Product(code="legacy-123", name="Second", cost=1),
        ]
    )
    db.commit()
    payload = parse_balance_response(
        _response("1004", [_item("000000000000000123", "1")]),
        requested_plant="1004",
        effective_date=date(2026, 10, 14),
    )
    with pytest.raises(ValueError, match="ambiguous internal product mappings"):
        canonical_rows_from_sap(db, payload)


@pytest.mark.parametrize(
    "payload, message",
    [
        (
            SapBalancePayload(
                plant="1004",
                effective_date=date(2026, 10, 14),
                items=(SapBalanceItem("0001", "1005", "1000", "", Decimal("1"), {}),),
                raw_row_count=1,
            ),
            "plant mismatch",
        ),
        (
            SapBalancePayload(
                plant="1004",
                effective_date=date(2026, 10, 14),
                items=(SapBalanceItem("0001", "1004", "4000", "", Decimal("1"), {}),),
                raw_row_count=1,
            ),
            "LGORT mismatch",
        ),
        (
            SapBalancePayload(
                plant="1004",
                effective_date=date(2026, 10, 14),
                items=(),
                raw_row_count=1,
            ),
            "incomplete",
        ),
    ],
)
def test_sap_canonicalization_revalidates_scope_and_completeness(payload, message):
    with pytest.raises(ValueError, match=message):
        canonical_rows_from_sap(_session(), payload)


@pytest.mark.parametrize("quantity", ["NaN", "Infinity", "-Infinity"])
def test_sap_parser_rejects_non_finite_quantities(quantity):
    with pytest.raises(ValueError, match="invalid endingStockQuantity"):
        parse_balance_response(
            _response("1004", [_item("0001", quantity)]),
            requested_plant="1004",
            effective_date=date(2026, 10, 14),
        )


def test_fetch_balance_reports_timeout_without_retry(monkeypatch):
    monkeypatch.setenv("SAP_BALANCE_ENABLED", "true")
    monkeypatch.setenv("SAP_BALANCE_ALLOW_PIQ_REQUESTS", "true")
    monkeypatch.setenv("SAP_BALANCE_USERNAME", "local-user")
    monkeypatch.setenv("SAP_BALANCE_PASSWORD", "local-pass")
    monkeypatch.setenv("SAP_BALANCE_COMPANY_CODE", "LOCAL")
    monkeypatch.setenv("SAP_BALANCE_TIMEOUT_SECONDS", "240")
    get_settings.cache_clear()
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        raise httpx.ReadTimeout("offline timeout fixture", request=request)

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await fetch_balance(
                get_settings(),
                plant="1004",
                date_from=date(2026, 10, 14),
                date_to=date(2026, 10, 14),
                client=client,
            )

    try:
        with pytest.raises(TimeoutError, match="240 seconds"):
            asyncio.run(run())
        assert calls == 1
    finally:
        get_settings.cache_clear()


def test_fetch_balance_rejects_malformed_json(monkeypatch):
    monkeypatch.setenv("SAP_BALANCE_ENABLED", "true")
    monkeypatch.setenv("SAP_BALANCE_ALLOW_PIQ_REQUESTS", "true")
    monkeypatch.setenv("SAP_BALANCE_USERNAME", "local-user")
    monkeypatch.setenv("SAP_BALANCE_PASSWORD", "local-pass")
    monkeypatch.setenv("SAP_BALANCE_COMPANY_CODE", "LOCAL")
    get_settings.cache_clear()

    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=b"not-json")

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await fetch_balance(
                get_settings(),
                plant="1004",
                date_from=date(2026, 10, 14),
                date_to=date(2026, 10, 14),
                client=client,
            )

    try:
        with pytest.raises(ValueError, match="invalid JSON"):
            asyncio.run(run())
    finally:
        get_settings.cache_clear()


def test_sap_missing_product_fails_without_replacing_active_snapshot():
    db = _session()
    product = Product(code="1005213".zfill(18), name="Known", cost=1)
    db.add(product)
    db.flush()
    manual = _job(db, "Есик")
    first = publish_stock_snapshot(
        db,
        branch_id="Есик",
        source_type="manual",
        effective_date=date(2026, 10, 14),
        rows=[CanonicalStockRow(product.code, product.code, product.id, Decimal("32"), {})],
        import_job=manual,
    )
    db.commit()
    payload = parse_balance_response(
        _response("1004", [_item("0000000009999999", "10")]),
        requested_plant="1004",
        effective_date=date(2026, 10, 15),
    )
    result = import_sap_balance_payload(
        db,
        branch_id="Есик",
        payload=payload,
        user_name="test",
        manual_precedence_hours=0,
    )
    assert result.status == "error"
    status = db.scalar(select(ReferenceUpdateStatus).where(ReferenceUpdateStatus.branch_id == "Есик"))
    assert status.active_stock_snapshot_id == first.id
    assert db.scalar(select(BranchStock.stock).where(BranchStock.branch_id == "Есик")) == Decimal("32.0000")


def test_sap_absence_carries_forward_active_stock_without_breaking_idempotency():
    db = _session()
    branch_id = "carry-forward-test"
    returned = Product(code="1005213".zfill(18), name="Returned", cost=1)
    absent = Product(code="1017769".zfill(18), name="Absent", cost=1)
    db.add_all([returned, absent])
    db.flush()
    manual = _job(db, branch_id)
    manual_snapshot = publish_stock_snapshot(
        db,
        branch_id=branch_id,
        source_type="manual",
        effective_date=date(2026, 10, 14),
        rows=[
            CanonicalStockRow(returned.code, returned.code, returned.id, Decimal("32"), {"manual": True}),
            CanonicalStockRow(absent.code, absent.code, absent.id, Decimal("3"), {"manual": True}),
        ],
        import_job=manual,
    )
    db.commit()
    payload = parse_balance_response(
        _response("1004", [_item(returned.code, "6")]),
        requested_plant="1004",
        effective_date=date(2026, 10, 15),
    )
    first_job = import_sap_balance_payload(
        db,
        branch_id=branch_id,
        payload=payload,
        user_name="worker",
        manual_precedence_hours=0,
    )
    first_status = db.scalar(select(ReferenceUpdateStatus).where(ReferenceUpdateStatus.branch_id == branch_id))
    first_sap_snapshot_id = first_status.active_stock_snapshot_id
    assert first_job.status == "success"
    assert first_sap_snapshot_id != manual_snapshot.id
    assert dict(
        db.execute(
            select(BranchStock.product_id, BranchStock.stock).where(BranchStock.branch_id == branch_id)
        ).all()
    ) == {returned.id: Decimal("6.0000"), absent.id: Decimal("3.0000")}
    first_sap_snapshot = db.get(StockSnapshot, first_sap_snapshot_id)
    metadata = json.loads(first_sap_snapshot.metadata_json)
    assert metadata["carried_forward_product_count"] == 1
    assert metadata["absent_material_policy"] == "carry_forward_pending_business_decision"

    second_job = import_sap_balance_payload(
        db,
        branch_id=branch_id,
        payload=payload,
        user_name="retry",
        manual_precedence_hours=0,
    )
    second_status = db.scalar(select(ReferenceUpdateStatus).where(ReferenceUpdateStatus.branch_id == branch_id))
    assert second_job.status == "success"
    assert second_status.active_stock_snapshot_id == first_sap_snapshot_id


def test_manual_csv_uses_same_atomic_snapshot_publication_path():
    db = _session()
    result = import_reference_excel(
        db=db,
        data_type="stock",
        branch_ids=["Есик"],
        content="Материал;Остаток\n1005213;32,75\n".encode("utf-8"),
        filename="Остаток_Есик.csv",
        user_name="manual-test",
    )
    assert result.status == "success"
    status = db.scalar(select(ReferenceUpdateStatus).where(ReferenceUpdateStatus.branch_id == "Есик"))
    assert status.active_source_type == "manual"
    assert status.active_stock_snapshot_id is not None
    assert db.scalar(select(BranchStock.stock).where(BranchStock.branch_id == "Есик")) == Decimal("32.7500")


def test_newer_manual_snapshot_blocks_scheduled_sap_publication():
    db = _session()
    product = Product(code="0000000001005213", name="Known", cost=1)
    db.add(product)
    db.flush()
    manual = _job(db, "Есик")
    first = publish_stock_snapshot(
        db,
        branch_id="Есик",
        source_type="manual",
        effective_date=date(2026, 10, 14),
        rows=[CanonicalStockRow(product.code, product.code, product.id, Decimal("32"), {})],
        import_job=manual,
    )
    db.commit()
    payload = parse_balance_response(
        _response("1004", [_item(product.code, "6")]),
        requested_plant="1004",
        effective_date=date(2026, 10, 15),
    )
    result = import_sap_balance_payload(
        db,
        branch_id="Есик",
        payload=payload,
        user_name="scheduler",
        manual_precedence_hours=24,
    )
    assert result.status == "error"
    assert "manual_snapshot_precedence" in result.error
    assert db.get(StockSnapshot, first.id).status == "active"


def test_identical_sap_retry_reuses_active_snapshot_without_rewriting_lineage():
    db = _session()
    product = Product(code="1005213".zfill(18), name="Known", cost=1)
    db.add(product)
    db.commit()
    payload = parse_balance_response(
        _response("1004", [_item(product.code, "32.75")]),
        requested_plant="1004",
        effective_date=date(2026, 10, 15),
    )
    first_job = import_sap_balance_payload(
        db, branch_id="Есик", payload=payload, user_name="worker", manual_precedence_hours=24
    )
    first_status = db.scalar(select(ReferenceUpdateStatus).where(ReferenceUpdateStatus.branch_id == "Есик"))
    first_snapshot_id = first_status.active_stock_snapshot_id
    second_job = import_sap_balance_payload(
        db, branch_id="Есик", payload=payload, user_name="retry", manual_precedence_hours=24
    )
    second_status = db.scalar(select(ReferenceUpdateStatus).where(ReferenceUpdateStatus.branch_id == "Есик"))
    assert first_job.status == second_job.status == "success", (first_job.error, second_job.error)
    assert second_status.active_stock_snapshot_id == first_snapshot_id
    assert db.scalar(select(func.count(StockSnapshot.id))) == 1


def test_generated_price_list_keeps_exact_historical_esik_snapshot_after_republish():
    db = _session()
    examples = {
        "1005213": (Decimal("32"), Decimal("6")),
        "1017769": (Decimal("3"), Decimal("1000")),
        "1003454": (Decimal("140"), Decimal("1473")),
        "1003450": (Decimal("4"), Decimal("27")),
        "1005967": (Decimal("31"), Decimal("83")),
        "1005754": (Decimal("18"), Decimal("16")),
    }
    products = []
    for sku in examples:
        product = Product(code=sku.zfill(18), name=sku, cost=1)
        db.add(product)
        products.append(product)
    db.flush()
    first_job = _job(db, "Есик")
    first = publish_stock_snapshot(
        db,
        branch_id="Есик",
        source_type="manual",
        effective_date=date(2026, 10, 14),
        rows=[
            CanonicalStockRow(product.code, product.code, product.id, examples[product.name][0], {})
            for product in products
        ],
        import_job=first_job,
    )
    pf = PriceFormat(code="GPL_1004_001", name="Esik", branch="Есик", reference_branch_id="Есик")
    db.add(pf)
    db.flush()
    price_list = PriceList(
        number="ГПЛ_1004_001_Есик_2026-10-15_wf1850",
        price_format_id=pf.id,
        run_reference_versions_json=json.dumps({"generationReconciliation": {"stock_snapshot_id": first.id}}),
    )
    db.add(price_list)
    db.flush()
    db.add_all(
        [
            CalculatedPrice(
                price_list_id=price_list.id,
                product_id=product.id,
                cost=Decimal("1"),
                base_price=Decimal("1"),
                final_price=Decimal("1"),
            )
            for product in products
        ]
    )
    db.commit()

    second_job = _job(db, "Есик", "sap")
    second = publish_stock_snapshot(
        db,
        branch_id="Есик",
        source_type="sap",
        effective_date=date(2026, 10, 15),
        rows=[
            CanonicalStockRow(product.code, product.code, product.id, examples[product.name][1], {})
            for product in products
        ],
        import_job=second_job,
    )
    db.commit()

    historical = _stock_quantities_for_price_list(
        db,
        pl=price_list,
        pf=pf,
        product_ids=[product.id for product in products],
    )
    assert {product.name: historical[product.id] for product in products} == {
        sku: expected for sku, (expected, _wrong) in examples.items()
    }
    assert compare_stock_quantities(
        {sku.zfill(18): expected for sku, (expected, _wrong) in examples.items()},
        {sku.zfill(18): wrong for sku, (_expected, wrong) in examples.items()},
    )["mismatchCount"] == 6
    assert compare_stock_quantities(
        {sku.zfill(18): expected for sku, (expected, _wrong) in examples.items()},
        {product.code: historical[product.id] for product in products},
    )["mismatchCount"] == 0

    admin = AppUser(username="stock-integrity-test", role="admin", is_active=True)
    historical_api = get_generated_price_list_items(
        price_list.number,
        q=None,
        zone=None,
        top_filter="all",
        sort=None,
        page=1,
        page_size=100,
        db=db,
        current_user=admin,
    )
    assert {row["sku"]: Decimal(str(row["stock"])) for row in historical_api["items"]} == {
        sku.zfill(18): expected for sku, (expected, _wrong) in examples.items()
    }

    current_price_list = PriceList(
        number="GPL_1004_001_Esik_2026-10-16_current",
        price_format_id=pf.id,
        run_reference_versions_json=json.dumps(
            {"generationReconciliation": {"stock_snapshot_id": second.id}}
        ),
    )
    db.add(current_price_list)
    db.flush()
    db.add_all(
        [
            CalculatedPrice(
                price_list_id=current_price_list.id,
                product_id=product.id,
                cost=Decimal("1"),
                base_price=Decimal("1"),
                final_price=Decimal("1"),
            )
            for product in products
        ]
    )
    db.commit()
    current_api = get_generated_price_list_items(
        current_price_list.number,
        q=None,
        zone=None,
        top_filter="all",
        sort=None,
        page=1,
        page_size=100,
        db=db,
        current_user=admin,
    )
    assert {row["sku"]: Decimal(str(row["stock"])) for row in current_api["items"]} == {
        sku.zfill(18): wrong for sku, (_expected, wrong) in examples.items()
    }

    response = export_price_list_xlsx(
        price_list.number,
        search=None,
        top_filter="all",
        top_sort="default",
        db=db,
    )

    async def response_bytes() -> bytes:
        return b"".join([chunk async for chunk in response.body_iterator])

    workbook = load_workbook(io.BytesIO(asyncio.run(response_bytes())), data_only=True)
    exported = {
        str(row[1].value): Decimal(str(row[3].value))
        for row in workbook.active.iter_rows(min_row=2)
    }
    workbook.close()
    assert exported == {
        sku.zfill(18): expected for sku, (expected, _wrong) in examples.items()
    }

    csv_response = export_price_list_csv(
        price_list.number,
        search=None,
        top_filter="all",
        top_sort="default",
        db=db,
    )

    async def response_text() -> str:
        chunks = [chunk async for chunk in csv_response.body_iterator]
        return "".join(chunk.decode("utf-8") if isinstance(chunk, bytes) else chunk for chunk in chunks)

    csv_rows = list(csv.reader(io.StringIO(asyncio.run(response_text()))))
    csv_exported = {row[1]: Decimal(row[3]) for row in csv_rows[1:]}
    assert csv_exported == {
        sku.zfill(18): expected for sku, (expected, _wrong) in examples.items()
    }

    other_branch_job = _job(db, "1", "manual")
    other_branch_snapshot = publish_stock_snapshot(
        db,
        branch_id="1",
        source_type="manual",
        effective_date=date(2026, 10, 15),
        rows=[
            CanonicalStockRow(product.code, product.code, product.id, Decimal("999"), {})
            for product in products
        ],
        import_job=other_branch_job,
    )
    price_list.run_reference_versions_json = json.dumps(
        {"generationReconciliation": {"stock_snapshot_id": other_branch_snapshot.id}}
    )
    db.commit()
    with pytest.raises(HTTPException, match="does not match generated price-list branch") as exc_info:
        _stock_quantities_for_price_list(
            db,
            pl=price_list,
            pf=pf,
            product_ids=[product.id for product in products],
        )
    assert exc_info.value.status_code == 409
